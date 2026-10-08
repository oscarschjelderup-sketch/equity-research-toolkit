"""Core helpers for the DCF template builder: styles, registry, writers."""
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string

FONT = "Arial"
NAVY, LBLUE, PALE = "003255", "81B0C0", "DCEBF2"
GREYBG, INFILL, FCFILL = "F2F2F2", "FEF4CE", None
TITLE, HISTGREY, KEYFILL, BANNER = "407061", "808080", "DCEBF2", "003255"
BLUE, GREEN, BLACK, GREY, WHITE = "0000FF", "008000", "000000", "7F7F7F", "FFFFFF"
OKGREEN, WARNAMB, ERRRED = "C6EFCE", "FFEB9C", "FFC7CE"

HC = list("EFGHIJK")          # historical columns (7)
FC = list("LMNOPQRS")         # forecast columns (8)
HIST_START = 0                # index of the first history column with data (set by build_model from the case module)
TC = "T"                      # terminal-year column
AC = HC + FC
NH, NF = len(HC), len(FC)

F_NUM = '#,##0_);(#,##0);"–"_)'
F_NUM1 = '#,##0.0_);(#,##0.0);"–"_)'
F_NOK = '#,##0.00_);(#,##0.00);"–"_)'
F_PCT = '0.0%_);(0.0%);"–"_)'
F_PCT2 = '0.00%_);(0.00%);"–"_)'
F_MULT = '0.0"x"_);(0.0"x");"–"_)'
F_MULT2 = '0.00"x"_);(0.00"x");"–"_)'
F_DATE = "dd.mm.yyyy"
F_FAC = "0.000"
F_INT = "0"
F_GEN = "General"

THIN = Side(style="thin", color="A6A6A6")
MED = Side(style="medium", color=NAVY)

ROWS = {}    # (sheet, key) -> row number
CELLS = {}   # (sheet, key) -> "C24"

PURE_LINK = re.compile(r"^=(?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!\$?[A-Z]{1,3}\$?\d+$")


def prev(col):
    return get_column_letter(column_index_from_string(col) - 1)


def nxt(col, k=1):
    return get_column_letter(column_index_from_string(col) + k)


def q(sheet):
    return f"'{sheet}'" if re.search(r"[^A-Za-z0-9_]", sheet) else sheet


def X(sheet, key, col, cur=None, absc=False, absr=False):
    """Reference to a time-series cell (sheet, row key, column)."""
    r = ROWS[(sheet, key)]
    a = f"{'$' if absc else ''}{col}{'$' if absr else ''}{r}"
    return a if sheet == cur else f"{q(sheet)}!{a}"


def XR(sheet, key, c1, c2, cur=None):
    """Absolute row range reference, e.g. Model!$L$32:$S$32."""
    r = ROWS[(sheet, key)]
    a = f"${c1}${r}:${c2}${r}"
    return a if sheet == cur else f"{q(sheet)}!{a}"


def S(sheet, key, cur=None):
    """Absolute reference to a scalar cell."""
    addr = CELLS[(sheet, key)]
    col = re.match(r"[A-Z]+", addr).group(0)
    a = f"${col}${addr[len(col):]}"
    return a if sheet == cur else f"{q(sheet)}!{a}"


def style(c, kind="calc", fmt=None, bold=False, italic=False, size=9, color=None, fill=None,
          align=None, wrap=False, indent=0):
    col = color
    if col is None:
        col = {"input": BLUE, "inputh": BLUE, "link": GREEN, "calc": BLACK, "label": BLACK, "note": GREY,
               "head": WHITE}.get(kind, BLACK)
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=col)
    if kind == "input":
        c.fill = PatternFill("solid", fgColor=INFILL)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    if align or wrap or indent:
        c.alignment = Alignment(horizontal=align, vertical="center" if wrap else None, wrap_text=wrap, indent=indent)


def auto_kind(v):
    if isinstance(v, str) and v.startswith("="):
        return "link" if PURE_LINK.match(v) else "calc"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return "input"
    return "label"


PLAIN_SHEETS = {"Output", "Analysis", "Dashboard", "Deck_Feed", "Football"}   # link-only pages shown in black


def put(ws, addr, v, kind=None, **kw):
    c = ws[addr]
    c.value = v
    k = kind or auto_kind(v)
    if k == "link" and (ws.title in PLAIN_SHEETS or ws.title.endswith(">>>")):
        k = "calc"
    style(c, k, **kw)
    return c


def band(ws, row, c1, c2, text=None, fill=NAVY, color=WHITE, size=10, bold=True, height=None):
    """Section header band across columns c1..c2."""
    for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
        cell = ws.cell(row=row, column=ci)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    if text is not None:
        ws[f"{c1}{row}"].value = text
    if height:
        ws.row_dimensions[row].height = height


def top_border(ws, row, c1, c2, side=THIN):
    for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
        cell = ws.cell(row=row, column=ci)
        b = cell.border
        cell.border = Border(top=side, bottom=b.bottom, left=b.left, right=b.right)


def bottom_border(ws, row, c1, c2, side=THIN):
    for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
        cell = ws.cell(row=row, column=ci)
        b = cell.border
        cell.border = Border(bottom=side, top=b.top, left=b.left, right=b.right)


def sheet_title(ws, title, subtitle=None, company_ref=True, last_col="T"):
    put(ws, "B1", title, kind="label", bold=True, size=15, color=TITLE)
    ws.row_dimensions[1].height = 24
    if subtitle:
        put(ws, "B2", subtitle, kind="note", italic=True)
    if company_ref:
        put(ws, f"{last_col}1", f"={S('Inputs', 'company')}&\"  |  Scenario: \"&{S('Inputs', 'scenario')}",
            kind="calc", bold=True, color=NAVY, align="right")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90


def fill_row(ws, row, c1, c2, color):
    for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
        ws.cell(row=row, column=ci).fill = PatternFill("solid", fgColor=color)


def MREF(key, col, absolute=False):
    """Reference a row by key wherever it lives (Model first, then Analysis / Output)."""
    for sh in ("Model", "Analysis", "Output"):
        if (sh, key) in ROWS:
            r = ROWS[(sh, key)]
            return f"{sh}!${col}${r}" if absolute else f"{sh}!{col}{r}"
    raise KeyError(key)


def MROW(key):
    for sh in ("Model", "Analysis", "Output"):
        if (sh, key) in ROWS:
            return sh, ROWS[(sh, key)]
    raise KeyError(key)


# ---------------------------------------------------------------- time series
class TS:
    """Row-spec driven time-series sheet (years across columns E..S, TY in T)."""

    def __init__(self, name, start_row=7, last_col="T"):
        self.name, self.start, self.last_col = name, start_row, last_col
        self.spec = []

    def section(self, title):
        self.spec.append(("section", title))
        self.spec.append(("spacer",))

    def blank(self):
        self.spec.append(("blank",))

    def sub(self, title):
        """Campari-style sub-header: bold label with a line underneath."""
        self.spec.append(("sub", title))

    def row(self, key, label, unit="", fmt=F_NUM, hist=None, fc=None, ty=None, bold=False,
            italic=False, line=False, hist_from=0, input_fc=False, note=None, indent=0, key_row=False,
            check=False):
        self.spec.append(("row", dict(key=key, label=label, unit=unit, fmt=fmt, hist=hist, fc=fc,
                                      ty=ty, bold=bold or key_row, italic=italic or check, line=line,
                                      hist_from=hist_from, input_fc=input_fc, note=note,
                                      indent=indent, key_row=key_row, check=check)))

    def layout(self):
        r = self.start
        for e in self.spec:
            if e[0] == "row":
                ROWS[(self.name, e[1]["key"])] = r
            r += 1
        self.end = r - 1

    def write(self, ws):
        r = self.start
        for e in self.spec:
            if e[0] == "section":
                band(ws, r, "B", self.last_col, e[1], fill=NAVY, size=9)
                ws.row_dimensions[r].height = 14.5
            elif e[0] == "spacer":
                ws.row_dimensions[r].height = 5
            elif e[0] == "blank":
                ws.row_dimensions[r].height = 8
            elif e[0] == "sub":
                lab = e[1]() if callable(e[1]) else e[1]
                put(ws, f"B{r}", lab, kind="calc" if str(lab).startswith("=") else "label", bold=True)
                bottom_border(ws, r, "B", self.last_col, THIN)
            elif e[0] == "row":
                d = e[1]
                lab = d["label"]() if callable(d["label"]) else d["label"]
                tone = GREY if d["check"] else None
                if d["key_row"]:
                    fill_row(ws, r, "B", self.last_col, KEYFILL)
                put(ws, f"B{r}", lab, kind="label" if not str(lab).startswith("=") else "calc",
                    bold=d["bold"], italic=d["italic"], indent=d["indent"], color=tone,
                    fill=KEYFILL if d["key_row"] else None)
                put(ws, f"C{r}", d["unit"], kind="note", italic=True, align="center",
                    fill=KEYFILL if d["key_row"] else None)
                for i, c in enumerate(HC):
                    if d["hist"] is None or i < d["hist_from"] + HIST_START:
                        continue
                    v = d["hist"](c, prev(c), i)
                    if v is None:
                        continue
                    if HIST_START and isinstance(v, str) and v.startswith("="):
                        if i == HIST_START and re.search(rf"(?<![A-Z!]){prev(c)}\d+", v):
                            continue                      # the previous year is blank – no growth or change for the first year
                        if not v.startswith("=IFERROR("):
                            v = '=IFERROR(' + v[1:] + ',"")'
                    kind = "inputh" if auto_kind(v) == "input" else None
                    put(ws, f"{c}{r}", v, kind=kind, fmt=d["fmt"], bold=d["bold"], italic=d["italic"], color=tone,
                        fill=KEYFILL if d["key_row"] else None, align="right")
                for i, c in enumerate(FC):
                    if d["fc"] is None:
                        continue
                    v = d["fc"](c, prev(c), i)
                    if v is None:
                        continue
                    put(ws, f"{c}{r}", v, fmt=d["fmt"], bold=d["bold"], italic=d["italic"], color=tone,
                        fill=KEYFILL if d["key_row"] else None, align="right")
                if d["ty"] is not None:
                    v = d["ty"]()
                    if v is not None:
                        put(ws, f"{TC}{r}", v, fmt=d["fmt"], bold=d["bold"], italic=d["italic"],
                            fill=KEYFILL if d["key_row"] else None)
                if d["note"]:
                    put(ws, f"V{r}", d["note"], kind="note", italic=True)
                if d["line"]:
                    top_border(ws, r, "B", self.last_col)
            r += 1


def ts_header(ws, sheet, hist=True, fc=True, ty=False, last_col="T"):
    """Rows 3-5: numeric years, Historical/Forecast band and year labels."""
    band(ws, 5, "B", "D", fill=NAVY, size=9)
    put(ws, "B5", f"=\"(\"&{S('Inputs', 'unit')}&\")\"", kind="calc", bold=True, color=WHITE, fill=NAVY)
    ws.row_dimensions[5].height = 15
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[6].height = 5
    for i, c in enumerate(HC):
        if hist:
            put(ws, f"{c}3", f"={S('Inputs', 'first_hist')}+{i}", kind="calc", fmt=F_INT, color=GREY, size=7,
                align="right")
            put(ws, f"{c}5", f"={c}3&\"A\"", kind="calc", bold=True, align="right", color=WHITE, fill=HISTGREY)
    for i, c in enumerate(FC):
        if fc:
            put(ws, f"{c}3", f"={S('Inputs', 'first_fc')}+{i}", kind="calc", fmt=F_INT, color=GREY, size=7,
                align="right")
            put(ws, f"{c}5", f"={c}3&\"E\"", kind="calc", bold=True, align="right", color=WHITE, fill=NAVY)
    if ty:
        put(ws, f"{TC}5", "TV year", kind="label", bold=True, align="right", color=WHITE, fill=HISTGREY)
    if hist:
        ws.merge_cells(f"{HC[0]}4:{HC[-1]}4")
        put(ws, f"{HC[0]}4", "Historicals", kind="label", bold=True, align="center")
        bottom_border(ws, 4, HC[0], HC[-1], Side(style="thin", color="000000"))
    if fc:
        ws.merge_cells(f"{FC[0]}4:{FC[-1]}4")
        put(ws, f"{FC[0]}4", "Explicit forecast", kind="label", bold=True, align="center")
        bottom_border(ws, 4, FC[0], FC[-1], Side(style="thin", color="000000"))
    if ty:
        put(ws, f"{TC}4", "Terminal", kind="label", bold=True, align="center")
        bottom_border(ws, 4, TC, TC, Side(style="thin", color="000000"))
    if last_col == "T":
        put(ws, "V5", "Notes", kind="label", bold=True, color=WHITE, fill=NAVY)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 2
    for c in AC + [TC]:
        ws.column_dimensions[c].width = 9.5
    ws.column_dimensions["U"].width = 1.5
    ws.column_dimensions["V"].width = 58
    ws.freeze_panes = "E6"


# ---------------------------------------------------------------- scalar sheets
class Scalars:
    """Label / value / unit / note list (values in column C by default)."""

    def __init__(self, name, start_row=4, vcol="C", ucol="D", ncol="E", lcol="B", thin=True):
        self.name, self.start, self.vcol, self.thin = name, start_row, vcol, thin
        self.ucol, self.ncol, self.lcol = ucol, ncol, lcol
        self.spec = []

    def section(self, t):
        self.spec.append(("section", t))
        self.spec.append(("spacer",))

    def blank(self):
        self.spec.append(("blank",))

    def item(self, key, label, value, fmt=F_GEN, unit="", note="", kind=None, bold=False,
             italic=False, color=None):
        self.spec.append(("item", dict(key=key, label=label, value=value, fmt=fmt, unit=unit,
                                       note=note, kind=kind, bold=bold, italic=italic,
                                       color=color)))

    def layout(self):
        r = self.start
        for e in self.spec:
            if e[0] == "item":
                CELLS[(self.name, e[1]["key"])] = f"{self.vcol}{r}"
            r += 1
        self.end = r - 1

    def write(self, ws, band_to="E"):
        r = self.start
        for e in self.spec:
            if e[0] == "section":
                band(ws, r, self.lcol, band_to, e[1], size=9)
                ws.row_dimensions[r].height = 14.5
            elif e[0] == "spacer" and self.thin:
                ws.row_dimensions[r].height = 5
            elif e[0] == "blank" and self.thin:
                ws.row_dimensions[r].height = 8
            elif e[0] == "item":
                d = e[1]
                v = d["value"]
                if callable(v):
                    v = v()
                put(ws, f"{self.lcol}{r}", d["label"], kind="label", bold=d["bold"],
                    italic=d["italic"], color=d["color"])
                kind = d["kind"] or auto_kind(v)
                if v is None:
                    style(ws[f"{self.vcol}{r}"], "input", fmt=d["fmt"])
                elif kind == "text":
                    put(ws, f"{self.vcol}{r}", v, kind="input", fmt=d["fmt"], align="right")
                else:
                    put(ws, f"{self.vcol}{r}", v, kind=kind, fmt=d["fmt"], bold=d["bold"],
                        italic=d["italic"], color=d["color"], align="right")
                if self.ucol:
                    put(ws, f"{self.ucol}{r}", d["unit"], kind="note", italic=True)
                if self.ncol:
                    put(ws, f"{self.ncol}{r}", d["note"], kind="note", italic=True)
            r += 1


def fx(sheet, expr):
    """Formula template -> builder(c, p, i).
    {k}  current column, row k on `sheet`      [k]  previous column, row k on `sheet`
    <k>  Drivers live row k, current column    %k%  Hist row k, current column
    ~k~  Model row k, current column           @k   Inputs scalar k
    """
    def build(c, p, i):
        s = re.sub(r"\{(\w+)\}", lambda m: f"{c}{ROWS[(sheet, m.group(1))]}", expr)
        s = re.sub(r"\[(\w+)\]", lambda m: f"{p}{ROWS[(sheet, m.group(1))]}", s)
        s = re.sub(r"<(\w+)>", lambda m: f"Drivers!{c}{ROWS[('Drivers', m.group(1))]}", s)
        s = re.sub(r"%(\w+)%", lambda m: f"Hist!{c}{ROWS[('Hist', m.group(1))]}", s)
        s = re.sub(r"~(\w+)~", lambda m: f"Model!{c}{ROWS[('Model', m.group(1))]}", s)
        s = re.sub(r"\^(\w+)\^", lambda m: f"Model!{p}{ROWS[('Model', m.group(1))]}", s)
        s = re.sub(r"@(\w+)", lambda m: S("Inputs", m.group(1)), s)
        return "=" + s
    return build

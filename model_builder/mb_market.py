"""Market view sheets.

Reverse_DCF – what the share price implies (EBIT margin, growth, WACC, terminal growth), solved without VBA or Goal Seek:
              margin by a closed form (the DCF is linear in a uniform margin shift), WACC and terminal growth by bisection
              on the closed-form DCF, growth by interpolation in Sensitivity table 5 (an Excel data table).
Consensus   – consensus inputs, our estimates vs. consensus and our variant perception ('where we differ').
Thesis      – scenario stories and kill criteria with a live status on the latest actuals.
"""
from datetime import date
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from mb_core import *
from mb_val import SENS, T5_STEPS, _sens_value_formula, t4_col
import case_data as ex

I = lambda k: S("Inputs", k)
D = lambda k: S("DCF", k)


def MS(key, col):
    return f"Model!${col}${ROWS[('Model', key)]}"


YR = lambda col: f"Model!${col}$5"
PCT_SIGN = '+0.0%;-0.0%;0.0%'

# ============================================================ REVERSE DCF
RD_ROWS = dict(m_last=6, m_avg=7, cagr=8, m_at_g=9, wacc=10, g=11)
PRICE_ROW, DCF_ROW, UP_ROW = 13, 14, 15
LINE_ROWS = (18, 19, 20)
MS_ROW = 23            # margin solver 23..28
GS_ROW = 31            # growth solver 31..37
WS_TOP = 40            # WACC solver: header 40, bounds 41-43, steps 44.., result
TS_TOP = 87            # terminal growth solver: header 87, bounds 88-90, steps 91.., result
NSTEP = 40
CH_ROW = 22            # chart data header (I:K), points 23..39
N5 = len(T5_STEPS)


def _rd_layout():
    for k, r in RD_ROWS.items():
        for suf, col in [("", "C"), ("_ours", "D"), ("_diff", "E"), ("_ref", "F")]:
            CELLS[("Reverse_DCF", f"rd_{k}{suf}")] = f"{col}{r}"
    for j, r in enumerate(LINE_ROWS):
        CELLS[("Reverse_DCF", f"rd_line{j + 1}")] = f"B{r}"
    CELLS[("Reverse_DCF", "rd_k")] = f"C{MS_ROW + 3}"
    CELLS[("Reverse_DCF", "rd_dm")] = f"C{MS_ROW + 4}"
    CELLS[("Reverse_DCF", "rd_dm_check")] = f"C{MS_ROW + 5}"
    CELLS[("Reverse_DCF", "rd_dg")] = f"C{GS_ROW + 4}"
    CELLS[("Reverse_DCF", "rd_wacc_ok")] = f"C{WS_TOP + 3}"
    CELLS[("Reverse_DCF", "rd_g_ok")] = f"C{TS_TOP + 3}"


def _bisection(ws, top, title, var_label, lo0, hi0, value_of):
    """Rows: band top-1, header top, lower bound top+1, upper bound top+2, bracket flag top+3, steps, result.
    value_of(cell) -> formula for the DCF value per share when the variable sits in `cell`."""
    P = I("price")
    band(ws, top - 1, "B", "G", title)
    for col, txt in zip("BCDEFG", ["Step", "Low", "High", "Mid", "Value (mid)", "Value (low)"]):
        put(ws, f"{col}{top}", txt, kind="label", bold=True, align="right" if col != "B" else None)
    bottom_border(ws, top, "B", "G")
    lo, hi, flag = top + 1, top + 2, top + 3
    put(ws, f"B{lo}", f"Lower bound for the {var_label}", kind="label")
    put(ws, f"C{lo}", lo0, fmt=F_PCT2)
    put(ws, f"F{lo}", value_of(f"C{lo}"), fmt=F_NOK)
    put(ws, f"B{hi}", f"Upper bound for the {var_label}", kind="label")
    put(ws, f"C{hi}", hi0, fmt=F_PCT2)
    put(ws, f"F{hi}", value_of(f"C{hi}"), fmt=F_NOK)
    put(ws, f"B{flag}", "Share price inside the bracket (1 = yes)", kind="label")
    put(ws, f"C{flag}", f"=IF((F{lo}-{P})*(F{hi}-{P})<=0,1,0)", fmt=F_INT)
    s0 = flag + 1
    for k in range(NSTEP):
        r = s0 + k
        put(ws, f"B{r}", k + 1, kind="calc", fmt=F_INT, color=GREY)
        if k == 0:
            put(ws, f"C{r}", f"=C{lo}", fmt=F_PCT2)
            put(ws, f"D{r}", f"=C{hi}", fmt=F_PCT2)
            put(ws, f"G{r}", f"=F{lo}", fmt=F_NOK)
        else:
            same = f"SIGN(F{r - 1}-{P})=SIGN(G{r - 1}-{P})"
            put(ws, f"C{r}", f"=IF({same},E{r - 1},C{r - 1})", fmt=F_PCT2)
            put(ws, f"D{r}", f"=IF({same},D{r - 1},E{r - 1})", fmt=F_PCT2)
            put(ws, f"G{r}", f"=IF({same},F{r - 1},G{r - 1})", fmt=F_NOK)
        put(ws, f"E{r}", f"=(C{r}+D{r})/2", fmt='0.0000%')
        put(ws, f"F{r}", value_of(f"E{r}"), fmt=F_NOK)
    ws.row_dimensions.group(s0, s0 + NSTEP - 1, outline_level=1, hidden=True)
    res = s0 + NSTEP
    put(ws, f"B{res}", f"Market-implied {var_label} (steps hidden – click + to expand)", kind="label", bold=True)
    put(ws, f"C{res}", f'=IF(C{flag}=1,E{res - 1},"outside range")', fmt=F_PCT2, bold=True, fill=PALE)
    top_border(ws, res, "B", "G")
    return res


def write_reverse_dcf(wb):
    ws = wb["Reverse_DCF"]
    sheet_title(ws, "Reverse DCF – what the share price implies",
                "Solves the DCF for today's share price: which margin, growth, WACC or terminal growth would justify it? "
                "All other assumptions are kept at the base case.", last_col="K")
    for k, v in dict(A=2, B=60, C=14, D=14, E=13, F=13, G=46, H=2, I=13, J=19, K=13).items():
        ws.column_dimensions[k].width = v
    P = I("price")
    rr, rm = ROWS[("Model", "is_rev")], ROWS[("Model", "ebit_adj_m")]
    # ---- A. results
    band(ws, 4, "B", "G", "A. What the share price implies – market-implied vs. our base case")
    for col, txt in zip("BCDEFG", ["Driver", "Market-implied", "Our base case", "Difference", "Reference", "  Comment"]):
        put(ws, f"{col}5", txt, kind="label", bold=True, align="right" if col in "CDEF" else None)
    bottom_border(ws, 5, "B", "G")
    m_rng = f"Model!${FC[0]}${rm}:${FC[-1]}${rm}"
    ms0 = MS_ROW
    wres, tres = WS_TOP + 3 + NSTEP + 1, TS_TOP + 3 + NSTEP + 1
    spec = [
        ("m_last", f'="EBIT adj. margin, "&{YR(FC[-1])}', f'=IF(ISNUMBER(C{ms0 + 4}),{MS("ebit_adj_m", FC[-1])}+C{ms0 + 4},"n.a.")',
         f"={MS('ebit_adj_m', FC[-1])}", f"={MS('ebit_adj_m', HC[-1])}", f'="Last actual year ("&{YR(HC[-1])}&")"', F_PCT),
        ("m_avg", "EBIT adj. margin, average of the forecast years", f'=IF(ISNUMBER(C{ms0 + 4}),AVERAGE({m_rng})+C{ms0 + 4},"n.a.")',
         f"=AVERAGE({m_rng})", f"=AVERAGE(Model!${HC[-3]}${rm}:${HC[-1]}${rm})", "Average of the last three actual years", F_PCT),
        ("cagr", f'="Revenue CAGR, "&{YR(HC[-1])}&"–"&{YR(FC[-1])}', f"=C{GS_ROW + 5}",
         f"=({MS('is_rev', FC[-1])}/{MS('is_rev', HC[-1])})^(1/{NF})-1",
         f"=({MS('is_rev', HC[-1])}/{MS('is_rev', HC[-4])})^(1/3)-1", "CAGR over the last three actual years", F_PCT),
        ("m_at_g", "… EBIT adj. margin in the final year at that growth", f"=C{GS_ROW + 6}",
         f"={MS('ebit_adj_m', FC[-1])}", None, "Less growth means less operating leverage", F_PCT),
        ("wacc", "WACC", f"=C{wres}", f"={S('WACC', 'wacc')}", None, "All other assumptions unchanged", F_PCT2),
        ("g", "Terminal growth rate", f"=C{tres}", f"={S('Drivers', 'tg')}", None, "All other assumptions unchanged", F_PCT2),
    ]
    for key, lab, imp, ours, ref, com, fmt in spec:
        r = RD_ROWS[key]
        put(ws, f"B{r}", lab, kind="calc" if str(lab).startswith("=") else "label", bold=key in ("m_last", "cagr"),
            indent=1 if key == "m_at_g" else 0)
        put(ws, f"C{r}", imp, fmt=fmt, bold=True, fill=PALE, align="right")
        put(ws, f"D{r}", ours, fmt=fmt, align="right")
        put(ws, f"E{r}", f'=IF(ISNUMBER(C{r}),C{r}-D{r},"n.a.")', fmt=PCT_SIGN, italic=True, align="right")
        if ref:
            put(ws, f"F{r}", ref, fmt=fmt, align="right")
        put(ws, f"G{r}", com, kind="calc" if str(com).startswith("=") else "note", italic=True, indent=1)
    put(ws, f"B{PRICE_ROW}", "Share price (NOK)", kind="label")
    put(ws, f"C{PRICE_ROW}", f"={P}", fmt=F_NOK)
    put(ws, f"B{DCF_ROW}", "DCF value per share, base case (NOK)", kind="label")
    put(ws, f"C{DCF_ROW}", f"={D('dcf_ps')}", fmt=F_NOK)
    put(ws, f"B{UP_ROW}", "Upside to the DCF value", kind="label")
    put(ws, f"C{UP_ROW}", f"={D('dcf_up')}", fmt=F_PCT)
    put(ws, f"G{PRICE_ROW}", "Differences are in percentage points. Each driver is solved on its own.", kind="note", italic=True)
    # ---- B. conclusion
    band(ws, 17, "B", "G", "B. Conclusion (for the pitch)")
    r6, r8, r9, r10, r11 = (RD_ROWS[k] for k in ("m_last", "cagr", "m_at_g", "wacc", "g"))
    lines = [
        f'=IF(ISNUMBER(C{r6}),"At NOK "&FIXED(C{PRICE_ROW},0)&" the market prices in an EBIT adj. margin of "&FIXED(C{r6}*100,1)'
        f'&"% in "&{YR(FC[-1])}&" – our base case is "&FIXED(D{r6}*100,1)&"%, and the company earned "&FIXED(F{r6}*100,1)&"% in "'
        f'&{YR(HC[-1])}&".","EBIT margin: the share price is outside the solver range.")',
        f'=IF(ISNUMBER(C{r8}),"Or revenue growth of "&FIXED(C{r8}*100,1)&"% a year to "&{YR(FC[-1])}&" (our base case "'
        f'&FIXED(D{r8}*100,1)&"%), which with our cost base also means an EBIT margin of "&FIXED(C{r9}*100,1)&"%.",'
        f'"Growth: the share price is outside the range of Sensitivity table 5.")',
        f'=IF(AND(ISNUMBER(C{r10}),ISNUMBER(C{r11})),"Or a WACC of "&FIXED(C{r10}*100,1)&"% (ours "&FIXED(D{r10}*100,1)'
        f'&"%), or terminal growth of "&FIXED(C{r11}*100,1)&"% (ours "&FIXED(D{r11}*100,1)&"%).",'
        f'"WACC / terminal growth: the share price is outside the solver range.")',
    ]
    for r, f in zip(LINE_ROWS, lines):
        put(ws, f"B{r}", f, bold=r == LINE_ROWS[0], color=NAVY, wrap=True, align="left")
        ws.merge_cells(f"B{r}:G{r}")
        ws.row_dimensions[r].height = 27
    # ---- C. margin solver (closed form)
    band(ws, ms0 - 1, "B", "G", "C. Margin solver – the same shift in the EBIT margin in every forecast year and the terminal year "
                                "(the DCF is linear in it, so the solution is exact)")
    rev = XR("DCF", "rev", FC[0], FC[-1])
    tax = XR("Drivers", "tax", FC[0], FC[-1])
    share = XR("DCF", "share", FC[0], FC[-1])
    df = XR("DCF", "df", FC[0], FC[-1])
    rT = f"DCF!${TC}${ROWS[('DCF', 'rev')]}"
    rS = f"DCF!${FC[-1]}${ROWS[('DCF', 'rev')]}"
    taxS = f"Drivers!${FC[-1]}${ROWS[('Drivers', 'tax')]}"
    g, ronic = D("tv_g"), D("tv_ronic")
    msolver = [
        ("PV of the explicit FCFF per 1.00 shift in the EBIT margin (NOKm)", f"=SUMPRODUCT({rev},1-{tax},{share},{df})", F_NUM),
        ("Terminal-year FCFF per 1.00 shift (NOKm)",
         f'=IF({I("tv_method")}=1,{rT}*(1-{S("Drivers", "lt_tax")})*(1-{g}/{ronic})*IF({S("Drivers", "ty_margin_in")}="",1,0),'
         f"{rS}*(1-{taxS})*(1+{g}))", F_NUM),
        ("PV of the terminal value per 1.00 shift (NOKm)",
         f"={I('w_gordon')}*C{ms0 + 1}/({D('tv_wacc')}-{g})*{D('tv_dfg')}+{I('w_exit')}*{rS}*{D('tv_mult')}*{D('tv_dfe')}", F_NUM),
        ("DCF value per share per 1 percentage point of margin (NOK)", f"=(C{ms0}+C{ms0 + 2})/{I('shares_dil')}/100", F_NOK),
        ("Market-implied shift in the EBIT margin (percentage points)",
         f'=IF(C{ms0 + 3}=0,"n.a.",({P}-{D("dcf_ps")})/C{ms0 + 3}/100)', '+0.00%;-0.00%;0.00%'),
        ("Check: DCF value at the implied shift (NOK) = share price",
         f'=IF(ISNUMBER(C{ms0 + 4}),{D("dcf_ps")}+C{ms0 + 4}*100*C{ms0 + 3},"n.a.")', F_NOK),
    ]
    for j, (lab, f, fmt) in enumerate(msolver):
        r = ms0 + j
        put(ws, f"B{r}", lab, kind="label", bold=j == 4)
        put(ws, f"C{r}", f, fmt=fmt, bold=j == 4, fill=PALE if j == 4 else None)
    put(ws, f"G{ms0 + 3}", "Same effect as the EBITDA margin override on Sensitivity (via other opex)", kind="note", italic=True)
    # ---- D. growth solver (data table on Sensitivity)
    band(ws, GS_ROW - 1, "B", "G", "D. Growth solver – the same shift in volume / like-for-like growth for all segments "
                                   "(interpolated in Sensitivity table 5)")
    c0, c1 = SENS["t5_rows"]
    rng = lambda col: f"Sensitivity!${col}${c0}:${col}${c1}"
    g0 = GS_ROW
    ok = f"AND(C{g0}>0,C{g0}<{N5})"
    gsolver = [
        ("Position of the share price in table 5 (row)", f"=IFERROR(MATCH({P},{rng('D')},1),0)", F_INT, None),
        ("Growth shift at the bracket (low / high)", f'=IF({ok},INDEX({rng("C")},C{g0}),"n.a.")', PCT_SIGN,
         f'=IF({ok},INDEX({rng("C")},C{g0}+1),"n.a.")'),
        ("DCF value at the bracket (low / high)", f'=IF({ok},INDEX({rng("D")},C{g0}),"n.a.")', F_NOK,
         f'=IF({ok},INDEX({rng("D")},C{g0}+1),"n.a.")'),
        ("Interpolation weight", f'=IF(ISNUMBER(C{g0 + 2}),({P}-C{g0 + 2})/(D{g0 + 2}-C{g0 + 2}),"n.a.")', F_FAC, None),
        ("Market-implied growth shift (percentage points a year)",
         f'=IF(ISNUMBER(C{g0 + 3}),C{g0 + 1}+C{g0 + 3}*(D{g0 + 1}-C{g0 + 1}),"outside range")', '+0.00%;-0.00%;0.00%', None),
        ("Revenue CAGR at the implied shift",
         f'=IF(ISNUMBER(C{g0 + 3}),INDEX({rng("E")},C{g0})+C{g0 + 3}*(INDEX({rng("E")},C{g0}+1)-INDEX({rng("E")},C{g0})),"n.a.")',
         F_PCT, None),
        ("EBIT adj. margin in the final year at the implied shift",
         f'=IF(ISNUMBER(C{g0 + 3}),INDEX({rng("F")},C{g0})+C{g0 + 3}*(INDEX({rng("F")},C{g0}+1)-INDEX({rng("F")},C{g0})),"n.a.")',
         F_PCT, None),
    ]
    for j, (lab, f, fmt, f2) in enumerate(gsolver):
        r = g0 + j
        put(ws, f"B{r}", lab, kind="label", bold=j == 4)
        put(ws, f"C{r}", f, fmt=fmt, bold=j == 4, fill=PALE if j == 4 else None)
        if f2:
            put(ws, f"D{r}", f2, fmt=fmt)
    put(ws, f"G{g0}", "Table 5 is an Excel data table – press F9 if it looks stale", kind="note", italic=True)
    # ---- E/F. bisection solvers
    wacc_val = lambda cell: _sens_value_formula(cell, S("Drivers", "tg"))
    g_val = lambda cell: _sens_value_formula(S("WACC", "wacc"), cell)
    _bisection(ws, WS_TOP, "E. WACC solver – bisection on the closed-form DCF (the same formula as Sensitivity table 1)",
               "WACC", f"={S('Drivers', 'tg')}+0.0025", 0.30, wacc_val)
    _bisection(ws, TS_TOP, "F. Terminal growth solver – bisection on the closed-form DCF", "terminal growth",
               -0.05, f"={S('WACC', 'wacc')}-0.0025", g_val)
    # ---- chart: DCF value vs. margin shift
    put(ws, f"I{CH_ROW}", "Margin shift", kind="label", bold=True, align="right")
    put(ws, f"J{CH_ROW}", "DCF value per share", kind="label", bold=True, align="right")
    put(ws, f"K{CH_ROW}", "Share price", kind="label", bold=True, align="right")
    for k in range(17):
        r = CH_ROW + 1 + k
        put(ws, f"I{r}", round(-0.05 + 0.005 * k, 4), fmt=PCT_SIGN, color=GREY)
        put(ws, f"J{r}", f"={D('dcf_ps')}+I{r}*100*$C${ms0 + 3}", fmt=F_NUM1)
        put(ws, f"K{r}", f"={P}", fmt=F_NUM1)
    ch = LineChart()
    ch.title = "DCF value per share vs. a uniform shift in the EBIT margin"
    ch.add_data(Reference(ws, min_col=10, max_col=11, min_row=CH_ROW, max_row=CH_ROW + 17), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=9, min_row=CH_ROW + 1, max_row=CH_ROW + 17))
    ch.series[0].graphicalProperties.line.solidFill = NAVY
    ch.series[0].graphicalProperties.line.width = 28000
    ch.series[1].graphicalProperties.line.solidFill = "C00000"
    ch.series[1].graphicalProperties.line.dashStyle = "dash"
    for s_ in ch.series:
        s_.smooth = False
        s_.marker.symbol = "none"
    ch.x_axis.number_format = PCT_SIGN
    ch.y_axis.number_format = "0"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines = None
    ch.legend.position = "t"
    ch.x_axis.tickLblSkip = 2
    ch.height, ch.width = 7.6, 15.5
    ws.add_chart(ch, "I4")
    ws.freeze_panes = "A4"


# ============================================================ CONSENSUS
METRICS = [("rev", "Revenue", "is_rev", F_NUM), ("ebitda", "EBITDA (reported, IFRS 16)", "ebitda", F_NUM),
           ("ebit", "EBIT (reported, IFRS 16)", "ebit", F_NUM), ("eps", "EPS (NOK)", "eps", F_NOK), ("dps", "DPS (NOK)", "dps", F_NOK)]
MLABEL = {"rev": "Revenue", "ebitda": "EBITDA", "ebit": "EBIT", "eps": "EPS", "dps": "DPS"}
GCOLS = [("C", "D", "E"), ("F", "G", "H"), ("I", "J", "K")]
CS_T0 = 14              # metric rows 14..18, EBIT margin row 19
CS_VP0 = 23             # variant perception: 3 topics x 7 rows
VP_FIELDS = [("topic", "Topic"), ("ours", "Our view"), ("cons", "Consensus view"), ("why", "Why we differ – the evidence"),
             ("when", "What proves it – and when"), ("num", "The number")]


def _cs_layout():
    for j, (m, *_) in enumerate(METRICS + [("ebitm", None, None, None)]):
        r = CS_T0 + j
        for y, cols in enumerate(GCOLS, start=1):
            for pre, col in zip(("our", "cons", "diff"), cols):
                CELLS[("Consensus", f"{pre}_{m}_{y}")] = f"{col}{r}"
    for k, addr in dict(cons_source="C5", cons_date="C6", cons_n="C7", cons_buy="C8", cons_hold="D8", cons_sell="E8",
                        cons_tp="C9", cons_tp_up="E9").items():
        CELLS[("Consensus", k)] = addr
    for j in range(len(ex.VARIANT)):
        for k, (f, _) in enumerate(VP_FIELDS):
            CELLS[("Consensus", f"vp{j + 1}_{f}")] = f"C{CS_VP0 + 7 * j + k}"
        CELLS[("Consensus", f"vp{j + 1}_metric")] = f"M{CS_VP0 + 7 * j + 5}"
        CELLS[("Consensus", f"vp{j + 1}_year")] = f"N{CS_VP0 + 7 * j + 5}"


def write_consensus(wb):
    ws = wb["Consensus"]
    sheet_title(ws, "Consensus vs. our estimates",
                "Type the latest consensus in the blue cells (reported IFRS 16 basis). Differences update automatically.",
                last_col="K")
    for k, v in dict(A=2, B=42, L=2).items():
        ws.column_dimensions[k].width = v
    for c in "CDEFGHIJK":
        ws.column_dimensions[c].width = 11.5
    C = ex.CONSENSUS
    band(ws, 4, "B", "K", "A. Consensus inputs")
    put(ws, "B5", "Source", kind="label")
    put(ws, "C5", C["source"], kind="input", align="left")
    ws.merge_cells("C5:G5")
    put(ws, "B6", "Date", kind="label")
    put(ws, "C6", date(*C["date"]), fmt=F_DATE, kind="input")
    put(ws, "B7", "Number of analysts", kind="label")
    put(ws, "C7", C["n"], fmt=F_INT)
    put(ws, "B8", "Recommendations: Buy / Hold / Sell", kind="label")
    for col, k in zip("CDE", ("buy", "hold", "sell")):
        put(ws, f"{col}8", C[k], fmt=F_INT)
    put(ws, "B9", "Consensus target price (NOK) – typed on Inputs", kind="label")
    put(ws, "C9", f"={I('cons_tp')}", fmt=F_NOK)
    put(ws, "D9", "Upside", kind="note", italic=True, align="right")
    put(ws, "E9", f"=C9/{I('price')}-1", fmt=F_PCT)
    # ---- B. table
    band(ws, 11, "B", "K", "B. Our estimates vs. consensus")
    put(ws, "B12", f'={S("Inputs", "unit")}', kind="calc", italic=True, color=GREY)
    for y, cols in enumerate(GCOLS):
        ws.merge_cells(f"{cols[0]}12:{cols[2]}12")
        put(ws, f"{cols[0]}12", f"=Model!{FC[y]}$5", bold=True, color=WHITE, fill=NAVY, align="center")
        for col, txt in zip(cols, ("Ours", "Consensus", "Difference")):
            put(ws, f"{col}13", txt, kind="label", bold=True, align="right")
    bottom_border(ws, 13, "B", "K")
    for j, (m, lab, key, fmt) in enumerate(METRICS):
        r = CS_T0 + j
        put(ws, f"B{r}", lab, kind="label", bold=m in ("ebit", "eps"))
        for y, cols in enumerate(GCOLS):
            put(ws, f"{cols[0]}{r}", f"={MS(key, FC[y])}", fmt=fmt)
            put(ws, f"{cols[1]}{r}", C[m][y], fmt=fmt)
            put(ws, f"{cols[2]}{r}", f'=IFERROR({cols[0]}{r}/{cols[1]}{r}-1,"n.a.")', fmt=PCT_SIGN, bold=True)
    r = CS_T0 + len(METRICS)
    put(ws, f"B{r}", "EBIT margin (reported)", kind="label", italic=True)
    rr, re_ = CS_T0, CS_T0 + 2
    for y, cols in enumerate(GCOLS):
        put(ws, f"{cols[0]}{r}", f"=IFERROR({cols[0]}{re_}/{cols[0]}{rr},0)", fmt=F_PCT, italic=True)
        put(ws, f"{cols[1]}{r}", f"=IFERROR({cols[1]}{re_}/{cols[1]}{rr},0)", fmt=F_PCT, italic=True)
        put(ws, f"{cols[2]}{r}", f"={cols[0]}{r}-{cols[1]}{r}", fmt=PCT_SIGN, italic=True)
    top_border(ws, r, "B", "K")
    diff_rng = ",".join(f"{cols[2]}{CS_T0}:{cols[2]}{CS_T0 + len(METRICS)}" for cols in GCOLS)
    for rng in diff_rng.split(","):
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0.02"],
                                                      fill=PatternFill("solid", fgColor=OKGREEN)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["-0.02"],
                                                      fill=PatternFill("solid", fgColor=ERRRED)))
    put(ws, f"B{r + 1}", "Green = we are more than 2% above consensus, red = more than 2% below. Margin row in percentage points.",
        kind="note", italic=True)
    # ---- C. variant perception
    band(ws, CS_VP0 - 2, "B", "K", "C. Where we differ from consensus – our variant perception")
    put(ws, f"M{CS_VP0 - 1}", "Metric", kind="label", bold=True, align="center")
    put(ws, f"N{CS_VP0 - 1}", "Year (1-3)", kind="label", bold=True, align="center")
    put(ws, f"O{CS_VP0 - 1}", "rev / ebitda / ebit / eps / dps – drives 'The number'", kind="note", italic=True)
    ws.column_dimensions["M"].width = 9
    ws.column_dimensions["N"].width = 9
    for j, v in enumerate(ex.VARIANT):
        topic, ours, cons, why, when, mkey, y = v
        top = CS_VP0 + 7 * j
        texts = [topic, ours, cons, why, when]
        for k, (f, lab) in enumerate(VP_FIELDS):
            r = top + k
            put(ws, f"B{r}", f"{j + 1}. {lab}" if k == 0 else lab, kind="label", bold=k == 0, color=NAVY if k == 0 else None,
                indent=0 if k == 0 else 1)
            if f == "num":
                put(ws, f"M{r}", mkey, kind="input", align="center")
                put(ws, f"N{r}", y, fmt=F_INT, align="center")
                keys = "{" + ",".join(f'"{m}"' for m, *_ in METRICS) + "}"
                labs = ",".join(f'"{MLABEL[m]}"' for m, *_ in METRICS)
                pos = f"MATCH($M{r},{keys},0)"
                diff = (f"INDEX(CHOOSE($N{r},$E${CS_T0}:$E${CS_T0 + 4},$H${CS_T0}:$H${CS_T0 + 4},$K${CS_T0}:$K${CS_T0 + 4}),{pos})")
                put(ws, f"C{r}", f'=IFERROR(CHOOSE({pos},{labs})&" "&INDEX(Model!${FC[0]}$5:${FC[2]}$5,$N{r})&": our estimate is "'
                                 f'&IF({diff}>=0,"+","")&FIXED({diff}*100,1)&"% vs. consensus","Check the metric and year in M:N")',
                    bold=True, color=NAVY)
            else:
                put(ws, f"C{r}", texts[k], kind="input", wrap=True, align="left", bold=k == 0)
                ws.row_dimensions[r].height = 26 if len(texts[k]) > 115 else 15
            ws.merge_cells(f"C{r}:K{r}")
    # ---- D. summary
    sr = CS_VP0 + 7 * len(ex.VARIANT)
    band(ws, sr, "B", "K", "D. Summary")
    for k, m in enumerate(("ebit", "eps")):
        r = sr + 1 + k
        put(ws, f"B{r}", f"Our {MLABEL[m]} vs. consensus (years 1 / 2 / 3)", kind="label")
        for y, col in enumerate("CDE"):
            put(ws, f"{col}{r}", f"={S('Consensus', f'diff_{m}_{y + 1}', 'Consensus')}", fmt=PCT_SIGN, bold=True)
    ws.freeze_panes = "A4"


# ============================================================ THESIS
SC_NAMES = ["bear", "base", "bull"]
TH_SC0 = 6              # scenario blocks: 5 rows each (6..20)
TH_KC0 = 24             # kill criteria rows 24..28
TH_CNT = TH_KC0 + len(ex.KILL) + 1


def _th_layout():
    for j, n in enumerate(SC_NAMES):
        r = TH_SC0 + 5 * j
        for k, col in zip(("prob", "value", "up", "tp", "rating"), "CDEFG"):
            CELLS[("Thesis", f"sc_{n}_{k}")] = f"{col}{r}"
        for k, off in zip(("story", "trigger", "signpost"), (1, 2, 3)):
            CELLS[("Thesis", f"sc_{n}_{k}")] = f"C{r + off}"
    for j in range(len(ex.KILL)):
        r = TH_KC0 + j
        for k, col in zip(("kpi", "dir", "th", "latest", "per", "base", "buf", "status", "basechk", "comment"), "BCDEFGHIJK"):
            CELLS[("Thesis", f"kc{j + 1}_{k}")] = f"{col}{r}"
    for k, off in zip(("n_broken", "n_watch", "n_incons"), (0, 1, 2)):
        CELLS[("Thesis", k)] = f"C{TH_CNT + off}"


def write_thesis(wb):
    ws = wb["Thesis"]
    sheet_title(ws, "Thesis tracker – scenario stories and kill criteria",
                "Stories and kill criteria are typed here (blue). Probabilities, values and actuals are linked from the model.",
                last_col="K")
    for k, v in dict(A=2, B=36, C=12, D=14, E=13, F=13, G=16, H=12, I=13, J=15, K=38).items():
        ws.column_dimensions[k].width = v
    P = I("price")
    t4 = lambda k, j: f"Sensitivity!${t4_col(k)}${SENS['t4_rows'][0] + j}"
    band(ws, 4, "B", "K", "A. Scenario stories – what has to happen, and what to watch")
    for col, txt in zip("BCDEFG", ["Scenario", "Probability", "DCF value", "vs. price", "Target price", "Rating"]):
        put(ws, f"{col}5", txt, kind="label", bold=True, align="right" if col != "B" else None)
    bottom_border(ws, 5, "B", "K")
    colors = {"bear": "C00000", "base": NAVY, "bull": "2E7D32"}
    for j, n in enumerate(SC_NAMES):
        r = TH_SC0 + 5 * j
        put(ws, f"B{r}", n.capitalize(), kind="label", bold=True, color=colors[n])
        put(ws, f"C{r}", f"={S('Drivers', f'adj_prob_{n}')}", fmt=F_PCT, align="right")
        put(ws, f"D{r}", f"={t4('dcf_ps', j)}", fmt=F_NOK, bold=True, align="right")
        put(ws, f"E{r}", f"=D{r}/{P}-1", fmt=PCT_SIGN, align="right")
        put(ws, f"F{r}", f"={t4('tp', j)}", fmt=F_NOK, align="right")
        put(ws, f"G{r}", f"={t4('rating', j)}", bold=True, align="right")
        top_border(ws, r, "B", "K")
        story = ex.SCENARIO_STORIES[n]
        for k, lab in enumerate(("Story", "What has to happen", "Signpost to watch")):
            rr = r + 1 + k
            put(ws, f"B{rr}", lab, kind="label", indent=1)
            put(ws, f"C{rr}", story[k], kind="input", wrap=True, align="left")
            ws.merge_cells(f"C{rr}:K{rr}")
            ws.row_dimensions[rr].height = 26 if len(story[k]) > 125 else 15
    # ---- kill criteria
    band(ws, TH_KC0 - 2, "B", "K", "B. Kill criteria – what would make us wrong (status on the latest actuals)")
    hdr = ["KPI", "Kill if", "Threshold", "Latest actual", "Period", "Base case (next year)", "Watch zone", "Status",
           "Base case check", "Source"]
    for col, txt in zip("BCDEFGHIJK", hdr):
        put(ws, f"{col}{TH_KC0 - 1}", txt, kind="label", bold=True, align="right" if col in "DEFGH" else "center" if col in "CIJ" else None,
            wrap=True)
    ws.row_dimensions[TH_KC0 - 1].height = 26
    bottom_border(ws, TH_KC0 - 1, "B", "K")
    dv = DataValidation(type="list", formula1='"below,above"', allow_blank=False)
    ws.add_data_validation(dv)
    for j, (lab, key, direction, th, buf, kind) in enumerate(ex.KILL):
        r = TH_KC0 + j
        fmt = F_PCT if kind == "pct" else F_MULT2
        put(ws, f"B{r}", lab, kind="input", align="left")
        put(ws, f"C{r}", direction, kind="input", align="center")
        dv.add(f"C{r}")
        put(ws, f"D{r}", th, fmt=fmt)
        put(ws, f"E{r}", f"={MS(key, HC[-1])}", fmt=fmt)
        put(ws, f"F{r}", f"={YR(HC[-1])}", align="right")
        put(ws, f"G{r}", f"={MS(key, FC[0])}", fmt=fmt)
        put(ws, f"H{r}", buf, fmt=fmt)
        put(ws, f"I{r}", f'=IF(ISNUMBER(E{r}),IF(C{r}="below",IF(E{r}<D{r},"BROKEN",IF(E{r}<D{r}+H{r},"WATCH","ON TRACK")),'
                         f'IF(E{r}>D{r},"BROKEN",IF(E{r}>D{r}-H{r},"WATCH","ON TRACK"))),"n.a.")', bold=True, align="center")
        put(ws, f"J{r}", f'=IF(ISNUMBER(G{r}),IF(C{r}="below",IF(G{r}<D{r},"INCONSISTENT","OK"),IF(G{r}>D{r},"INCONSISTENT","OK")),"n.a.")',
            bold=True, align="center")
        put(ws, f"K{r}", "Quarterly report / LTM", kind="input", align="left")
    last = TH_KC0 + len(ex.KILL) - 1
    for rng, pairs in [(f"I{TH_KC0}:I{last}", [("ON TRACK", OKGREEN), ("WATCH", WARNAMB), ("BROKEN", ERRRED)]),
                       (f"J{TH_KC0}:J{last}", [("OK", OKGREEN), ("INCONSISTENT", ERRRED)])]:
        for txt, colr in pairs:
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                          fill=PatternFill("solid", fgColor=colr)))
    for off, (lab, f) in enumerate([("Kill criteria breached on the latest actuals", f'=COUNTIF(I{TH_KC0}:I{last},"BROKEN")'),
                                    ("Kill criteria in the watch zone", f'=COUNTIF(I{TH_KC0}:I{last},"WATCH")'),
                                    ("Base case breaches a kill criterion", f'=COUNTIF(J{TH_KC0}:J{last},"INCONSISTENT")')]):
        r = TH_CNT + off
        put(ws, f"B{r}", lab, kind="label", bold=off == 0)
        put(ws, f"C{r}", f, fmt=F_INT, bold=True, align="right")
    put(ws, f"B{TH_CNT + 4}", "Latest actual links to the last actual year. Overwrite it with the latest reported quarter (LTM) "
                              "as new reports arrive – the Checks sheet flags any breach.", kind="note", italic=True)
    ws.freeze_panes = "A4"


# ============================================================ entry points
def layout():
    _rd_layout()
    _cs_layout()
    _th_layout()


def write_all(wb):
    write_reverse_dcf(wb)
    write_consensus(wb)
    write_thesis(wb)

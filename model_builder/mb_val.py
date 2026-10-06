"""Sensitivity, Comps and Football sheets."""
from openpyxl.formatting.rule import ColorScaleRule
from mb_core import *
import case_data as ex

# ------------------------------------------------------------ fixed layout (Sensitivity)
SENS = dict(t1_hdr=12, t1_rows=(13, 17), t2_hdr=22, t2_rows=(23, 27), t3_hdr=32, t3_rows=(33, 37),
            t4_lab=40, t4_frm=41, t4_rows=(42, 44), t4_pw=46,
            t5_band=51, t5_lab=53, t5_frm=54, t5_rows=(55, 87))
T5_STEPS = [round(-0.10 + 0.005 * k, 4) for k in range(33)]     # growth shift -10pp .. +6pp in 0.5pp steps
WCOLS = list("DEFGHIJ")          # 7 WACC columns
T3COLS = list("DEFGH")            # 5 margin columns
T4_OUT = []                      # (col, key, label, fmt) filled in layout

CELLS[("Sensitivity", "ovr_scen")] = "D4"
CELLS[("Sensitivity", "ovr_growth")] = "D5"
CELLS[("Sensitivity", "ovr_margin")] = "D6"
CELLS[("Sensitivity", "wacc_step")] = "C9"
CELLS[("Sensitivity", "g_step")] = "C10"
CELLS[("Sensitivity", "m_step")] = "C20"
CELLS[("Sensitivity", "center1")] = f"G{SENS['t1_rows'][0] + 2}"
CELLS[("Sensitivity", "center2")] = f"G{SENS['t2_rows'][0] + 2}"
CELLS[("Sensitivity", "center3")] = f"F{SENS['t3_rows'][0] + 2}"


def _sens_layout():
    cols = []
    from openpyxl.utils import get_column_letter as gl
    ci = 4  # D
    outs = [("cagr", "Revenue CAGR (last actual → final year)", F_PCT),
            ("rev_last", "Revenue, final forecast year (NOKm)", F_NUM),
            ("ebit_m", "EBIT adj. margin, final year", F_PCT),
            ("fcff_last", "FCFF, final forecast year (NOKm)", F_NUM),
            ("ev", "Enterprise value (NOKm)", F_NUM),
            ("dcf_ps", "DCF value per share (NOK)", F_NOK),
            ("dcf_up", "Upside to DCF value", F_PCT),
            ("fair", "Blended fair value (NOK)", F_NOK),
            ("tp", "Target price (NOK)", F_NOK),
            ("tp_up", "Upside to target price", F_PCT),
            ("rating", "Recommendation", F_GEN)]
    for i in range(NF):
        outs.append((f"rev{i + 1}", f"Revenue {i + 1}", F_NUM))
    for i in range(NF):
        outs.append((f"ebit{i + 1}", f"EBIT adj. {i + 1}", F_NUM))
    for k, lab, fmt in outs:
        col = gl(ci)
        T4_OUT.append((col, k, lab, fmt))
        CELLS[("Sensitivity", "sc_" + k)] = f"{col}{SENS['t4_rows'][0]}"   # bear row cell (base = +1, bull = +2)
        ci += 1


_sens_layout()


def t4_col(key):
    return [c for c, k, _, _ in T4_OUT if k == key][0]


def _sens_value_formula(w, g, method="gordon"):
    """Closed-form DCF value per share for WACC cell w and growth / multiple cell g."""
    fc = XR("DCF", "fcff_incl", FC[0], FC[-1])
    per = XR("DCF", "period", FC[0], FC[-1])
    S8 = f"DCF!$S${ROWS[('DCF', 'rev')]}"
    TYm = f"DCF!$T${ROWS[('DCF', 'ebit_m')]}"
    S20 = f"DCF!$S${ROWS[('DCF', 'fcff')]}"
    S10 = f"DCF!$S${ROWS[('DCF', 'ebitdaal')]}"
    tS = f"DCF!$S${ROWS[('DCF', 'period')]}"
    nS = f"DCF!$S${ROWS[('DCF', 'n')]}"
    stub = S("Inputs", "stub")
    bridge = f"SUM({S('DCF', 'b_nibd')}:{S('DCF', 'b_other').split('!')[1]})"
    shares = S("Inputs", "shares_dil")
    pv = f"SUMPRODUCT({fc},(1+{w})^(-{per}))"
    pv_exit_std = f"{S10}*{S('Drivers', 'exit')}*(1+{w})^(-({stub}+{nS}-1))"
    if method == "gordon":
        tfcf = (f"IF({S('Inputs', 'tv_method')}=1,{S8}*(1+{g})*{TYm}*(1-{S('Drivers', 'lt_tax')})"
                f"*(1-{g}/{S('Drivers', 'ronic')}),{S20}*(1+{g}))")
        pv_g = f"{tfcf}/({w}-{g})*(1+{w})^(-{tS})"
        tv = f"{S('Inputs', 'w_gordon')}*{pv_g}+{S('Inputs', 'w_exit')}*{pv_exit_std}"
    else:
        tv = f"{S10}*{g}*(1+{w})^(-({stub}+{nS}-1))"
    return f"=({pv}+{tv}+{bridge})/{shares}"


def write_sens(wb):
    ws = wb["Sensitivity"]
    sheet_title(ws, "Sensitivity & scenario analysis",
                "Tables 1-2 are live formulas. Tables 3-4 are Excel data tables (Data > What-If Analysis).",
                last_col="N")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 17
    from openpyxl.utils import get_column_letter as gl
    for i in range(4, 32):
        ws.column_dimensions[gl(i)].width = 11
    band(ws, 3, "B", "N", "Data-table input cells – must stay BLANK (Excel substitutes values here when tables recalculate)")
    for r, lab in [(4, "Scenario override (1-3)"), (5, "Revenue growth adjustment (pp p.a.)"),
                   (6, "EBITDA margin adjustment (pp)")]:
        put(ws, f"B{r}", lab, kind="label")
        c = ws[f"D{r}"]
        style(c, "calc", fmt=F_PCT if r > 4 else F_INT, fill=GREYBG)
    put(ws, "F4", "Do not type in the grey cells – they feed the data tables in sections 3 and 4.", kind="note", italic=True)

    # ---- table 1: WACC x g
    band(ws, 8, "B", "N", "1. DCF value per share (NOK) – WACC vs. terminal growth")
    put(ws, "B9", "WACC step", kind="label")
    put(ws, "C9", 0.005, fmt=F_PCT2)
    put(ws, "B10", "Terminal growth step", kind="label")
    put(ws, "C10", 0.0025, fmt=F_PCT2)
    h = SENS["t1_hdr"]
    wacc, tg = S("WACC", "wacc"), S("Drivers", "tg")
    put(ws, f"C{h}", "g  \\  WACC", kind="label", bold=True, align="center", fill=PALE)
    for k, c in enumerate(WCOLS):
        put(ws, f"{c}{h}", f"={wacc}+({k}-3)*$C$9", fmt=F_PCT2, bold=True, color=WHITE, fill=NAVY, align="center")
    r0, r1 = SENS["t1_rows"]
    for k, r in enumerate(range(r0, r1 + 1)):
        put(ws, f"C{r}", f"={tg}+({k}-2)*$C$10", fmt=F_PCT2, bold=True, color=WHITE, fill=NAVY, align="center")
        for c in WCOLS:
            put(ws, f"{c}{r}", _sens_value_formula(f"{c}${h}", f"$C{r}"), fmt=F_NOK, align="center")
    ws.conditional_formatting.add(f"D{r0}:J{r1}", ColorScaleRule(start_type="min", start_color="F8CBAD",
                                  mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                  end_type="max", end_color="C6E0B4"))
    ctr = ws[CELLS[("Sensitivity", "center1")]]
    ctr.font = Font(name=FONT, size=9, bold=True)
    ctr.border = Border(left=MED, right=MED, top=MED, bottom=MED)
    put(ws, f"L{h}", "Centre cell = model value (checked on the Checks sheet)", kind="note", italic=True)

    # ---- table 2: WACC x exit multiple
    band(ws, 19, "B", "N", "2. DCF value per share (NOK) – WACC vs. exit multiple (exit-multiple method only)")
    put(ws, "B20", "Exit multiple step", kind="label")
    put(ws, "C20", 1.0, fmt=F_MULT)
    h = SENS["t2_hdr"]
    put(ws, f"C{h}", "Multiple  \\  WACC", kind="label", bold=True, align="center", fill=PALE)
    for k, c in enumerate(WCOLS):
        put(ws, f"{c}{h}", f"={wacc}+({k}-3)*$C$9", fmt=F_PCT2, bold=True, color=WHITE, fill=NAVY, align="center")
    r0, r1 = SENS["t2_rows"]
    for k, r in enumerate(range(r0, r1 + 1)):
        put(ws, f"C{r}", f"={S('Drivers', 'exit')}+({k}-2)*$C$20", fmt=F_MULT, bold=True, color=WHITE,
            fill=NAVY, align="center")
        for c in WCOLS:
            put(ws, f"{c}{r}", _sens_value_formula(f"{c}${h}", f"$C{r}", method="exit"), fmt=F_NOK, align="center")
    ws.conditional_formatting.add(f"D{r0}:J{r1}", ColorScaleRule(start_type="min", start_color="F8CBAD",
                                  mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                  end_type="max", end_color="C6E0B4"))
    ctr = ws[CELLS[("Sensitivity", "center2")]]
    ctr.font = Font(name=FONT, size=9, bold=True)
    ctr.border = Border(left=MED, right=MED, top=MED, bottom=MED)

    # ---- table 3: data table growth x margin
    band(ws, 29, "B", "N", "3. DCF value per share (NOK) – change in revenue growth vs. change in EBITDA margin (data table)")
    put(ws, "B30", "Rows: change in annual volume growth (segment mode) or like-for-like growth (location mode) for all segments. "
                   "Columns: change in EBITDA margin (via other opex).", kind="note", italic=True)
    h = SENS["t3_hdr"]
    put(ws, f"C{h}", f"={S('DCF', 'dcf_ps')}", fmt='"growth ↓ margin →"', bold=True, fill=PALE, align="center")
    steps = [-0.02, -0.01, 0.0, 0.01, 0.02]
    for c, v in zip(T3COLS, steps):
        put(ws, f"{c}{h}", v, fmt='+0.0%;-0.0%;0.0%', bold=True, color=WHITE, fill=NAVY, align="center")
    r0, r1 = SENS["t3_rows"]
    for r, v in zip(range(r0, r1 + 1), steps):
        put(ws, f"C{r}", v, fmt='+0.0%;-0.0%;0.0%', bold=True, color=WHITE, fill=NAVY, align="center")
        for c in T3COLS:
            style(ws[f"{c}{r}"], "calc", fmt=F_NOK, align="center")
    ws.conditional_formatting.add(f"D{r0}:H{r1}", ColorScaleRule(start_type="min", start_color="F8CBAD",
                                  mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                  end_type="max", end_color="C6E0B4"))

    # ---- table 4: scenario data table
    band(ws, 39, "B", gl(3 + len(T4_OUT)), "4. Scenario analysis – Bear / Base / Bull (data table on the scenario switch)")
    lab, frm = SENS["t4_lab"], SENS["t4_frm"]
    put(ws, f"C{lab}", "Scenario", kind="label", bold=True, align="center")
    src = {
        "cagr": f"=(Model!$S${ROWS[('Model', 'is_rev')]}/Model!$K${ROWS[('Model', 'is_rev')]})^(1/{NF})-1",
        "rev_last": f"=Model!$S${ROWS[('Model', 'is_rev')]}",
        "ebit_m": f"=Model!$S${ROWS[('Model', 'ebit_adj_m')]}",
        "fcff_last": f"=DCF!$S${ROWS[('DCF', 'fcff')]}",
        "ev": f"={S('DCF', 'ev')}", "dcf_ps": f"={S('DCF', 'dcf_ps')}", "dcf_up": f"={S('DCF', 'dcf_up')}",
        "fair": f"={S('DCF', 'tp_fair')}", "tp": f"={S('DCF', 'tp')}", "tp_up": f"={S('DCF', 'tp_up')}",
        "rating": f"={S('DCF', 'rating')}",
    }
    for i, c in enumerate(FC):
        src[f"rev{i + 1}"] = f"=Model!${c}${ROWS[('Model', 'is_rev')]}"
        src[f"ebit{i + 1}"] = f"=Model!${c}${ROWS[('Model', 'ebit_adj')]}"
    for col, k, labtxt, fmt in T4_OUT:
        if k.startswith("rev") and k[3:].isdigit():
            labtxt = f'="Revenue "&Model!{FC[int(k[3:]) - 1]}$5'
        if k.startswith("ebit") and k[4:].isdigit():
            labtxt = f'="EBIT adj. "&Model!{FC[int(k[4:]) - 1]}$5'
        put(ws, f"{col}{lab}", labtxt, kind="calc" if str(labtxt).startswith("=") else "label", bold=True,
            align="center", wrap=True)
        put(ws, f"{col}{frm}", src[k], fmt=fmt, color=GREY, italic=True, align="center")
    ws.row_dimensions[lab].height = 36
    put(ws, f"B{frm}", "Live model (current scenario)", kind="note", italic=True)
    r0, r1 = SENS["t4_rows"]
    for j, (r, name) in enumerate(zip(range(r0, r1 + 1), ["Bear", "Base", "Bull"])):
        put(ws, f"B{r}", name, kind="label", bold=True)
        put(ws, f"C{r}", j + 1, fmt=F_INT, align="center")
        for col, k, labtxt, fmt in T4_OUT:
            style(ws[f"{col}{r}"], "calc", fmt=fmt, align="center", bold=(k in ("dcf_ps", "tp", "rating")))
    bottom_border(ws, lab, "B", gl(3 + len(T4_OUT)))
    pw = SENS["t4_pw"]
    prob = [S("Drivers", "adj_prob_bear"), S("Drivers", "adj_prob_base"), S("Drivers", "adj_prob_bull")]
    for col_key, lab_txt in [("dcf_ps", "Probability-weighted DCF value per share (NOK)"),
                             ("tp", "Probability-weighted target price (NOK)")]:
        col = t4_col(col_key)
        rr = pw if col_key == "dcf_ps" else pw + 1
        put(ws, f"B{rr}", lab_txt, kind="label", bold=True)
        put(ws, f"{col}{rr}", f"={col}{r0}*{prob[0]}+{col}{r0 + 1}*{prob[1]}+{col}{r0 + 2}*{prob[2]}",
            fmt=F_NOK, bold=True, fill=PALE)
        CELLS[("Sensitivity", "pw_" + col_key)] = f"{col}{rr}"
    put(ws, f"B{pw + 2}", "Probabilities are set on the Drivers sheet (section B).", kind="note", italic=True)

    # ---- table 5: growth shift -> DCF value, revenue CAGR and final-year margin (feeds the reverse DCF)
    b5, lab5, frm5 = SENS["t5_band"], SENS["t5_lab"], SENS["t5_frm"]
    c0, c1 = SENS["t5_rows"]
    band(ws, b5, "B", "N", "5. Reverse DCF input – value vs. a uniform shift in volume / like-for-like growth (data table)")
    put(ws, f"B{b5 + 1}", "Rows: change in annual growth for all segments. The Reverse_DCF sheet finds the shift that matches "
                          "the share price.", kind="note", italic=True)
    rr = ROWS[("Model", "is_rev")]
    heads = [("C", "Growth shift (pp p.a.)"), ("D", "DCF value per share (NOK)"), ("E", "Revenue CAGR"),
             ("F", f'="EBIT adj. margin "&Model!{FC[-1]}$5')]
    for col, txt in heads:
        put(ws, f"{col}{lab5}", txt, kind="calc" if txt.startswith("=") else "label", bold=True, align="center", wrap=True)
    ws.row_dimensions[lab5].height = 26
    bottom_border(ws, lab5, "B", "F")
    put(ws, f"B{frm5}", "Live model (current settings)", kind="note", italic=True)
    put(ws, f"D{frm5}", f"={S('DCF', 'dcf_ps')}", fmt=F_NOK, color=GREY, italic=True, align="center")
    put(ws, f"E{frm5}", f"=(Model!$S${rr}/Model!$K${rr})^(1/{NF})-1", fmt=F_PCT, color=GREY, italic=True, align="center")
    put(ws, f"F{frm5}", f"=Model!$S${ROWS[('Model', 'ebit_adj_m')]}", fmt=F_PCT, color=GREY, italic=True, align="center")
    for r, v in zip(range(c0, c1 + 1), T5_STEPS):
        put(ws, f"C{r}", v, fmt='+0.0%;-0.0%;0.0%', align="center")
        style(ws[f"D{r}"], "calc", fmt=F_NOK, align="center")
        style(ws[f"E{r}"], "calc", fmt=F_PCT, align="center")
        style(ws[f"F{r}"], "calc", fmt=F_PCT, align="center")
    ws.freeze_panes = "C3"


# ============================================================ COMPS
COMPS_PEER0 = 6
for j in range(8):
    ROWS[("Comps", f"peer{j + 1}")] = COMPS_PEER0 + j
CR = dict(mean=15, median=16, q1=17, q3=18, high=19, low=20, company=22, prem=23,
          yr=26, bridge=27, ihdr=28, evs=29, evebitda=30, evebit=31, pe=32, avg=33, vd_hdr=36, vd_co=37, vd_peer=38)
GROUPS = [("EV / Sales (x)", "FGH", "evs", F_MULT2), ("EV / EBITDA (x)", "IJK", "evebitda", F_MULT),
          ("EV / EBIT (x)", "LMN", "evebit", F_MULT), ("P / E (x)", "OPQ", "pe", F_MULT)]
CELLS[("Comps", "peer_val")] = f"H{CR['avg']}"
CELLS[("Comps", "peer_low")] = f"G{CR['avg']}"
CELLS[("Comps", "peer_high")] = f"I{CR['avg']}"
CELLS[("Comps", "year")] = f"C{CR['yr']}"
CELLS[("Comps", "bridge")] = f"C{CR['bridge']}"


def write_comps(wb):
    ws = wb["Comps"]
    sheet_title(ws, "Peer group & relative valuation",
                ex.MARKET.get("cons_note", "Source: Bloomberg consensus as provided in the case material (example data – replace). Multiples at current prices."),
                last_col="S")
    widths = dict(A=2, B=38, C=10, D=11, E=11, R=11, S=11)
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for c in "FGHIJKLMNOPQ":
        ws.column_dimensions[c].width = 8.5
    put(ws, "D4", "Mcap", kind="label", bold=True, align="right")
    put(ws, "E4", "EV", kind="label", bold=True, align="right")
    put(ws, "D5", "NOKm", kind="note", italic=True, align="right")
    put(ws, "E5", "NOKm", kind="note", italic=True, align="right")
    put(ws, "B5", "Company", kind="label", bold=True)
    put(ws, "C5", "Country", kind="label", bold=True)
    for title, cols, key, fmt in GROUPS:
        ws.merge_cells(f"{cols[0]}4:{cols[-1]}4")
        put(ws, f"{cols[0]}4", title, kind="label", bold=True, color=WHITE, fill=NAVY, align="center")
        for j, c in enumerate(cols):
            put(ws, f"{c}5", f"=Model!{FC[j]}$5", bold=True, align="right")
    put(ws, "R4", "EBIT margin", kind="label", bold=True, align="right")
    put(ws, "S4", "Rev. growth", kind="label", bold=True, align="right")
    put(ws, "R5", f"=Model!{FC[1]}$5", bold=True, align="right")
    put(ws, "S5", f"=Model!{FC[1]}$5", bold=True, align="right")
    bottom_border(ws, 5, "B", "S", MED)
    for j, pr in enumerate(ex.PEERS):
        r = COMPS_PEER0 + j
        name, ctry, mcap, ev, evs, eve, evb, pe, em, gr = pr[:10]
        put(ws, f"B{r}", name, kind="input")
        put(ws, f"C{r}", ctry, kind="input", align="center")
        put(ws, f"D{r}", mcap, fmt=F_NUM)
        put(ws, f"E{r}", ev, fmt=F_NUM)
        for (title, cols, key, fmt), vals in zip(GROUPS, [evs, eve, evb, pe]):
            for c, v in zip(cols, vals):
                put(ws, f"{c}{r}", v, fmt=fmt)
        put(ws, f"R{r}", em, fmt=F_PCT)
        put(ws, f"S{r}", gr, fmt=F_PCT)
    p0, p1 = COMPS_PEER0, COMPS_PEER0 + 7
    stats = [("mean", "Mean", "AVERAGE({0}{1}:{0}{2})"), ("median", "Median", "MEDIAN({0}{1}:{0}{2})"),
             ("q1", "25th percentile", "QUARTILE({0}{1}:{0}{2},1)"), ("q3", "75th percentile", "QUARTILE({0}{1}:{0}{2},3)"),
             ("high", "High", "MAX({0}{1}:{0}{2})"), ("low", "Low", "MIN({0}{1}:{0}{2})")]
    for key, lab, tmpl in stats:
        r = CR[key]
        put(ws, f"B{r}", lab, kind="label", bold=key == "median")
        for c in "DEFGHIJKLMNOPQRS":
            fmt = F_NUM if c in "DE" else (F_PCT if c in "RS" else (F_MULT2 if c in "FGH" else F_MULT))
            put(ws, f"{c}{r}", f'=IFERROR({tmpl.format(c, p0, p1)},"")', fmt=fmt, bold=key == "median",
                fill=PALE if key == "median" else None)      # blank when no peer has the multiple
    top_border(ws, CR["mean"], "B", "S")
    # subject company
    r = CR["company"]
    put(ws, f"B{r}", f"={S('Inputs', 'company')}&\" (at current price)\"", bold=True, color=NAVY)
    put(ws, f"C{r}", "", kind="label")
    put(ws, f"D{r}", f"={S('Inputs', 'mcap')}", fmt=F_NUM, bold=True)
    put(ws, f"E{r}", f"=IF({S('Inputs', 'peer_lease')}=1,{S('Inputs', 'ev_incl')},{S('Inputs', 'ev_ex')})", fmt=F_NUM, bold=True)
    for title, cols, key, fmt in GROUPS:
        for j, c in enumerate(cols):
            put(ws, f"{c}{r}", "=" + MREF(key, FC[j]), fmt=fmt, bold=True)
    put(ws, f"R{r}", "=" + MREF("pm_ebit_m", FC[1]), fmt=F_PCT, bold=True)
    put(ws, f"S{r}", f"=Model!{FC[1]}{ROWS[('Model', 'grev')]}", fmt=F_PCT, bold=True)
    top_border(ws, r, "B", "S", MED)
    r = CR["prem"]
    put(ws, f"B{r}", "Premium / (discount) vs. peer median", kind="label", italic=True)
    for c in "FGHIJKLMNOPQ":
        put(ws, f"{c}{r}", f'=IFERROR({c}{CR["company"]}/{c}{CR["median"]}-1,"n.m.")', fmt=F_PCT, italic=True)
    for c in "RS":
        put(ws, f"{c}{r}", f'=IFERROR({c}{CR["company"]}-{c}{CR["median"]},"n.m.")', fmt='+0.0%;-0.0%;0.0%', italic=True)
    put(ws, f"B{r + 1}", "Margin and growth columns show the difference in percentage points.", kind="note", italic=True)

    # ---- implied valuation
    band(ws, CR["yr"] - 1, "B", "S", "Implied valuation from peer multiples")
    r = CR["yr"]
    put(ws, f"B{r}", "Forecast year used (from Inputs)", kind="label")
    put(ws, f"C{r}", f"={S('Inputs', 'peer_year')}", fmt=F_INT)
    put(ws, f"D{r}", f"=INDEX(Model!${FC[0]}$5:${FC[1]}$5,1,C{r})", bold=True, indent=1, align="left")
    r = CR["bridge"]
    I = lambda k: S("Inputs", k)
    put(ws, f"B{r}", "EV-to-equity bridge on peer basis (NOKm)", kind="label")
    put(ws, f"C{r}", f"=-({I('nibd')}+IF({I('peer_lease')}=1,{I('leases')},0))-{I('minority')}+{I('associates')}-{I('pension')}+{I('other_adj')}",
        fmt=F_NUM)
    put(ws, f"E{r}", "Net debt (incl. leases if IFRS 16 basis), minorities, associates, pensions", kind="note", italic=True)
    h = CR["ihdr"]
    for c, txt in zip("BCDEFGHIJ", ["Multiple", "Company metric", "Peer 25th", "Peer median", "Peer 75th",
                                     "Value @25th", "Value @median", "Value @75th", "Include (1/0)"]):
        put(ws, f"{c}{h}", txt, kind="label", bold=True, align="right" if c != "B" else None, wrap=True)
    ws.row_dimensions[h].height = 26
    bottom_border(ws, h, "B", "J")
    yr = f"$C${CR['yr']}"
    shares = I("shares_dil")
    rows = [("evs", "EV / Sales", "pm_rev", "FG", 0, F_MULT2),
            ("evebitda", "EV / EBITDA", "pm_ebitda", "IJ", 1, F_MULT),
            ("evebit", "EV / EBIT", "pm_ebit", "LM", 1, F_MULT),
            ("pe", "P / E", "pm_eps", "OP", 1, F_MULT)]
    for key, lab, mkey, cols, inc, fmt in rows:
        r = CR[key]
        put(ws, f"B{r}", lab, kind="label", bold=True)
        msheet, mrow = MROW(mkey)
        put(ws, f"C{r}", f"=INDEX({msheet}!${FC[0]}${mrow}:${FC[1]}${mrow},1,{yr})", fmt=F_NOK if key == "pe" else F_NUM)
        for c, st in zip("DEF", ["q1", "median", "q3"]):
            put(ws, f"{c}{r}", f"=INDEX(${cols[0]}${CR[st]}:${cols[1]}${CR[st]},1,{yr})", fmt=fmt)
        for c, m in zip("GHI", "DEF"):
            if key == "pe":
                put(ws, f"{c}{r}", f'=IF({m}{r}="","",C{r}*{m}{r})', fmt=F_NOK)
            else:
                put(ws, f"{c}{r}", f'=IF({m}{r}="","",(C{r}*{m}{r}+$C${CR["bridge"]})/{shares})', fmt=F_NOK)
        put(ws, f"J{r}", inc, fmt=F_INT, align="center")
    r = CR["avg"]
    put(ws, f"B{r}", "Peer multiples value per share (average of included)", kind="label", bold=True)
    j0, j1 = CR["evs"], CR["pe"]
    for c in "GHI":
        n_inc = f"SUMPRODUCT(--ISNUMBER({c}{j0}:{c}{j1}),$J${j0}:$J${j1})"      # multiples without peer data are skipped
        put(ws, f"{c}{r}", f"=IF({n_inc}=0,0,SUMPRODUCT({c}{j0}:{c}{j1},$J${j0}:$J${j1})/{n_inc})",
            fmt=F_NOK, bold=True, fill=PALE)
    top_border(ws, r, "B", "J")
    put(ws, f"K{r}", "← feeds the target price (weight on Inputs) and the football field", kind="note", italic=True)

    # ---- value drivers vs peers
    band(ws, CR["vd_hdr"] - 1, "B", "S", "Company vs. peers – value drivers (second forecast year)")
    h = CR["vd_hdr"]
    for c, txt in zip("BCDEF", ["", "EBIT margin", "Revenue growth", "EV / EBIT", "P / E"]):
        put(ws, f"{c}{h}", txt, kind="label", bold=True, align="right", wrap=True)
    ws.row_dimensions[h].height = 26
    bottom_border(ws, h, "B", "F")
    co, pm = CR["vd_co"], CR["vd_peer"]
    put(ws, f"B{co}", f"={S('Inputs', 'company')}", bold=True)
    put(ws, f"B{pm}", "Peer median", kind="label", bold=True)
    for c, src in zip("CDEF", ["R", "S", "M", "P"]):
        fmt = F_PCT if c in "CD" else F_MULT
        put(ws, f"{c}{co}", f"={src}{CR['company']}", fmt=fmt)
        put(ws, f"{c}{pm}", f"={src}{CR['median']}", fmt=fmt)
    ws.freeze_panes = "C6"


# ============================================================ FOOTBALL
FB0 = 5
FB_METHODS = ["w52", "dcf_wg", "dcf_sc", "evebitda", "evebit", "pe", "evs"]
for j, k in enumerate(FB_METHODS):
    ROWS[("Football", k)] = FB0 + j
FB_CHART = (FB0, FB0 + len(FB_METHODS) + 2)      # chart helper rows (methods + 3 markers: price, fair value, target price)
CELLS[("Football", "price")] = "C13"
CELLS[("Football", "fair")] = "C14"
CELLS[("Football", "dcf")] = "C15"
CELLS[("Football", "tp")] = "C16"
CELLS[("Football", "cons")] = "C17"


def write_football(wb):
    ws = wb["Football"]
    sheet_title(ws, "Valuation summary – football field", "Ranges in NOK per share. The chart updates automatically.",
                last_col="H")
    for k, v in dict(A=2, B=40, C=11, D=11, E=11, F=11, G=46).items():
        ws.column_dimensions[k].width = v
    for c, txt in zip("BCDEFG", ["Valuation method", "Low", "High", "Spread", "Mid-point", "Basis"]):
        put(ws, f"{c}4", txt, kind="label", bold=True, align="right" if c in "CDEF" else None)
    bottom_border(ws, 4, "B", "G", MED)
    I = lambda k: S("Inputs", k)
    s1r0, s1r1 = SENS["t1_rows"]
    sc0, sc1 = SENS["t4_rows"]
    dcol = t4_col("dcf_ps")
    rows = {
        "w52": ("52-week trading range", f"={I('low52')}", f"={I('high52')}", "Market data"),
        "dcf_wg": ("DCF – WACC ±0.5pp / g ±0.25pp", f"=MIN(Sensitivity!$F${s1r0 + 1}:$H${s1r0 + 3})",
                   f"=MAX(Sensitivity!$F${s1r0 + 1}:$H${s1r0 + 3})", "Sensitivity table 1 (inner 3x3)"),
        "dcf_sc": ("DCF – Bear to Bull", f"=MIN(Sensitivity!${dcol}${sc0}:${dcol}${sc1})",
                   f"=MAX(Sensitivity!${dcol}${sc0}:${dcol}${sc1})", "Sensitivity table 4"),
        "evebitda": ("EV/EBITDA – peers (25th–75th pct)", f"=Comps!$G${CR['evebitda']}", f"=Comps!$I${CR['evebitda']}", "Comps"),
        "evebit": ("EV/EBIT – peers (25th–75th pct)", f"=Comps!$G${CR['evebit']}", f"=Comps!$I${CR['evebit']}", "Comps"),
        "pe": ("P/E – peers (25th–75th pct)", f"=Comps!$G${CR['pe']}", f"=Comps!$I${CR['pe']}", "Comps"),
        "evs": ("EV/Sales – peers (25th–75th pct)", f"=Comps!$G${CR['evs']}", f"=Comps!$I${CR['evs']}", "Comps"),
    }
    for k in FB_METHODS:
        r = ROWS[("Football", k)]
        lab, lo, hi, basis = rows[k]
        put(ws, f"B{r}", lab, kind="label")
        put(ws, f"C{r}", lo, fmt=F_NOK)
        put(ws, f"D{r}", hi, fmt=F_NOK)
        put(ws, f"E{r}", f"=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),MAX(0,D{r}-C{r}),0)", fmt=F_NOK)
        put(ws, f"F{r}", f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),AVERAGE(C{r}:D{r}),"")', fmt=F_NOK)
        put(ws, f"G{r}", basis, kind="note", italic=True, indent=1)
    put(ws, "B13", "Current share price", kind="label", bold=True)
    put(ws, "C13", f"={I('price')}", fmt=F_NOK, bold=True)
    put(ws, "B14", "Fair value today (blended DCF / peers)", kind="label", bold=True)
    put(ws, "C14", f"={S('DCF', 'tp_fair')}", fmt=F_NOK, bold=True, fill=PALE)
    put(ws, "G14", "Comparable with the ranges above – all are values today", kind="note", italic=True, indent=1)
    put(ws, "B15", "DCF value per share", kind="label")
    put(ws, "C15", f"={S('DCF', 'dcf_ps')}", fmt=F_NOK)
    put(ws, "B16", "12-month target price (shown separately)", kind="label")
    put(ws, "C16", f"={S('DCF', 'tp')}", fmt=F_NOK)
    put(ws, "G16", "Fair value rolled forward at the cost of equity, less the dividend – not a value today", kind="note", italic=True, indent=1)
    put(ws, "B17", "Consensus target price (if available)", kind="label")
    put(ws, "C17", f"={I('cons_tp')}", fmt=F_NOK)
    last = FB0 + len(FB_METHODS) - 1
    put(ws, "B18", "Chart baseline (hidden axis starts here)", kind="label")
    put(ws, "C18", f"=ROUNDDOWN(MIN(C{FB0}:C{last},C13:C16)*0.85/10,0)*10", fmt=F_NUM)
    put(ws, "B19", "Marker width for price / fair value / target price", kind="label")
    put(ws, "C19", f"=MAX(0.5,(MAX(D{FB0}:D{last},C13:C16)-C18)/90)", fmt=F_NOK)
    ws.column_dimensions["H"].width = 3
    ws.column_dimensions["I"].width = 44
    ws.column_dimensions["J"].width = 10
    ws.column_dimensions["K"].width = 10
    for c, txt in zip("IJK", ["Chart label", "Offset", "Bar"]):
        put(ws, f"{c}4", txt, kind="label", bold=True, align="right" if c != "I" else None)
    bottom_border(ws, 4, "I", "K", MED)
    for k in FB_METHODS:
        r = ROWS[("Football", k)]
        put(ws, f"I{r}", f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),B{r}&" ("&FIXED(C{r},0)&"–"&FIXED(D{r},0)&")",B{r}&" (n.a.)")')
        put(ws, f"J{r}", f"=IF(ISNUMBER(C{r}),MAX(0,C{r}-$C$18),0)", fmt=F_NOK)
        put(ws, f"K{r}", f"=E{r}", fmt=F_NOK)
    for k, (cell, lab) in enumerate([("C13", "Current share price"), ("C14", "Fair value today"),
                                     ("C16", "12m target price – rolled forward")]):
        r = last + 1 + k
        put(ws, f"I{r}", f'="{lab} ("&FIXED({cell},0)&")"')
        put(ws, f"J{r}", f"=MAX(0,{cell}-$C$18-$C$19/2)", fmt=F_NOK)
        put(ws, f"K{r}", "=$C$19", fmt=F_NOK)
    put(ws, "I19", "Helper columns for the Dashboard chart – bars start at the baseline in C18.", kind="note", italic=True)

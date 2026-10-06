"""Model (forecast), DCF and WACC sheets."""
from mb_core import *
import case_data as ex

# ============================================================ MODEL
MODEL = TS("Model", start_row=7)


def HL(hkey):
    return lambda c, p, i: f"=Hist!{c}{ROWS[('Hist', hkey)]}"


def _model_spec():
    t = MODEL
    m = lambda e: fx("Model", e)
    I = lambda k: S("Inputs", k)
    # ------------------------------------------------------------------ 1
    t.section("1. Revenue build – locations x volume per location x price/mix, per segment")
    R1, R2, R3 = (S("Drivers", k) for k in ("ramp1", "ramp2", "ramp3"))

    def meq(s_):
        """Mature-equivalent locations: last year + closures + this year's openings ramping up (Drivers F)."""
        def build(c, p, i):
            r = lambda k: ROWS[("Model", k)]
            c2 = prev(p)
            c3 = prev(c2)
            o = r(f"open_{s_}")
            return (f"={p}{r(f'meq_{s_}')}+MIN(0,{c}{r(f'net_{s_}')})+{c}{o}*{R1}+{p}{o}*({R2}-{R1})"
                    f"+{c2}{o}*({R3}-{R2})+{c3}{o}*(1-{R3})")
        return build

    for s_ in "abc":
        t.sub((lambda s_=s_: f'={I("seg_" + s_)}'))
        t.row(f"loc_{s_}", "Locations (end of year)", "#", fmt=F_NUM, hist=HL(f"loc_{s_}"),
              fc=m(f"[loc_{s_}]+{{net_{s_}}}"))
        t.row(f"net_{s_}", "Net new locations", "#", fmt=F_NUM, italic=True, indent=1, hist_from=1,
              hist=m(f"{{loc_{s_}}}-[loc_{s_}]"), fc=m(f"<net_{s_}>"))
        t.row(f"open_{s_}", "Openings still ramping up (forecast openings only)", "#", fmt=F_NUM, italic=True, indent=1,
              hist=lambda c, p, i: "=0", fc=m(f"MAX(0,{{net_{s_}}})"))
        t.row(f"meq_{s_}", "Mature-equivalent locations", "#", fmt=F_NUM1, hist=m(f"{{loc_{s_}}}"), fc=meq(s_),
              note="Location mode: existing locations + new ones weighted by the ramp-up curve (Drivers F)" if s_ == "a" else None)
        t.row(f"mpl_{s_}", (lambda: f'="Volume per mature location ("&{I("vol_unit")}&")"'), "#", fmt=F_NOK,
              hist=m(f"IF({{loc_{s_}}}=0,0,{{vol_{s_}}}/{{loc_{s_}}})"), fc=m(f"[mpl_{s_}]*(1+<lfl_{s_}>)"),
              note="Grows with like-for-like growth (Drivers A)" if s_ == "a" else None)
        t.row(f"glfl_{s_}", "Like-for-like growth", "%", fmt=F_PCT, italic=True, indent=1, hist_from=1,
              hist=m(f"IF([mpl_{s_}]=0,0,{{mpl_{s_}}}/[mpl_{s_}]-1)"), fc=m(f"IF([mpl_{s_}]=0,0,{{mpl_{s_}}}/[mpl_{s_}]-1)"))
        t.row(f"vol_{s_}", (lambda: f'="Volume ("&{I("vol_unit")}&")"'), "#",
              fmt=F_NUM1, hist=HL(f"vol_{s_}"), fc=m(f"IF(@rev_mode=2,{{meq_{s_}}}*{{mpl_{s_}}},[vol_{s_}]*(1+<vol_{s_}>))"),
              note="Location mode: mature-equivalent locations x volume per location. Segment mode: volume growth (Drivers A)"
              if s_ == "a" else None)
        t.row(f"gvol_{s_}", "Volume growth", "%", fmt=F_PCT, italic=True, indent=1, hist_from=1,
              hist=m(f"{{vol_{s_}}}/[vol_{s_}]-1"), fc=m(f"{{vol_{s_}}}/[vol_{s_}]-1"))
        t.row(f"rpu_{s_}", "Revenue per unit", "NOK '000", fmt=F_NOK,
              hist=m(f"IF({{vol_{s_}}}=0,0,{{rev_{s_}}}/{{vol_{s_}}})"), fc=m(f"IF({{vol_{s_}}}=0,0,{{rev_{s_}}}/{{vol_{s_}}})"))
        t.row(f"px_{s_}", "Price/mix growth", "%", fmt=F_PCT, italic=True, indent=1, hist_from=1,
              hist=m(f"(1+{{grev_{s_}}})/(1+{{gvol_{s_}}})-1"), fc=m(f"(1+{{grev_{s_}}})/(1+{{gvol_{s_}}})-1"))
        t.row(f"rev_{s_}", "Revenue", "NOKm", bold=True, line=True, hist=HL(f"rev_{s_}"),
              fc=m(f"IF([vol_{s_}]=0,[rev_{s_}]*(1+<vol_{s_}>),[rev_{s_}]*{{vol_{s_}}}/[vol_{s_}])*(1+<px_{s_}>)"))
        t.row(f"grev_{s_}", "Revenue growth", "%", fmt=F_PCT, italic=True, indent=1, hist_from=1,
              hist=m(f"{{rev_{s_}}}/[rev_{s_}]-1"), fc=m(f"{{rev_{s_}}}/[rev_{s_}]-1"))
        t.blank()
    t.sub("Group – locations")
    t.row("loc", "Locations (end of year)", "#", fmt=F_NUM, bold=True,
          hist=m("{loc_a}+{loc_b}+{loc_c}"), fc=m("{loc_a}+{loc_b}+{loc_c}"))
    t.row("open", "Openings still ramping up", "#", fmt=F_NUM, italic=True, indent=1,
          hist=lambda c, p, i: "=0", fc=m("{open_a}+{open_b}+{open_c}"))
    t.row("avg_loc", "Average number of locations", "#", fmt=F_NUM1, hist_from=1,
          hist=m("AVERAGE([loc],{loc})"), fc=m("AVERAGE([loc],{loc})"))
    t.row("loc_g", "Growth in the average number of locations", "%", fmt=F_PCT, italic=True, indent=1, hist_from=2,
          hist=m("{avg_loc}/[avg_loc]-1"), fc=m("{avg_loc}/[avg_loc]-1"),
          note="Location mode: drives fixed costs (Drivers C) and rent per location (section 2)")
    t.row("rev_loc", "Revenue per average location", "NOKm", fmt=F_NUM1, hist_from=1,
          hist=m("IF({avg_loc}=0,0,{rev}/{avg_loc})"), fc=m("IF({avg_loc}=0,0,{rev}/{avg_loc})"))
    t.blank()
    t.row("rev", "Total revenue", "NOKm", key_row=True,
          hist=m("{rev_a}+{rev_b}+{rev_c}"), fc=m("{rev_a}+{rev_b}+{rev_c}"))
    t.row("grev", "Revenue growth", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=m("{rev}/[rev]-1"), fc=m("{rev}/[rev]-1"))
    vc = "([rev_a]*{gvol_a}+[rev_b]*{gvol_b}+[rev_c]*{gvol_c})/[rev]"
    t.row("g_volc", "– of which volume", "%", fmt=F_PCT, italic=True, indent=1, hist_from=1, hist=m(vc), fc=m(vc))
    cap = lambda k: "(" + "+".join(f"IF([{k}_{x}]=0,0,[rev_{x}]*({{{k}_{x}}}/[{k}_{x}]-1))" for x in "abc") + ")/[rev]"
    t.row("g_newc", "of which new locations (capacity)", "%", fmt=F_PCT, italic=True, indent=2, hist_from=1,
          hist=m(cap("loc")), fc=m(f"IF(@rev_mode=2,{cap('meq')},0)"),
          note="History: growth in locations (incl. acquired). Forecast: growth in mature-equivalent locations (location mode)")
    t.row("g_lflc", "of which like-for-like volume", "%", fmt=F_PCT, italic=True, indent=2, hist_from=1,
          hist=m("{g_volc}-{g_newc}"), fc=m("{g_volc}-{g_newc}"))
    t.row("g_pxc", "– of which price/mix", "%", fmt=F_PCT, italic=True, indent=1, hist_from=1,
          hist=m("{grev}-{g_volc}"), fc=m("{grev}-{g_volc}"))
    t.blank()
    # ------------------------------------------------------------------ 2
    t.section("2. Earnings build – from revenue to net profit")
    t.row("is_rev", "Total revenue", "NOKm", bold=True, hist=m("{rev}"), fc=m("{rev}"))
    t.row("cogs", "Cost of goods sold", "NOKm", hist=HL("cogs"), fc=m("-{is_rev}*<cogs>"))
    t.row("gp", "Gross profit", "NOKm", bold=True, line=True, hist=m("{is_rev}+{cogs}"), fc=m("{is_rev}+{cogs}"))
    t.row("gm", "Gross margin", "%", fmt=F_PCT, italic=True, hist=m("{gp}/{is_rev}"), fc=m("{gp}/{is_rev}"))
    t.blank()
    DR = lambda k: S("Drivers", k)
    fixed = lambda key, sh: (lambda c, p, i: f"={c}{ROWS[('Model', key)]}*{DR(sh)}")
    t.row("pers_fix", "Personnel expenses – fixed part", "NOKm", italic=True, indent=1, hist=fixed("pers", "pers_fixsh"),
          fc=m("[pers_fix]*(1+<cpi>)*(1+<fix_real>)"),
          note="History: total x fixed share (Drivers E). Forecast: grows with cost inflation and capacity (Drivers C) – not with revenue")
    t.row("pers_var", "Personnel expenses – variable part", "NOKm", italic=True, indent=1,
          hist=m("{pers}-{pers_fix}"), fc=m("-{is_rev}*<pers_var>"), note="Scales with revenue (Drivers A)")
    t.row("pers", "Personnel expenses", "NOKm", hist=HL("pers"), fc=m("{pers_fix}+{pers_var}"))
    t.row("oth_fix", "Other opex – fixed part", "NOKm", italic=True, indent=1, hist=fixed("oth", "oth_fixsh"),
          fc=m("[oth_fix]*(1+<cpi>)*(1+<fix_real>)"))
    t.row("oth_var", "Other opex – variable part", "NOKm", italic=True, indent=1,
          hist=m("{oth}-{oth_fix}"), fc=m("-{is_rev}*<oth_var>"))
    t.row("oth", "Other operating expenses", "NOKm", hist=HL("oth"), fc=m("{oth_fix}+{oth_var}"),
          note="Operating leverage is a result, not an input: fixed costs grow slower than revenue")
    t.row("ebitda", "EBITDA (reported, IFRS 16)", "NOKm", bold=True, line=True,
          hist=m("{gp}+{pers}+{oth}"), fc=m("{gp}+{pers}+{oth}"))
    t.row("ebitda_m", "EBITDA margin", "%", fmt=F_PCT, italic=True, hist=m("{ebitda}/{is_rev}"), fc=m("{ebitda}/{is_rev}"))
    t.blank()
    t.row("lease", "Lease payments (IFRS 16 adjustment)", "NOKm", hist=HL("lease_pay"),
          fc=m("IF(@rev_mode=2,[lease]*(1+<cpi>)*(1+{loc_g}),-{is_rev}*<lease>)"),
          note="Rent. Location mode: rent per location grows with cost inflation; segment mode: % of revenue. An operating cost in the DCF")
    t.row("ebitdaal", "EBITDAaL (pre-IFRS 16)", "NOKm", bold=True, line=True,
          hist=m("{ebitda}+{lease}"), fc=m("{ebitda}+{lease}"))
    t.row("ebitdaal_m", "EBITDAaL margin", "%", fmt=F_PCT, italic=True,
          hist=m("{ebitdaal}/{is_rev}"), fc=m("{ebitdaal}/{is_rev}"))
    t.row("da", "D&A – owned assets", "NOKm", hist=HL("da"), fc=m("-{is_rev}*<da>"))
    t.row("ebit_adj", "EBIT adj. (pre-IFRS 16, excl. special items)", "NOKm", key_row=True,
          hist=m("{ebitdaal}+{da}"), fc=m("{ebitdaal}+{da}"), note="Basis for NOPAT and FCFF in the DCF")
    t.row("ebit_adj_m", "EBIT adj. margin", "%", fmt=F_PCT, italic=True,
          hist=m("{ebit_adj}/{is_rev}"), fc=m("{ebit_adj}/{is_rev}"))
    t.blank()
    t.sub("Reconciliation to reported EBIT (IFRS 16)")
    t.row("special", "Impairments and special items", "NOKm", hist=HL("special"),
          fc=lambda c, p, i: 0, input_fc=True, note="Forecast normally zero – enter known one-offs in the yellow cells")
    t.row("rou_dep", "Depreciation – right-of-use assets", "NOKm", hist=HL("rou_dep"),
          fc=m("{lease}*(1-<lint>)"))
    t.row("ebit", "EBIT (reported, IFRS 16)", "NOKm", bold=True, line=True,
          hist=m("{ebitda}+{da}+{rou_dep}+{special}"), fc=m("{ebitda}+{da}+{rou_dep}+{special}"))
    t.row("ebit_m", "EBIT margin (reported)", "%", fmt=F_PCT, italic=True, hist=m("{ebit}/{is_rev}"), fc=m("{ebit}/{is_rev}"))
    t.blank()
    t.row("netfin", "Net financial items excl. leases", "NOKm", hist=HL("netfin"),
          fc=m("-[nibd]*IF([nibd]>0,<rdebt>,<rcash>)"), note="Interest on opening net debt (avoids circularity)")
    t.row("lease_int", "Interest on lease liabilities", "NOKm", hist=HL("lease_int"), fc=m("{lease}*<lint>"))
    t.row("pbt", "Profit before tax", "NOKm", bold=True, line=True,
          hist=m("{ebit}+{netfin}+{lease_int}"), fc=m("{ebit}+{netfin}+{lease_int}"))
    t.row("tax", "Income tax", "NOKm", hist=HL("tax"), fc=m("-{pbt}*<tax>"))
    t.row("np", "Net profit", "NOKm", bold=True, line=True, hist=m("{pbt}+{tax}"), fc=m("{pbt}+{tax}"))
    t.row("min_pl", "Minority interests", "NOKm", hist=HL("min_pl"), fc=m("-{np}*<mins>"))
    t.row("np_sh", "Net profit to shareholders", "NOKm", key_row=True,
          hist=m("{np}+{min_pl}"), fc=m("{np}+{min_pl}"))
    t.row("np_m", "Net profit margin", "%", fmt=F_PCT, italic=True, hist=m("{np_sh}/{is_rev}"), fc=m("{np_sh}/{is_rev}"))
    t.blank()
    t.sub("Per share")
    sh_fc = lambda c, p, i: ((f"={I('shares_dil')}" if i == 0 else f"={p}{ROWS[('Model', 'shares')]}")
                             + f"+{c}{ROWS[('Model', 'buyback')]}/{c}{ROWS[('Model', 'bb_price')]}")
    t.row("shares", "Diluted shares", "m", fmt=F_NUM1, hist=HL("shares"), fc=sh_fc, note="Reduced by buybacks (section 4)")
    t.row("eps", "EPS", "NOK", fmt=F_NOK, bold=True, hist=m("{np_sh}/{shares}"), fc=m("{np_sh}/{shares}"))
    t.row("g_eps", "EPS growth", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=m('IF([eps]<=0,"n.m.",{eps}/[eps]-1)'), fc=m('IF([eps]<=0,"n.m.",{eps}/[eps]-1)'))
    t.row("dps", "DPS", "NOK", fmt=F_NOK, hist=HL("dps"), fc=m("MAX(0,{eps}*<payout>)"))
    t.row("payout", "Payout ratio", "%", fmt=F_PCT, italic=True,
          hist=m("IF({eps}<=0,0,{dps}/{eps})"), fc=m("IF({eps}<=0,0,{dps}/{eps})"))
    t.blank()
    # ------------------------------------------------------------------ 3
    t.section("3. Investments, working capital and free cash flow")
    t.row("nwc", "Net working capital", "NOKm", hist=HL("nwc"), fc=m("{is_rev}*<nwc>"))
    t.row("dnwc", "Change in NWC (increase = cash outflow)", "NOKm", hist_from=1,
          hist=m("{nwc}-[nwc]"), fc=m("{nwc}-[nwc]"))
    t.row("capex_m", "Maintenance capex", "NOKm", italic=True, indent=1, fc=m("-{is_rev}*<mcapex>"),
          note="≈ D&A – keeps the existing asset base intact (Drivers A)")
    t.row("capex_g", "Growth capex", "NOKm", italic=True, indent=1,
          fc=m("IF(@rev_mode=2,-{open}*<capex_loc>,-MAX(0,{is_rev}-[is_rev])/<s2c>)"),
          note="Location mode: openings x capex per new location. Segment mode: incremental revenue / sales-to-capital")
    t.row("capex", "Capital expenditure", "NOKm", hist=HL("capex"), fc=m("{capex_m}+{capex_g}"))
    t.row("ppe", "PP&E and intangibles excl. goodwill", "NOKm", hist=m("%ppe%+%intang%"),
          fc=m("[ppe]-{capex}+{da}"), note="Opening balance + capex – D&A")
    t.row("gw", "Goodwill", "NOKm", hist=HL("gw"), fc=m("[gw]"), note="Held flat – no acquisitions in the forecast")
    t.row("ic", "Invested capital (pre-IFRS 16)", "NOKm", key_row=True,
          hist=m("{nwc}+{ppe}+{gw}"), fc=m("{nwc}+{ppe}+{gw}"))
    t.blank()
    t.row("nopat", "NOPAT (pre-IFRS 16)", "NOKm", hist=m("{ebit_adj}*(1-@norm_tax)"), fc=m("{ebit_adj}*(1-<tax>)"))
    t.row("roic", "ROIC (average invested capital)", "%", fmt=F_PCT, bold=True, hist_from=1,
          hist=m("{nopat}/AVERAGE([ic],{ic})"), fc=m("{nopat}/AVERAGE([ic],{ic})"))
    t.row("net_inv", "Net reinvestment (capex – D&A + increase in NWC)", "NOKm", italic=True, hist_from=1,
          hist=m("-{capex}+{da}+{dnwc}"), fc=m("-{capex}+{da}+{dnwc}"),
          note="The return on this capital is compared with today's ROIC and the terminal RONIC on the DCF sheet")
    t.blank()
    t.row("fcff", "Free cash flow to firm (FCFF)", "NOKm", key_row=True, hist_from=1,
          hist=m("{nopat}-{da}+{capex}-{dnwc}"), fc=m("{nopat}-{da}+{capex}-{dnwc}"),
          note="NOPAT + D&A – capex – increase in NWC")
    t.row("fcff_m", "FCFF margin", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=m("{fcff}/{is_rev}"), fc=m("{fcff}/{is_rev}"))
    t.row("cconv", "Cash conversion (FCFF / EBITDAaL)", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=m("IF({ebitdaal}=0,0,{fcff}/{ebitdaal})"), fc=m("IF({ebitdaal}=0,0,{fcff}/{ebitdaal})"))
    t.blank()
    # ------------------------------------------------------------------ 4
    t.section("4. Financing and balance sheet roll-forward")
    t.row("div_paid", "Dividends paid", "NOKm", hist=HL("div_paid"), fc=m("-[dps]*[shares]"),
          note="Dividend for year t-1 is paid in year t")
    t.row("nibd_pre", "NIBD before buybacks", "NOKm", italic=True, indent=1,
          fc=m("[nibd]-({fcff}+({netfin}+{special})*(1-<tax>))-{div_paid}"),
          note="Opening NIBD less free cash flow to equity plus dividends")
    t.row("bb_price", "Assumed buyback price", "NOK", fmt=F_NOK, italic=True, indent=1,
          fc=lambda c, p, i: f"={I('price')}*(1+{DR('bb_g')})^{i + 1}", note="Current share price grown at the rate on Drivers E")
    t.row("buyback", "Share buybacks", "NOKm", indent=1,
          fc=m(f"IF({DR('bb_on')}=1,-MAX(0,{DR('lev_t')}*{{ebitdaal}}-{{nibd_pre}}),0)"),
          note="Cash above the target leverage (Drivers E) is returned when the switch is on")
    t.row("nibd", "Net interest-bearing debt excl. leases", "NOKm", key_row=True, hist=HL("nibd"),
          fc=m("{nibd_pre}-{buyback}"), note="NIBD before buybacks plus cash used for buybacks")
    t.row("leases", "Lease liabilities", "NOKm", hist=m("%lease_nc%+%lease_c%"),
          fc=m("IF([lease]=0,[leases],[leases]*{lease}/[lease])"), note="Grows with lease payments")
    t.row("nibd_incl", "Net debt incl. leases", "NOKm", bold=True, line=True, hist=m("{nibd}+{leases}"), fc=m("{nibd}+{leases}"))
    t.row("lev", "NIBD / EBITDAaL", "x", fmt=F_MULT, italic=True,
          hist=m("IF({ebitdaal}=0,0,{nibd}/{ebitdaal})"), fc=m("IF({ebitdaal}=0,0,{nibd}/{ebitdaal})"))
    t.blank()
    t.row("rou", "Right-of-use assets", "NOKm", hist=HL("rou"), fc=m("[rou]+{leases}-[leases]"),
          note="Moves with lease liabilities (new leases add to both sides)")
    t.row("other_nca", "Other non-current assets", "NOKm", hist=HL("other_nca"), fc=m("[other_nca]"), note="Held flat")
    t.row("deftax", "Deferred tax and other non-current liabilities", "NOKm", hist=HL("deftax"), fc=m("[deftax]"),
          note="Held flat")
    t.row("equity", "Equity to shareholders (book)", "NOKm", hist=HL("equity"), fc=m("[equity]+{np_sh}+{div_paid}+{buyback}"))
    t.row("min_bs", "Minority interests (book)", "NOKm", hist=HL("min_bs"), fc=m("[min_bs]-{min_pl}"))
    t.row("roe", "Return on equity", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=m("{np_sh}/AVERAGE([equity],{equity})"), fc=m("{np_sh}/AVERAGE([equity],{equity})"))


_model_spec()


def write_model(wb):
    ws = wb["Model"]
    sheet_title(ws, "Operating Model – Calculations",
                "Calculation engine: history is linked from Hist, forecasts are driven by the live drivers on Drivers. Results are presented on Output.")
    ts_header(ws, "Model")
    MODEL.write(ws)


# ============================================================ DCF
DCF = TS("DCF", start_row=21)
DCFB = Scalars("DCF", start_row=0, vcol="E", ucol=None, ncol="G")


def _dcf_spec():
    t = DCF
    d = lambda e: fx("DCF", e)
    B = lambda k: S("DCF", k, "DCF")
    t.section("Free cash flow to firm (FCFF, pre-IFRS 16 – lease payments treated as operating costs)")
    t.row("rev", "Total revenue", "NOKm", hist=d("~is_rev~"), fc=d("~is_rev~"),
          ty=lambda: f"=S{ROWS[('DCF', 'rev')]}*(1+{B('tv_g')})")
    t.row("grev", "Revenue growth", "%", fmt=F_PCT, italic=True, hist_from=1, hist=d("{rev}/[rev]-1"),
          fc=d("{rev}/[rev]-1"), ty=lambda: f"=T{ROWS[('DCF', 'rev')]}/S{ROWS[('DCF', 'rev')]}-1")
    t.row("ebitdaal", "EBITDAaL", "NOKm", hist=d("~ebitdaal~"), fc=d("~ebitdaal~"))
    t.row("ebitdaal_m", "EBITDAaL margin", "%", fmt=F_PCT, italic=True, hist=d("{ebitdaal}/{rev}"), fc=d("{ebitdaal}/{rev}"))
    t.row("da", "D&A – owned assets", "NOKm", hist=d("~da~"), fc=d("~da~"))
    t.row("ebit", "EBIT adj. (pre-IFRS 16)", "NOKm", bold=True, line=True, hist=d("~ebit_adj~"),
          fc=d("~ebit_adj~"), ty=lambda: f"=T{ROWS[('DCF', 'rev')]}*T{ROWS[('DCF', 'ebit_m')]}")
    t.row("ebit_m", "EBIT adj. margin", "%", fmt=F_PCT, italic=True, hist=d("{ebit}/{rev}"), fc=d("{ebit}/{rev}"),
          ty=lambda: f"={S('Drivers', 'ty_margin')}")
    t.row("tax", "Taxes on EBIT", "NOKm", hist=d("-{ebit}*@norm_tax"), fc=d("-{ebit}*<tax>"),
          ty=lambda: f"=-T{ROWS[('DCF', 'ebit')]}*{S('Drivers', 'lt_tax')}")
    t.row("nopat", "NOPAT", "NOKm", bold=True, line=True, hist=d("{ebit}+{tax}"), fc=d("{ebit}+{tax}"),
          ty=lambda: f"=T{ROWS[('DCF', 'ebit')]}+T{ROWS[('DCF', 'tax')]}")
    t.row("addda", "Add back: D&A", "NOKm", hist=d("-{da}"), fc=d("-{da}"))
    t.row("capex", "Less: capital expenditure", "NOKm", hist=d("~capex~"), fc=d("~capex~"))
    t.row("dnwc", "Less: increase in net working capital", "NOKm", hist_from=1, hist=d("-~dnwc~"), fc=d("-~dnwc~"))
    t.row("fcff", "Unlevered free cash flow (FCFF)", "NOKm", key_row=True, hist_from=1,
          hist=d("{nopat}+{addda}+{capex}+{dnwc}"), fc=d("{nopat}+{addda}+{capex}+{dnwc}"),
          ty=lambda: f"={B('tv_fcff')}", note="TY column: normalised terminal-year FCFF (see terminal value)")
    t.row("gfcff", "FCFF growth", "%", fmt=F_PCT, italic=True, hist_from=2,
          hist=d('IF([fcff]<=0,"n.m.",{fcff}/[fcff]-1)'), fc=d('IF([fcff]<=0,"n.m.",{fcff}/[fcff]-1)'))
    t.row("conv", "Cash conversion (FCFF / EBITDAaL)", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=d("IF({ebitdaal}=0,0,{fcff}/{ebitdaal})"), fc=d("IF({ebitdaal}=0,0,{fcff}/{ebitdaal})"))
    t.row("reinv", "Reinvestment rate (1 – FCFF / NOPAT)", "%", fmt=F_PCT, italic=True, hist_from=1,
          hist=d("IF({nopat}=0,0,1-{fcff}/{nopat})"), fc=d("IF({nopat}=0,0,1-{fcff}/{nopat})"),
          ty=lambda: f"=1-T{ROWS[('DCF', 'fcff')]}/T{ROWS[('DCF', 'nopat')]}")
    t.row("roic", "ROIC", "%", fmt=F_PCT, italic=True, hist_from=1, hist=d("~roic~"), fc=d("~roic~"))
    t.blank()
    t.section("Discounting (valuation date and stub period from Inputs)")
    t.row("n", "Forecast year number", "#", fmt=F_INT, fc=lambda c, p, i: f"=COLUMN()-COLUMN(${FC[0]}$1)+1")
    t.row("share", "Share of year included (stub)", "%", fmt=F_PCT, fc=d("IF({n}=1,@stub,1)"))
    t.row("fcff_incl", "FCFF included in valuation", "NOKm", fc=d("{fcff}*{share}"))
    t.row("period", "Discount period (years)", "yrs", fmt="0.00",
          fc=d("IF(@midyear=1,IF({n}=1,@stub/2,@stub+{n}-1.5),@stub+{n}-1)"))
    t.row("df", "Discount factor", "", fmt=F_FAC, fc=lambda c, p, i: f"=1/(1+{S('WACC', 'wacc')})^{c}{ROWS[('DCF', 'period')]}")
    t.row("pv", "Present value of FCFF", "NOKm", key_row=True, fc=d("{fcff_incl}*{df}"))
    t.layout()

    b = DCFB
    b.start = t.end + 2
    R = lambda k: f"$S${ROWS[('DCF', k)]}"
    T_ = lambda k: f"$T${ROWS[('DCF', k)]}"
    b.section("Terminal value")
    b.item("tv_g", "Terminal growth rate (g)", lambda: f"={S('Drivers', 'tg')}", fmt=F_PCT)
    b.item("tv_wacc", "WACC", lambda: f"={S('WACC', 'wacc')}", fmt=F_PCT)
    b.item("tv_nopat", "Terminal-year NOPAT (final-year revenue x (1+g) x terminal margin x (1-t))",
           lambda: f"={T_('nopat')}", fmt=F_NUM)
    b.item("tv_ronic", "Return on new invested capital (RONIC)", lambda: f"={S('Drivers', 'ronic')}", fmt=F_PCT)
    b.item("tv_reinv", "Net reinvestment (NOPAT x g / RONIC)", lambda: f"=-{B('tv_nopat')}*{B('tv_g')}/{B('tv_ronic')}", fmt=F_NUM)
    b.item("tv_fcff_vd", "Terminal FCFF – value driver method", lambda: f"={B('tv_nopat')}+{B('tv_reinv')}", fmt=F_NUM)
    b.item("tv_fcff_gr", "Terminal FCFF – final-year FCFF x (1+g)", lambda: f"={R('fcff')}*(1+{B('tv_g')})", fmt=F_NUM)
    b.item("tv_fcff", "Terminal FCFF used", lambda: f"=IF({S('Inputs', 'tv_method')}=1,{B('tv_fcff_vd')},{B('tv_fcff_gr')})",
           fmt=F_NUM, bold=True)
    b.item("tv_gordon", "Terminal value – Gordon growth (end of final year)",
           lambda: f"={B('tv_fcff')}/({B('tv_wacc')}-{B('tv_g')})", fmt=F_NUM, note="TV = FCFF(TY) / (WACC – g)")
    b.item("tv_mult", "Exit multiple (EV / EBITDAaL)", lambda: f"={S('Drivers', 'exit')}", fmt=F_MULT)
    b.item("tv_exit", "Terminal value – exit multiple (end of final year)", lambda: f"={R('ebitdaal')}*{B('tv_mult')}", fmt=F_NUM)
    b.item("tv_dfg", "Discount factor – Gordon TV (same timing as final-year FCFF)", lambda: f"={R('df')}", fmt=F_FAC,
           note="With mid-year convention the Gordon TV is discounted half a year less")
    b.item("tv_dfe", "Discount factor – exit-multiple TV (end of final year)",
           lambda: f"=1/(1+{B('tv_wacc')})^({S('Inputs', 'stub')}+{R('n')}-1)", fmt=F_FAC)
    b.item("tv_pvg", "PV of terminal value – Gordon growth", lambda: f"={B('tv_gordon')}*{B('tv_dfg')}", fmt=F_NUM)
    b.item("tv_pve", "PV of terminal value – exit multiple", lambda: f"={B('tv_exit')}*{B('tv_dfe')}", fmt=F_NUM)
    b.item("tv_pv", "PV of terminal value – weighted (weights on Inputs)",
           lambda: f"={S('Inputs', 'w_gordon')}*{B('tv_pvg')}+{S('Inputs', 'w_exit')}*{B('tv_pve')}", fmt=F_NUM, bold=True)
    b.blank()
    b.section("Enterprise value to equity value")
    b.item("sum_pv", "Sum of PV of FCFF (explicit forecast)", lambda: f"=SUM($L${ROWS[('DCF', 'pv')]}:$S${ROWS[('DCF', 'pv')]})", fmt=F_NUM)
    b.item("pv_tv", "PV of terminal value", lambda: f"={B('tv_pv')}", fmt=F_NUM)
    b.item("ev", "Enterprise value (DCF)", lambda: f"={B('sum_pv')}+{B('pv_tv')}", fmt=F_NUM, bold=True)
    b.item("tv_share", "Terminal value in % of EV", lambda: f"={B('pv_tv')}/{B('ev')}", fmt=F_PCT, italic=True)
    b.item("b_nibd", "Less: net interest-bearing debt excl. leases", lambda: f"=-{S('Inputs', 'nibd')}", fmt=F_NUM)
    b.item("b_min", "Less: minority interests", lambda: f"=-{S('Inputs', 'minority')}", fmt=F_NUM)
    b.item("b_assoc", "Add: associates and financial investments", lambda: f"={S('Inputs', 'associates')}", fmt=F_NUM)
    b.item("b_pens", "Less: pension deficit and other debt-like items", lambda: f"=-{S('Inputs', 'pension')}", fmt=F_NUM)
    b.item("b_other", "Add: other adjustments", lambda: f"={S('Inputs', 'other_adj')}", fmt=F_NUM)
    b.item("eq", "Equity value", lambda: f"={B('ev')}+SUM({B('b_nibd')}:{B('b_other')})", fmt=F_NUM, bold=True)
    b.item("memo_lease", "Memo: lease liabilities – not deducted, FCFF is already after lease payments",
           lambda: f"={S('Inputs', 'leases')}", fmt=F_NUM, italic=True, color=GREY)
    b.item("shares", "Diluted shares (m)", lambda: f"={S('Inputs', 'shares_dil')}", fmt=F_NUM1)
    b.item("dcf_ps", "DCF value per share (NOK)", lambda: f"={B('eq')}/{B('shares')}", fmt=F_NOK, bold=True)
    b.item("price", "Current share price (NOK)", lambda: f"={S('Inputs', 'price')}", fmt=F_NOK)
    b.item("dcf_up", "Upside / (downside) to DCF value", lambda: f"={B('dcf_ps')}/{B('price')}-1", fmt=F_PCT)
    b.blank()
    b.section("Sanity checks")
    b.item("s_ev1", "Implied EV / EBITDAaL – first forecast year", lambda: f"={B('ev')}/$L${ROWS[('DCF', 'ebitdaal')]}", fmt=F_MULT)
    b.item("s_ev2", "Implied EV / EBITDAaL – second forecast year", lambda: f"={B('ev')}/$M${ROWS[('DCF', 'ebitdaal')]}", fmt=F_MULT)
    b.item("s_exit", "Implied exit multiple from Gordon TV (EV / EBITDAaL)", lambda: f"={B('tv_gordon')}/{R('ebitdaal')}", fmt=F_MULT,
           note="Compare with peer EV/EBITDA and the company's own trading history")
    b.item("s_g", "Implied perpetual growth from exit-multiple TV",
           lambda: f"=({B('tv_exit')}*{B('tv_wacc')}-{R('fcff')})/({B('tv_exit')}+{R('fcff')})", fmt=F_PCT)
    b.item("s_roic", "Terminal-year ROIC (NOPAT / final-year invested capital)",
           lambda: f"={B('tv_nopat')}/Model!$S${ROWS[('Model', 'ic')]}", fmt=F_PCT)
    ni = lambda: ROWS[("Model", "net_inv")]
    nop = lambda: ROWS[("DCF", "nopat")]
    b.item("s_ronic_fc", "Return on new capital, explicit period (ΔNOPAT / cumulative net reinvestment)",
           lambda: f"=IF(SUM(Model!$L${ni()}:$R${ni()})<=0,9.99,($S${nop()}-$L${nop()})/SUM(Model!$L${ni()}:$R${ni()}))",
           fmt=F_PCT, note="Shown as 999% if revenue grows without net investment")
    b.item("s_ronic_ref", "Benchmark: the higher of ROIC in the last actual year and the terminal RONIC",
           lambda: f"=MAX({B('tv_ronic')},Model!$K${ROWS[('Model', 'roic')]})", fmt=F_PCT,
           note="Explicit-period returns are judged against what the company earns today, not only the long-run fade")
    b.item("s_ronic_ratio", "… relative to the benchmark (x) – tolerance on Inputs", lambda: f"={B('s_ronic_fc')}/{B('s_ronic_ref')}",
           fmt=F_MULT, note="Price increases on existing locations need no capital; far above the tolerance means growth is almost free")
    b.item("s_spread", "RONIC minus WACC", lambda: f"={B('tv_ronic')}-{B('tv_wacc')}", fmt=F_PCT,
           note="Positive = growth creates value in the terminal period")
    b.item("s_wg", "WACC minus terminal growth", lambda: f"={B('tv_wacc')}-{B('tv_g')}", fmt=F_PCT)
    b.blank()
    b.section("Target price and recommendation")
    b.item("tp_dcf", "DCF value per share (NOK)", lambda: f"={B('dcf_ps')}", fmt=F_NOK)
    b.item("tp_peer", "Peer multiples value per share (NOK)", lambda: f"={S('Comps', 'peer_val')}", fmt=F_NOK)
    b.item("tp_fair", "Blended fair value per share (NOK)",
           lambda: f"={B('tp_dcf')}*{S('Inputs', 'w_dcf')}+{B('tp_peer')}*{S('Inputs', 'w_peers')}", fmt=F_NOK, bold=True)
    b.item("tp_ke", "Cost of equity", lambda: f"={S('WACC', 'ke')}", fmt=F_PCT)
    b.item("tp_roll", "Roll-forward factor (12 months)", lambda: f"=IF({S('Inputs', 'roll12')}=1,1+{B('tp_ke')},1)", fmt=F_FAC)
    b.item("tp_div", "Less: expected dividend next 12 months (NOK)",
           lambda: f"=IF({S('Inputs', 'roll12')}=1,Model!$L${ROWS[('Model', 'dps')]},0)", fmt=F_NOK)
    b.item("tp_raw", "12-month target price, unrounded (NOK)", lambda: f"={B('tp_fair')}*{B('tp_roll')}-{B('tp_div')}", fmt=F_NOK)
    b.item("tp", "Target price (NOK)", lambda: f"=ROUND({B('tp_raw')}/{S('Inputs', 'round_to')},0)*{S('Inputs', 'round_to')}",
           fmt=F_NOK, bold=True)
    b.item("tp_up", "Upside / (downside) to target price", lambda: f"={B('tp')}/{B('price')}-1", fmt=F_PCT)
    b.item("tp_dy", "Expected dividend yield (next 12 months)", lambda: f"=Model!$L${ROWS[('Model', 'dps')]}/{B('price')}", fmt=F_PCT)
    b.item("tp_tr", "Expected total return", lambda: f"={B('tp_up')}+{B('tp_dy')}", fmt=F_PCT, bold=True)
    b.item("tp_excess", "Excess return: expected total return minus cost of equity", lambda: f"={B('tp_tr')}-{B('tp_ke')}", fmt=F_PCT,
           bold=True, note="A fairly priced share returns its cost of equity – only the excess justifies a rating")
    b.item("rating", "Recommendation", lambda: f'=IF({B("tp_excess")}>={S("Inputs", "buy_th")},"BUY",IF({B("tp_excess")}<={S("Inputs", "sell_th")},"SELL","HOLD"))',
           bold=True)
    b.layout()


def write_dcf(wb):
    ws = wb["DCF"]
    sheet_title(ws, "DCF Valuation",
                "FCFF discounted at WACC with stub period and mid-year convention; terminal value by value-driver Gordon growth.")
    ts_header(ws, "DCF", ty=True)
    # Campari-style assumptions box at the top of the sheet
    band(ws, 7, "B", "T", "Key assumptions (linked – change them on Inputs, Drivers and WACC)", size=9)
    ws.row_dimensions[7].height = 14.5
    ws.row_dimensions[8].height = 5
    box = [("WACC", f"={S('WACC', 'wacc')}", F_PCT2),
           ("Terminal growth rate", f"={S('Drivers', 'tg')}", F_PCT2),
           ("Return on new invested capital (RONIC)", f"={S('Drivers', 'ronic')}", F_PCT),
           ("Long-term tax rate", f"={S('Drivers', 'lt_tax')}", F_PCT),
           ("Valuation date", f"={S('Inputs', 'val_date')}", F_DATE),
           ("Share of first forecast year included (stub)", f"={S('Inputs', 'stub')}", F_PCT),
           ("Mid-year discounting", f'=IF({S("Inputs", "midyear")}=1,"On","Off")', F_GEN),
           ("Active scenario", f"={S('Inputs', 'scenario')}", F_GEN),
           ("Revenue build", f'=IF({S("Inputs", "rev_mode")}=2,"Locations: like-for-like + new","Segments: volume x price")', F_GEN),
           ("Growth capex (first forecast year)",
            f'=IF({S("Inputs", "rev_mode")}=2,"NOK "&FIXED(Drivers!{FC[0]}{ROWS[("Drivers", "capex_loc")]},1)&"m per new location",'
            f'"Sales-to-capital "&FIXED(Drivers!{FC[0]}{ROWS[("Drivers", "s2c")]},1)&"x")', F_GEN),
           ("Capital structure in the WACC", f'=IF({S("WACC", "cs_mode")}=1,"Target","Modelled")&" D/(D+E) "&FIXED({S("WACC", "dv")}*100,0)&"%"', F_GEN)]
    for j, (lab, f, fmt) in enumerate(box):
        put(ws, f"B{9 + j}", lab, kind="label")
        put(ws, f"E{9 + j}", f, fmt=fmt, align="right")
    side = Side(style="thin", color="000000")
    for r in range(9, 9 + len(box)):
        for col in "BCDE":
            c = ws[f"{col}{r}"]
            c.border = Border(left=side if col == "B" else None, right=side if col == "E" else None,
                              top=side if r == 9 else None, bottom=side if r == 8 + len(box) else None)
    DCF.write(ws)
    DCFB.write(ws, band_to="K")
    put(ws, f"F{ROWS[('DCF', 'fcff')] - 0}", None) if False else None
    for k in ["tp_dcf", "tp_peer"]:
        r = int(CELLS[("DCF", k)][1:])
        put(ws, f"F{r}", f"={S('Inputs', 'w_dcf' if k == 'tp_dcf' else 'w_peers')}", fmt=F_PCT, italic=True)
        put(ws, f"G{r}", "weight", kind="note", italic=True)
    for k in ["ev", "eq", "dcf_ps", "tv_pv", "tp", "rating", "tp_tr", "tp_excess"]:
        r = int(CELLS[("DCF", k)][1:])
        ws[f"E{r}"].fill = PatternFill("solid", fgColor=PALE)
        top_border(ws, r, "B", "E")
    r = int(CELLS[("DCF", "rating")][1:])
    ws[f"E{r}"].font = Font(name=FONT, size=11, bold=True, color=NAVY)


# ============================================================ WACC
WACC = Scalars("WACC", start_row=4, thin=False)
BETA_TOP = 5          # header row of the peer beta table (columns G..M)


def _wacc_spec():
    s = WACC
    W = lambda k: S("WACC", k, "WACC")
    w = ex.WACC_INPUTS
    s.section("Cost of equity (CAPM)")
    s.item("rf", "Risk-free rate (10-year Norwegian government bond)", w["rf"], fmt=F_PCT2,
           note="Example value – update to the current 10y yield (Norges Bank)")
    s.item("erp", "Equity risk premium", w["erp"], fmt=F_PCT2, note="NFF/PwC Norwegian risk premium survey: typically about 5%")
    s.item("bu_med", "Unlevered beta – peer median", f"=MEDIAN($M${BETA_TOP + 1}:$M${BETA_TOP + 8})", fmt=F_FAC,
           note="From the peer beta table to the right")
    s.item("bu_ovr", "Unlevered beta – override (optional)", None, fmt=F_FAC, note="Leave blank to use the peer median")
    s.item("bu", "Unlevered beta used", lambda: f'=IF({W("bu_ovr")}="",{W("bu_med")},{W("bu_ovr")})', fmt=F_FAC)
    s.item("de", "Target debt / equity (D/E)", lambda: f"={W('dv')}/{W('ev_w')}", fmt=F_PCT)
    s.item("tax", "Tax rate", lambda: f"={S('Drivers', 'lt_tax')}", fmt=F_PCT)
    s.item("bl", "Relevered beta = βu x (1 + (1 – t) x D/E)", lambda: f"={W('bu')}*(1+(1-{W('tax')})*{W('de')})", fmt=F_FAC,
           note="Hamada formula")
    s.item("size", "Size premium", w["size"], fmt=F_PCT2, note="Optional small/mid-cap premium")
    s.item("specific", "Company-specific risk premium", w["specific"], fmt=F_PCT2, note="Use sparingly – prefer scenarios")
    s.item("ke", "Cost of equity (Ke)", lambda: f"={W('rf')}+{W('bl')}*{W('erp')}+{W('size')}+{W('specific')}",
           fmt=F_PCT2, bold=True)
    s.blank()
    s.section("Cost of debt")
    s.item("spread", "Credit spread over the risk-free rate", w["spread"], fmt=F_PCT2,
           note="From loan/bond terms, or a synthetic rating based on interest cover")
    s.item("kd", "Pre-tax cost of debt (Kd)", lambda: f"={W('rf')}+{W('spread')}", fmt=F_PCT2)
    s.item("kd_at", "After-tax cost of debt", lambda: f"={W('kd')}*(1-{W('tax')})", fmt=F_PCT2, bold=True)
    s.blank()
    s.section("Capital structure and WACC")
    s.item("cs_mode", "Capital structure basis (1 = target, 2 = modelled from the forecast balance sheet)", 1, fmt=F_INT,
           note="2 = average forecast D/(D+E) with NIBD from the Model and equity at the current market cap")
    s.item("dv_target", "Target D / (D + E)", w["target_dv"], fmt=F_PCT, note="Long-term target, excl. leases (FCFF is pre-IFRS 16)")
    nr = lambda: ROWS[("Model", "nibd")]
    rng = lambda: f"Model!$L${nr()}:$S${nr()}"
    s.item("dv_model", "Modelled average D / (D + E) over the forecast",
           lambda: f"=SUMPRODUCT(({rng()}>0)*{rng()}/(({rng()}>0)*{rng()}+{S('Inputs', 'mcap')}))/{len(FC)}", fmt=F_PCT, italic=True,
           note="Net cash counts as zero debt. Buybacks (Drivers E) keep this close to the target – a gap above 10pp is flagged on Checks")
    s.item("dv", "D / (D + E) used", lambda: f"=IF({W('cs_mode')}=2,{W('dv_model')},{W('dv_target')})", fmt=F_PCT, bold=True)
    s.item("ev_w", "E / (D + E) used", lambda: f"=1-{W('dv')}", fmt=F_PCT)
    s.item("dv_mkt", "Current market D / (D + E), excl. leases – reference",
           lambda: f"={S('Inputs', 'nibd')}/({S('Inputs', 'nibd')}+{S('Inputs', 'mcap')})", fmt=F_PCT, italic=True)
    s.item("wacc_calc", "WACC (calculated)", lambda: f"={W('ev_w')}*{W('ke')}+{W('dv')}*{W('kd_at')}", fmt=F_PCT2, bold=True)
    s.item("wacc_ovr", "WACC override (optional)", None, fmt=F_PCT2, note="Leave blank to use the calculated WACC")
    s.item("wacc", "WACC used in the valuation", lambda: f'=IF({W("wacc_ovr")}="",{W("wacc_calc")},{W("wacc_ovr")})',
           fmt=F_PCT2, bold=True)
    CELLS[("WACC", "blume")] = f"M{BETA_TOP + 12}"
    s.layout()


def write_wacc(wb):
    ws = wb["WACC"]
    sheet_title(ws, "Cost of capital (WACC)", "CAPM with peer-based beta. Only blue cells are inputs.", last_col="M")
    WACC.write(ws, band_to="E")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 11
    ws.column_dimensions["D"].width = 3
    ws.column_dimensions["E"].width = 52
    ws.column_dimensions["F"].width = 3
    ws.column_dimensions["G"].width = 16
    for c in "HIJKLM":
        ws.column_dimensions[c].width = 11
    for k in ["ke", "wacc", "kd_at"]:
        c = ws[CELLS[("WACC", k)]]
        c.fill = PatternFill("solid", fgColor=PALE)
    from openpyxl.worksheet.datavalidation import DataValidation
    dv12 = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv12)
    dv12.add(CELLS[("WACC", "cs_mode")])
    top = BETA_TOP
    band(ws, top - 1, "G", "M", "Peer beta analysis")
    for col, txt in zip("GHIJKLM", ["Peer", "Country", "Raw beta", "Adj. beta", "D/E (mkt)", "Tax rate", "Unlev. beta"]):
        put(ws, f"{col}{top}", txt, kind="label", bold=True, align="right" if col not in "GH" else None)
    bottom_border(ws, top, "G", "M")
    blume = S("WACC", "blume", "WACC")
    for j, pr in enumerate(ex.PEERS):
        r = top + 1 + j
        crow = ROWS[("Comps", "peer1")] + j
        put(ws, f"G{r}", f"=Comps!B{crow}")
        put(ws, f"H{r}", f"=Comps!C{crow}")
        put(ws, f"I{r}", pr[10], fmt=F_FAC)
        put(ws, f"J{r}", f"=0.67*I{r}+0.33", fmt=F_FAC)
        put(ws, f"K{r}", pr[11], fmt=F_PCT)
        put(ws, f"L{r}", pr[12], fmt=F_PCT)
        put(ws, f"M{r}", f"=IF({blume}=1,J{r},I{r})/(1+(1-L{r})*K{r})", fmt=F_FAC)
    r = top + 9
    put(ws, f"G{r}", "Median", kind="label", bold=True)
    put(ws, f"G{r + 1}", "Mean", kind="label", bold=True)
    for col in "IJKLM":
        f = F_PCT if col in "KL" else F_FAC
        put(ws, f"{col}{r}", f"=MEDIAN({col}{top + 1}:{col}{top + 8})", fmt=f, bold=True)
        put(ws, f"{col}{r + 1}", f"=AVERAGE({col}{top + 1}:{col}{top + 8})", fmt=f)
    top_border(ws, r, "G", "M")
    put(ws, f"G{top + 12}", "Apply Blume adjustment (1 = yes)", kind="label")
    put(ws, f"M{top + 12}", ex.WACC_INPUTS["blume"], fmt=F_INT)
    put(ws, f"G{top + 13}", ex.MARKET.get("beta_note", "Adjusted beta = 0.67 x raw + 0.33 x 1.0 (Bloomberg convention). Source: example data – replace with 2y weekly betas."),
        kind="note", italic=True)
    put(ws, f"G{top + 14}", "Unlevered beta = beta / (1 + (1 – tax) x D/E)", kind="note", italic=True)

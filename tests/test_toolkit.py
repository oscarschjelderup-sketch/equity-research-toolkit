"""Checks that run anywhere (no Excel): the case data is consistent, both workbooks build from code with the layout the
decks expect, the committed workbooks pass their own integrity checks, and every deck builds on the neutral layout.

Finalising the workbook (data tables, full recalculation), verify_dcf.py and stress_test.py need Excel on Windows and
are run locally – see model_builder/README.md."""
import importlib
import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
MB, DB = ROOT / "model_builder", ROOT / "deck_builder"
SATS_XLSX = ROOT / "examples" / "SATS" / "SATS_ASA_DCF_Model.xlsx"
EXAMPLE_XLSX = ROOT / "Equity_Research_DCF_Toolkit.xlsx"


def run(args, cwd, **env):
    res = subprocess.run([sys.executable, *map(str, args)], cwd=cwd, env={**os.environ, **env},
                         capture_output=True, text=True)
    assert res.returncode == 0, res.stdout + res.stderr
    return res.stdout


@pytest.mark.parametrize("case", ["example_data", "sats_data"])
def test_history_balances(case):
    """simulate() asserts that every balance sheet balances and every cash flow ties to the change in cash."""
    sys.path.insert(0, str(MB))
    try:
        d = importlib.import_module(case).simulate()
    finally:
        sys.path.remove(str(MB))
    assert len(d["rev"]) == 7


@pytest.mark.parametrize("case, committed_map", [("example_data", "model_map.json"), ("sats_data", "model_map_sats.json")])
def test_workbook_builds_with_the_committed_layout(tmp_path, case, committed_map):
    out, cell_map = tmp_path / "model.xlsx", tmp_path / "map.json"
    run(["build_model.py", out], MB, EQR_CASE=case, EQR_MAP=str(cell_map))
    assert out.exists()
    fresh = json.loads(cell_map.read_text(encoding="utf8"))
    assert fresh == json.loads((MB / committed_map).read_text(encoding="utf8")), \
        "the cell layout changed – rebuild and commit the workbook and its model_map"


def test_deck_builder_uses_the_same_map():
    assert json.loads((DB / "model_map.json").read_text(encoding="utf8")) == \
        json.loads((MB / "model_map.json").read_text(encoding="utf8"))


@pytest.mark.parametrize("xlsx, committed_map, dcf, tp", [
    (EXAMPLE_XLSX, "model_map.json", 139.11, 146),
    (SATS_XLSX, "model_map_sats.json", 54.26, 57),
])
def test_committed_workbooks_pass_their_checks(xlsx, committed_map, dcf, tp):
    cells = json.loads((MB / committed_map).read_text(encoding="utf8"))["cells"]
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    master = str(wb["Checks"][cells["Checks|master"]].value)
    assert master.startswith(("ALL CHECKS OK", "OK")), master
    assert wb["DCF"][cells["DCF|dcf_ps"]].value == pytest.approx(dcf, abs=0.01)
    assert wb["DCF"][cells["DCF|tp"]].value == tp


def slides(path):
    return len(Presentation(path).slides)


def test_example_deck(tmp_path):
    out = tmp_path / "deck.pptx"
    run(["build_deck.py", "--template", "none", "--out", out], DB, EQR_CASE="example_data")
    assert slides(out) == 17


def test_sats_decks(tmp_path):
    full, short = tmp_path / "full.pptx", tmp_path / "submission.pptx"
    common = ["--template", "none", "--model", SATS_XLSX, "--map", MB / "model_map_sats.json"]
    run(["build_deck.py", *common, "--out", full], DB, EQR_CASE="sats_data")
    run(["build_deck.py", *common, "--main-only", "--out", short], DB, EQR_CASE="sats_data")
    assert slides(full) == 16
    assert slides(short) == 6          # Pareto rule: four content slides plus cover and team


def test_chart_library(tmp_path):
    out = tmp_path / "library.pptx"
    run(["build_chartlib.py", out], DB, EQR_TEMPLATE="none", EQR_CASE="example_data")
    assert slides(out) == 23

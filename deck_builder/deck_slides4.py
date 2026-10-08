"""Appendix slides for the market view: reverse DCF, thesis tracker, consensus and the growth engine."""
from pptx.enum.chart import XL_MARKER_STYLE
from deck_core import *
from deck_data import num, pct, mult, HC, FC
from deck_slides1 import std
from deck_case import ct

MLABEL = {"rev": "Revenue", "ebitda": "EBITDA", "ebit": "EBIT", "eps": "EPS", "dps": "DPS"}
STATUS_FILL = {"ON TRACK": "C6EFCE", "WATCH": "FFEB9C", "BROKEN": "FFC7CE", "OK": "C6EFCE", "INCONSISTENT": "FFC7CE"}


def _isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _p(v, dec=1):
    return pct(v, dec) if _isnum(v) else "n.a."


def _pp(a, b):
    return f"{(a - b) * 100:+.1f}pp".replace("-", "−") if _isnum(a) and _isnum(b) else "–"


# ------------------------------------------------------------------ reverse DCF
def reverse_dcf_slide(prs, d):
    R = lambda k: d.c("Reverse_DCF", k)
    price, dcf = d.c("Inputs", "price"), d.c("DCF", "dcf_ps")
    yl, ya = d.years([FC[-1]])[0], d.years([HC[-1]])[0]
    ya3 = d.years([HC[max(3, ct("hist_first", 0))]])[0]
    m_imp, m_our, m_ref = R("rd_m_last"), R("rd_m_last_ours"), R("rd_m_last_ref")
    rel = "below" if _isnum(m_imp) and m_imp < m_ref else "above"
    s = std(prs, "Appendix #4.2", "What the market is pricing in",
            f"At NOK {price:.0f} the market prices in an EBIT margin of {_p(m_imp)} in {yl} – {rel} the {pct(m_ref)} earned in "
            f"{ya}; our base case is {pct(m_our)}",
            "Sources: Case team reverse DCF (Excel model, Reverse_DCF sheet). Each driver is solved on its own with all other "
            "assumptions at the base case; growth is interpolated in Sensitivity table 5. Figures are illustrative.")
    x, w = 0.47, 6.55
    panel_header(s, x, 1.38, w, "Market-implied vs. our base case", None)
    rows = [dict(cells=["", "Market-implied", "Our base case", "Difference", "Reference"], fill=NAVY, color=WHITE, bold=True,
                 size=9, h=0.27, align={j: "c" for j in range(1, 5)})]
    spec = [(f"EBIT adj. margin, {yl}", "m_last", 1, f"{pct(m_ref)} ({ya})", True),
            ("EBIT adj. margin, forecast average", "m_avg", 1, f"{_p(R('rd_m_avg_ref'))} (3y avg.)", False),
            (f"Revenue CAGR, {ya}–{yl}", "cagr", 1, f"{_p(R('rd_cagr_ref'))} ({ya3[:4]}–{ya[2:]})", True),
            ("  … EBIT margin at that growth", "m_at_g", 1, "", False),
            ("WACC", "wacc", 2, "", False), ("Terminal growth", "g", 2, "", False)]
    for lab, k, dec, ref, bold in spec:
        imp, our = R(f"rd_{k}"), R(f"rd_{k}_ours")
        rows.append(dict(cells=[lab, _p(imp, dec), _p(our, dec), _pp(imp, our), ref], size=9, h=0.25, line_bottom="E1E5EA",
                         align={j: "c" for j in range(1, 5)}, bolds={0: bold, 1: True}, fills={1: "E4EEF2"},
                         italic=lab.startswith("  ")))
    table(s, x, 1.82, w, rows, [2.25, 1.05, 1.05, 0.9, 1.3])
    # bullets
    k = R("rd_k")
    cagr_i, cagr_o, m_g = R("rd_cagr"), R("rd_cagr_ours"), R("rd_m_at_g")
    w_i, w_o = R("rd_wacc"), R("rd_wacc_ours")
    panel(s, x, 3.75, w, 3.05)
    bullets = [
        ("How to read it", {"bold": True, "color": NAVY, "size": 10.5, "space_after": 4}),
        (f"**Margins:** the share price discounts an EBIT margin of {_p(m_imp)} in {yl} – {rel} today's {pct(m_ref)}. "
         f"Our {pct(m_our)} rests on operating leverage from a largely fixed cost base.", {"bullet": True, "space_after": 5}),
        (f"**Growth:** or revenue growth of {_p(cagr_i)} a year instead of our {_p(cagr_o)} – with our cost base that also "
         f"means a margin of {_p(m_g)}.", {"bullet": True, "space_after": 5}),
        (f"**Discount rate:** or a WACC of {_p(w_i)} (ours {_p(w_o)}).", {"bullet": True, "space_after": 5}),
        (f"**What it is worth:** each percentage point of EBIT margin in every year is worth **NOK {k:.1f} per share** "
         f"– the upside rests on margins, not on multiples.", {"bullet": True}),
    ]
    text(s, x + 0.15, 3.85, w - 0.3, 2.9, bullets, size=10)
    # chart: value vs. margin
    rx, rw = 7.27, 5.60
    panel_header(s, rx, 1.38, rw, f"DCF value per share vs. EBIT adj. margin in {yl} (NOK)", None)
    grid = [m_our + sh / 100 for sh in range(-5, 4)]
    if _isnum(m_imp):
        near = min(range(len(grid)), key=lambda i: abs(grid[i] - m_imp))
        if abs(grid[near] - m_imp) < 0.0015:
            grid[near] = m_imp                       # replace the neighbouring grid point by the exact implied margin
        else:
            grid = sorted(grid + [m_imp])
    val = lambda m: dcf + (m - m_our) * 100 * k
    cats = [f"{m * 100:.1f}%" for m in grid]
    pick = lambda m0: [val(m) if _isnum(m0) and abs(m - m0) < 1e-12 else None for m in grid]
    series = [("DCF value per share", [val(m) for m in grid]), ("Share price", [price] * len(grid)),
              ("Market-implied", pick(m_imp)), ("Our base case", pick(m_our))]
    gf = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, rx, 1.85, rw, 3.9, cats, series, [NAVY, RED, RED, NAVY], size=8, legend="t",
                   labels=False, num_fmt="0", val_axis=True, gridlines=True)
    ch = gf.chart
    ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = "0", False
    for j, ser in enumerate(ch.plots[0].series):
        if j < 2:
            ser.marker.style = XL_MARKER_STYLE.NONE
            ser.smooth = False
            if j == 1:
                ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                ser.format.line.width = Pt(1.5)
        else:
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 11
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(RED if j == 2 else NAVY)
            ser.marker.format.line.color.rgb = rgb(WHITE)
            ser.format.line.fill.background()
    text(s, rx, 5.82, rw, 0.9, [
        (f"The line crosses today's share price at an EBIT margin of {_p(m_imp)} (red point). Our base case of "
         f"{pct(m_our)} (navy point) gives NOK {dcf:.0f} per share. Each point on the line shifts the margin by the same amount in "
         f"every forecast year and the terminal year.", {"size": 8.5, "color": MUTED}),
    ], italic=True)
    notes(s, "TEMPLATE: Numbers come from the Reverse_DCF sheet. Lead with the margin: it is the one driver where we have a "
             "view. Be ready to explain why the market is too cautious – and what would prove it (Thesis sheet).")
    return s


# ------------------------------------------------------------------ thesis tracker
def thesis_slide(prs, d):
    T = lambda k: d.c("Thesis", k)
    n = sum(1 for j in range(1, 11) if f"Thesis|kc{j}_kpi" in d.reg["cells"])
    ya, y1 = d.years([HC[-1]])[0], d.years([FC[0]])[0]
    nb, nw = int(T("n_broken") or 0), int(T("n_watch") or 0)
    sub = (f"All {n} kill criteria are on track on the {ya} actuals – we would change our view if any of them is breached"
           if nb == 0 and nw == 0 else
           f"{nb} kill criteria breached and {nw} in the watch zone on the {ya} actuals – the case needs a second look")
    s = std(prs, "Appendix #4.4", "Thesis tracker – what would make us wrong", sub,
            "Sources: Case team (Excel model, Thesis sheet). Latest actuals link to the last reported year – update them with each "
            "quarterly report. Figures are illustrative.")
    x, w = 0.47, 12.40
    panel_header(s, x, 1.38, w, "Kill criteria", None)
    hdr = dict(cells=["KPI", "Kill if", "Threshold", f"Latest ({ya})", f"Base case ({y1})", "Watch zone", "Status", "Base case check"],
               fill=NAVY, color=WHITE, bold=True, size=9, h=0.28, align={j: "c" for j in range(1, 8)})
    rows = [hdr]
    for j in range(1, n + 1):
        addr = d.reg["cells"][f"Thesis|kc{j}_th"]
        is_mult = '"x"' in (d.wb["Thesis"][addr].number_format or "")
        f_ = (lambda v: f"{v:.2f}x") if is_mult else (lambda v: pct(v, 1))
        st, bc = T(f"kc{j}_status"), T(f"kc{j}_basechk")
        rows.append(dict(cells=[T(f"kc{j}_kpi"), T(f"kc{j}_dir"), f_(T(f"kc{j}_th")), f_(T(f"kc{j}_latest")), f_(T(f"kc{j}_base")),
                                f_(T(f"kc{j}_buf")), st, bc], size=9, h=0.27, line_bottom="E1E5EA",
                         align={j2: "c" for j2 in range(1, 8)}, bolds={0: True, 6: True, 7: True},
                         fills={6: STATUS_FILL.get(st, "FFFFFF"), 7: STATUS_FILL.get(bc, "FFFFFF")}))
    table(s, x, 1.82, w, rows, [3.3, 1.0, 1.1, 1.3, 1.5, 1.2, 1.5, 1.5])
    text(s, x, 1.82 + 0.28 + 0.27 * n + 0.08, w, 0.22,
         "Watch zone: the distance to the threshold that turns the status from ON TRACK to WATCH. Base case check: our own "
         "forecast for next year must not breach the criterion.", size=8, italic=True, color=MUTED)
    # signposts per scenario
    y0 = 4.05
    panel_header(s, x, y0, w, "Scenario signposts – what has to happen, and what we watch", None)
    t4 = {k: c for c, k, _, _ in d.reg["t4"]}
    r0 = d.reg["sens"]["t4_rows"][0]
    bw = (w - 0.3) / 3
    for j, (name, col) in enumerate([("bear", RED), ("base", NAVY), ("bull", GREEN)]):
        bx = x + j * (bw + 0.15)
        v = d.v("Sensitivity", f"{t4['dcf_ps']}{r0 + j}")
        hb = rect(s, bx, y0 + 0.45, bw, 0.32, fill=col)
        shape_text(hb, f"{name.capitalize()}  ·  NOK {v:.0f} per share  ·  {T(f'sc_{name}_prob') * 100:.0f}%", size=10, bold=True)
        panel(s, bx, y0 + 0.81, bw, 1.95)
        text(s, bx + 0.1, y0 + 0.88, bw - 0.2, 1.85, [
            (f"**Story:** {T(f'sc_{name}_story')}", {"bullet": True, "space_after": 5}),
            (f"**What has to happen:** {T(f'sc_{name}_trigger')}", {"bullet": True, "space_after": 5}),
            (f"**Watch:** {T(f'sc_{name}_signpost')}", {"bullet": True}),
        ], size=9)
    notes(s, "TEMPLATE: Kill criteria and signposts live on the Thesis sheet. Update the latest actuals after every report – "
             "a BROKEN status is a reason to revisit the recommendation, not to argue with the data.")
    return s


# ------------------------------------------------------------------ consensus
def consensus_slide(prs, d):
    C = lambda k: d.c("Consensus", k)
    yrs = d.years(FC[:3])
    e = [C(f"diff_ebit_{y}") for y in (1, 2, 3)]
    p = [C(f"diff_eps_{y}") for y in (1, 2, 3)]
    head = "EBIT"
    if not any(_isnum(v) for v in e):                 # no EBIT consensus (e.g. Yahoo Finance): lead with revenue
        e, head = [C(f"diff_rev_{y}") for y in (1, 2, 3)], "revenue"
    li = max([i for i in range(3) if _isnum(e[i]) and _isnum(p[i])] or [0])      # last year with a consensus estimate
    widening = head == "EBIT" and all(_isnum(v) for v in e) and e[0] < e[1] < e[2]
    _pc = lambda v: f"{v * 100:+.1f}%" if _isnum(v) else "n.a."
    s = std(prs, "Appendix #4.5", "Where we differ from consensus",
            f"Our {yrs[li]} {head} is {_pc(e[li])} vs. consensus and EPS {_pc(p[li])}"
            + (" – the gap widens every year as the operating leverage builds" if widening else ""),
            f"Sources: {C('cons_source')}, {C('cons_n')} analysts; case team estimates (Excel model, Consensus sheet). "
            "Reported IFRS 16 basis. Figures are illustrative.")
    x, w = 0.47, 12.40
    panel_header(s, x, 1.38, w, "Our estimates vs. consensus (NOKm, EPS and DPS in NOK)", None)
    hdr = dict(cells=[""] + [f"{y} {k}" for y in yrs for k in ("ours", "cons.", "diff.")], fill=NAVY, color=WHITE, bold=True,
               size=8.5, h=0.27, align={j: "c" for j in range(1, 10)})
    rows = [hdr]
    for m, lab in [("rev", "Revenue"), ("ebitda", "EBITDA (reported)"), ("ebit", "EBIT (reported)"), ("eps", "EPS"), ("dps", "DPS")]:
        dec = 2 if m in ("eps", "dps") else 0
        cells, colors = [lab], {}
        for y in (1, 2, 3):
            o, c_, df = C(f"our_{m}_{y}"), C(f"cons_{m}_{y}"), C(f"diff_{m}_{y}")
            cells += [f"{o:,.{dec}f}", f"{c_:,.{dec}f}" if _isnum(c_) else "n.a.", f"{df * 100:+.1f}%" if _isnum(df) else "n.a."]
            colors[3 * y] = (GREEN if df > 0.005 else RED if df < -0.005 else DARK) if _isnum(df) else MUTED
        rows.append(dict(cells=cells, size=9, h=0.26, line_bottom="E1E5EA", colors=colors, bolds={3: True, 6: True, 9: True, 0: m in ("ebit", "eps")},
                         fills={3: "EEF4F7", 6: "EEF4F7", 9: "EEF4F7"}, align={j: "c" for j in range(1, 10)}))
    table(s, x, 1.82, w, rows, [2.2] + [1.13] * 9)
    text(s, x, 3.48, w, 0.22, f"Consensus recommendations: {C('cons_buy')} Buy / {C('cons_hold')} Hold / {C('cons_sell')} Sell · "
                              f"consensus target price NOK {C('cons_tp'):.0f} ({C('cons_tp_up') * 100:+.0f}% upside). Green = we are above consensus.",
         size=8, italic=True, color=MUTED)
    # variant perception cards
    y0 = 3.85
    panel_header(s, x, y0, w, "Our variant perception", None)
    n = sum(1 for j in range(1, 6) if f"Consensus|vp{j}_topic" in d.reg["cells"])
    bw = (w - 0.15 * (n - 1)) / n
    for j in range(1, n + 1):
        bx = x + (j - 1) * (bw + 0.15)
        mkey, yi = C(f"vp{j}_metric"), int(C(f"vp{j}_year") or 3)
        df = C(f"diff_{mkey}_{yi}") if mkey in MLABEL else None
        hb = rect(s, bx, y0 + 0.45, bw, 0.34, fill=NAVY)
        shape_text(hb, f"{j}. {C(f'vp{j}_topic')}", size=10.5, bold=True, align="l", margin=0.1)
        panel(s, bx, y0 + 0.83, bw, 2.0)
        paras = [(f"**We:** {C(f'vp{j}_ours')}", {"space_after": 3}), (f"**Consensus:** {C(f'vp{j}_cons')}", {"space_after": 3}),
                 (f"**Why:** {C(f'vp{j}_why')}", {"space_after": 3}), (f"**Proof – and when:** {C(f'vp{j}_when')}", {})]
        text(s, bx + 0.1, y0 + 0.9, bw - 0.2, 1.55, paras, size=8.5)
        if _isnum(df):
            text(s, bx + 0.1, y0 + 2.45, bw - 0.2, 0.35, f"{MLABEL[mkey]} {yrs[yi - 1]}: {df * 100:+.1f}% vs. consensus", size=11,
                 bold=True, color=GREEN if df >= 0 else RED)
    notes(s, "TEMPLATE: Consensus inputs and the variant perception live on the Consensus sheet. A pitch without a clear "
             "difference to consensus has no edge – say what the market misses, why, and when it will see it.")
    return s


# ------------------------------------------------------------------ growth engine
def growth_engine_slide(prs, d):
    cols5 = FC[:5]
    ya, y1, y5 = d.years([HC[-1]])[0], d.years([FC[0]])[0], d.years([FC[4]])[0]
    rev = d.row("Model", "is_rev", [HC[-1], FC[4]])
    cagr = (rev[1] / rev[0]) ** 0.2 - 1
    opens = sum(d.row("Model", "open", cols5))
    newc, lflc, pxc = (sum(d.row("Model", k, cols5)) / 5 for k in ("g_newc", "g_lflc", "g_pxc"))
    s = std(prs, "Appendix #4.6", "Growth engine – like-for-like and new locations",
            f"{opens:.0f} new locations over {y1}–{y5} add ~{newc * 100:.1f}pp a year, like-for-like ~{lflc * 100:.1f}pp and "
            f"price/mix ~{pxc * 100:.1f}pp – {cagr * 100:.1f}% revenue CAGR",
            "Sources: Case team estimates (Excel model, Model and Drivers sheets – location mode). New locations ramp up over "
            "three years; fixed costs and rent grow with inflation and the number of locations. Figures are illustrative.")
    # left: growth decomposition
    gcols = HC[max(3, ct("hist_first", 0)):] + FC[:5]
    panel_header(s, 0.47, 1.38, 6.1, "Revenue growth by source", None)
    gf = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, 0.47, 1.80, 6.1, 2.35, d.years(gcols),
                   [("New locations", d.row("Model", "g_newc", gcols)), ("Like-for-like volume", d.row("Model", "g_lflc", gcols)),
                    ("Price/mix", d.row("Model", "g_pxc", gcols))], [NAVY, MIDBLUE, LBLUE], size=7.5, legend="t", labels=True,
                   num_fmt="0%", label_pos=XL_LABEL_POSITION.CENTER, gap=45, overlap=100, label_color=WHITE)
    series_labels_off(gf.chart, 1)
    # right: locations and revenue per location
    lcols = HC[max(1, ct("hist_first", 0)):] + FC[:5]
    panel_header(s, 6.77, 1.38, 6.1, "Locations (year end) and revenue per location (NOKm)", None)
    gf2 = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 6.77, 1.80, 6.1, 2.35, d.years(lcols),
                    [("Locations", d.row("Model", "loc", lcols)), ("Revenue per average location (rhs)", [v if _isnum(v) else None for v in d.row("Model", "rev_loc", lcols)])],
                    [LBLUE, NAVY], size=7.5, legend="t", labels=True, num_fmt="0", label_pos=XL_LABEL_POSITION.INSIDE_END,
                    gap=40, val_min=0, val_max=max(d.row("Model", "loc", lcols)) * 1.45, label_color=NAVY)
    _rl = [v for v in d.row("Model", "rev_loc", lcols) if _isnum(v)]      # the first history year has no average-location figure
    _rg = (max(_rl) - min(_rl)) or 1.0            # keep the line in the band above the columns (no label collisions)
    to_combo(gf2.chart, 1, color=NAVY, fmt="0", size=7.5, val_min=min(_rl) - 4 * _rg, val_max=max(_rl) + 0.56 * _rg)
    # bottom: location economics
    y0 = 4.40
    panel_header(s, 0.47, y0, 12.40, "Location economics – base case", None)
    tcols = [HC[-1]] + cols5
    rows = [dict(cells=[""] + d.years(tcols), fill=NAVY, color=WHITE, bold=True, size=8.5, h=0.25,
                 fills={1: "8C96A0"}, align={j: "c" for j in range(1, 7)})]
    capg = [-v if _isnum(v) else None for v in d.row("Model", "capex_g", tcols)]
    capl = d.row("Drivers", "capex_loc", tcols)
    spec = [("Locations, year end", d.row("Model", "loc", tcols), lambda v: f"{v:,.0f}", True),
            ("Net new locations", d.row("Model", "net_a", tcols), None, False),
            ("Mature-equivalent locations (all segments)", None, lambda v: f"{v:,.1f}", False),
            ("Revenue per average location (NOKm)", d.row("Model", "rev_loc", tcols), lambda v: f"{v:,.1f}", False),
            ("Capex per new location (NOKm)", capl, lambda v: f"{v:,.1f}", False),
            ("Growth capex (NOKm)", capg, lambda v: f"{v:,.0f}", True)]
    netall = [sum(x or 0 for x in vals) for vals in zip(*(d.row("Model", f"net_{s_}", tcols) for s_ in "abc"))]
    meqall = [sum(x or 0 for x in vals) for vals in zip(*(d.row("Model", f"meq_{s_}", tcols) for s_ in "abc"))]
    for lab, vals, f_, bold in spec:
        if lab.startswith("Net new"):
            vals, f_ = netall, (lambda v: f"{v:+.0f}")
        if lab.startswith("Mature"):
            vals = meqall
        cells = [lab] + [("–" if not _isnum(v) else f_(v)) for v in vals]
        rows.append(dict(cells=cells, size=9, h=0.24, line_bottom="E1E5EA", bolds={0: bold}, align={j: "c" for j in range(1, 7)},
                         fills={1: "F2F3F5"}))
    table(s, 0.47, y0 + 0.44, 8.4, rows, [3.0] + [0.9] * 6)
    r1, r2, r3 = (d.c("Drivers", k) for k in ("ramp1", "ramp2", "ramp3"))
    panel(s, 9.10, y0 + 0.44, 3.77, 1.73)
    text(s, 9.22, y0 + 0.50, 3.55, 1.65, [
        ("How a new location matures", {"bold": True, "color": NAVY, "size": 10, "space_after": 3}),
        (f"Volume vs. a mature location: **{r1 * 100:.0f}%** in the opening year, **{r2 * 100:.0f}%** in year two, "
         f"**{r3 * 100:.0f}%** in year three, 100% from year four.", {"bullet": True, "space_after": 3}),
        ("Rent and staff are paid from day one – openings dilute margins first and lift them as they mature.",
         {"bullet": True}),
    ], size=9)
    notes(s, "TEMPLATE: Location mode (Inputs: revenue build = 2). Openings, like-for-like growth and capex per location are on "
             "Drivers A; the ramp-up curve on Drivers F. Check the ramp-up against the company's own disclosures.")
    return s

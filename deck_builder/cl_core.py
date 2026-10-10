"""Chart library helpers (build on deck_core). All charts are native, editable PowerPoint charts."""
from lxml import etree
from pptx.chart.data import XyChartData, BubbleChartData
from pptx.enum.chart import XL_MARKER_STYLE
from deck_core import *
from deck_slides1 import std

SOFT = "A9C8D3"
LGREY = "C9D1D9"
PAL = [NAVY, MIDBLUE, LBLUE, SOFT, LGREY, "407061"]
PLOT = (0.03, 0.06, 0.94, 0.78)          # default inner plot area (x, y, w, h) as fractions of the chart frame


def cpanel(s, x, y, w, h, title, tip, num=None, fill=None):
    """Panel header + 'use for' tip. Returns the chart box (x, y, w, h)."""
    panel_header(s, x, y, w, title, num, size=11)
    if fill:
        panel(s, x, y + 0.42, w, h - 0.70, fill=fill)
    text(s, x, y + h - 0.24, w, 0.22, [("Use for: " + tip, {})], size=7.5, italic=True, color=MUTED)
    return x, y + 0.44, w, h - 0.72


def geo(box, plot=PLOT, n=None, vmax=None, vmin=0.0, horizontal=False):
    """Geometry helpers for annotations on top of a chart with a manual plot layout."""
    x, y, w, h = box
    px, py, pw, ph = plot
    if horizontal:
        X = lambda v: x + (px + pw * (v - vmin) / (vmax - vmin)) * w
        Y = lambda i: y + (py + ph * (i + 0.5) / n) * h
    else:
        X = lambda i: x + (px + pw * (i + 0.5) / n) * w
        Y = lambda v: y + (py + ph * (1 - (v - vmin) / (vmax - vmin))) * h
    return X, Y


def cagr_arrow(s, X, Y, i0, v0, i1, v1, label, lift=0.30):
    x1, y1, x2, y2 = X(i0), Y(v0) - lift, X(i1), Y(v1) - lift
    line(s, x1, y1, x2, y2, color=DARK, width=1.25, arrow_end=True)
    b = rect(s, (x1 + x2) / 2 - 0.5, (y1 + y2) / 2 - 0.36, 1.0, 0.24, fill=WHITE, line=DARK,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    shape_text(b, label, size=8, color=DARK, bold=True, margin=0)


def no_markers(chart):
    for srs in chart.plots[0].series:
        srs.marker.style = XL_MARKER_STYLE.NONE


def label_skip(chart, every):
    ax = chart._chartSpace.find(".//" + qn("c:catAx"))
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        for e in ax.findall(qn(tag)):
            ax.remove(e)
    anchor = ax.find(qn("c:noMultiLvlLbl"))
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        el = etree.Element(qn(tag), val=str(every))
        if anchor is not None:
            anchor.addprevious(el)
        else:
            ax.append(el)


def series_dash(chart, idx, dash="sysDash", width=1.5):
    ln = chart.plots[0].series[idx].format.line
    ln.width = Pt(width)
    lnel = ln._get_or_add_ln()
    for e in lnel.findall(qn("a:prstDash")):
        lnel.remove(e)
    d = etree.SubElement(lnel, qn("a:prstDash"), val=dash)
    sf = lnel.find(qn("a:solidFill"))
    if sf is not None:
        sf.addnext(d)


def waterfall(s, box, steps, fmt="{:,.0f}", size=8, vmax=None, colors=(NAVY, GREEN, RED), plot=PLOT, connectors=True):
    """steps: [(label, value, 'total'|'delta')]. Native stacked column with an invisible base series."""
    x, y, w, h = box
    base, tot, up, down, tops, run = [], [], [], [], [], 0.0
    for lab, v, kind in steps:
        if kind == "total":
            base.append(0); tot.append(v); up.append(0); down.append(0); run = v; tops.append(v)
        elif v >= 0:
            base.append(run); tot.append(0); up.append(v); down.append(0); run += v; tops.append(run)
        else:
            base.append(run + v); tot.append(0); up.append(0); down.append(-v); tops.append(run); run += v
    vmax = vmax or max(tops) * 1.18
    gf = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, x, y, w, h, [st[0] for st in steps],
                   [("Base", base), ("Total", tot), ("Increase", up), ("Decrease", down)],
                   [WHITE, colors[0], colors[1], colors[2]], size=size, labels=False, gap=35, overlap=100,
                   val_min=0, val_max=vmax)
    hide_series(gf.chart, 0)
    manual_plot_layout(gf.chart, *plot)
    X, Y = geo(box, plot, n=len(steps), vmax=vmax)
    run = 0.0
    for i, (lab, v, kind) in enumerate(steps):
        txt = fmt.format(v) if kind == "total" else ("+" if v >= 0 else "−") + fmt.format(abs(v))
        col = DARK if kind == "total" else (colors[1] if v >= 0 else colors[2])
        text(s, X(i) - 0.5, Y(tops[i]) - 0.22, 1.0, 0.2, txt, size=size, bold=True, color=col, align="c")
        if connectors and i < len(steps) - 1:
            level = v if kind == "total" else run + v
            bw = plot[2] * w / len(steps)
            line(s, X(i) + bw * 0.36, Y(level), X(i + 1) - bw * 0.36, Y(level), color=GRIDGREY, width=0.75)
        run = v if kind == "total" else run + v
    return gf


def scatter(s, box, points, xfmt="0%", yfmt='0.0"x"', xtitle="", ytitle="", xlim=None, ylim=None, size=8,
            highlight=None, trend=True, plot=(0.10, 0.05, 0.86, 0.74), positions=None):
    """points: [(name, x, y)]. highlight = name drawn in a different colour. positions: {name: XL_LABEL_POSITION} for
    labels that would otherwise overlap (default RIGHT)."""
    x, y, w, h = box
    cd = XyChartData()
    a = cd.add_series("Peers")
    b = cd.add_series("Company")
    for name, px_, py_ in points:
        (b if name == highlight else a).add_data_point(px_, py_)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name, ch.font.size = FONT, Pt(size)
    names = {0: [p for p in points if p[0] != highlight], 1: [p for p in points if p[0] == highlight]}
    for k, srs in enumerate(ch.plots[0].series):
        srs.marker.style = XL_MARKER_STYLE.CIRCLE
        srs.marker.size = 9 if k == 0 else 11
        srs.marker.format.fill.solid()
        srs.marker.format.fill.fore_color.rgb = rgb(LBLUE if k == 0 else NAVY)
        srs.marker.format.line.fill.background()
        srs.format.line.fill.background()
        for i, p in enumerate(names[k]):
            dl = srs.points[i].data_label
            dl.position = (positions or {}).get(p[0], XL_LABEL_POSITION.RIGHT)
            tf = dl.text_frame
            tf.text = p[0]
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.name = Pt(size - 0.5), k == 1, FONT
            r.font.color.rgb = rgb(NAVY if k == 1 else DARK)
    xs, ys = [p[1] for p in points], [p[2] for p in points]
    xlim = xlim or (min(xs) * 0.8, max(xs) * 1.15)
    ylim = ylim or (min(ys) * 0.8, max(ys) * 1.15)
    for ax, lim, f, ttl in ((ch.category_axis, xlim, xfmt, xtitle), (ch.value_axis, ylim, yfmt, ytitle)):
        ax.minimum_scale, ax.maximum_scale = lim
        ax.has_major_gridlines = ax is ch.value_axis
        if ax.has_major_gridlines:
            ax.major_gridlines.format.line.color.rgb = rgb("E3E6EA")
        ax.tick_labels.number_format, ax.tick_labels.number_format_is_linked = f, False
        ax.tick_labels.font.size = Pt(size)
        ax.format.line.color.rgb = rgb(GRIDGREY)
        ax.major_tick_mark = XL_TICK_MARK.NONE
        if ttl:
            ax.has_title = True
            ax.axis_title.text_frame.text = ttl
            rr = ax.axis_title.text_frame.paragraphs[0].runs[0]
            rr.font.size, rr.font.bold, rr.font.name = Pt(size), False, FONT
    manual_plot_layout(ch, *plot)
    if trend and len(points) > 2:
        n = len(points)
        mx, my = sum(xs) / n, sum(ys) / n
        beta = sum((a_ - mx) * (b_ - my) for a_, b_ in zip(xs, ys)) / sum((a_ - mx) ** 2 for a_ in xs)
        alpha = my - beta * mx
        PX = lambda v: x + (plot[0] + plot[2] * (v - xlim[0]) / (xlim[1] - xlim[0])) * w
        PY = lambda v: y + (plot[1] + plot[3] * (1 - (v - ylim[0]) / (ylim[1] - ylim[0]))) * h
        x0, x1 = xlim[0] + 0.04 * (xlim[1] - xlim[0]), xlim[1] - 0.04 * (xlim[1] - xlim[0])
        line(s, PX(x0), PY(alpha + beta * x0), PX(x1), PY(alpha + beta * x1), color=MUTED, width=1,
             dash=MSO_LINE_DASH_STYLE.DASH)
        ss_res = sum((b_ - (alpha + beta * a_)) ** 2 for a_, b_ in zip(xs, ys))
        ss_tot = sum((b_ - my) ** 2 for b_ in ys)
        text(s, x + w - 1.3, y + plot[1] * h, 1.25, 0.2, f"R² = {1 - ss_res / ss_tot:.2f}", size=size, italic=True,
             color=MUTED, align="r")
    return gf


def bubble(s, box, points, xfmt="0%", yfmt="0%", xtitle="", ytitle="", xlim=None, ylim=None, size=8, highlight=None,
           plot=(0.10, 0.05, 0.86, 0.74), scale=60):
    """points: [(name, x, y, size)]."""
    x, y, w, h = box
    cd = BubbleChartData()
    a = cd.add_series("Peers")
    b = cd.add_series("Company")
    for name, px_, py_, sz in points:
        (b if name == highlight else a).add_data_point(px_, py_, sz)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.name, ch.font.size = FONT, Pt(size)
    bc = ch._chartSpace.find(".//" + qn("c:bubbleChart"))
    bs = bc.find(qn("c:bubbleScale"))
    if bs is None:
        bs = etree.Element(qn("c:bubbleScale"))
        bc.find(qn("c:axId")).addprevious(bs)
    bs.set("val", str(scale))
    names = {0: [p for p in points if p[0] != highlight], 1: [p for p in points if p[0] == highlight]}
    for k, srs in enumerate(ch.plots[0].series):
        srs.format.fill.solid()
        srs.format.fill.fore_color.rgb = rgb(LBLUE if k == 0 else NAVY)
        srs.format.line.color.rgb = rgb(WHITE)
        for i, p in enumerate(names[k]):
            dl = srs.points[i].data_label
            dl.position = XL_LABEL_POSITION.RIGHT
            tf = dl.text_frame
            tf.text = p[0]
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.name = Pt(size - 0.5), k == 1, FONT
            r.font.color.rgb = rgb(NAVY if k == 1 else DARK)
    xs, ys = [p[1] for p in points], [p[2] for p in points]
    xlim = xlim or (min(xs) * 0.7, max(xs) * 1.2)
    ylim = ylim or (min(ys) * 0.7, max(ys) * 1.2)
    for ax, lim, f, ttl in ((ch.category_axis, xlim, xfmt, xtitle), (ch.value_axis, ylim, yfmt, ytitle)):
        ax.minimum_scale, ax.maximum_scale = lim
        ax.has_major_gridlines = ax is ch.value_axis
        if ax.has_major_gridlines:
            ax.major_gridlines.format.line.color.rgb = rgb("E3E6EA")
        ax.tick_labels.number_format, ax.tick_labels.number_format_is_linked = f, False
        ax.tick_labels.font.size = Pt(size)
        ax.format.line.color.rgb = rgb(GRIDGREY)
        ax.major_tick_mark = XL_TICK_MARK.NONE
        if ttl:
            ax.has_title = True
            ax.axis_title.text_frame.text = ttl
            rr = ax.axis_title.text_frame.paragraphs[0].runs[0]
            rr.font.size, rr.font.bold, rr.font.name = Pt(size), False, FONT
    manual_plot_layout(ch, *plot)
    return gf


def pie(s, box, cats, vals, colors, size=8, doughnut=False, hole=55, centre=None, labels="pct", legend=None,
        label_colors=None):
    x, y, w, h = box
    tot = float(sum(vals))
    vals = [v / tot for v in vals]          # fractions, so value labels formatted 0% are correct
    gf = add_chart(s, XL_CHART_TYPE.DOUGHNUT if doughnut else XL_CHART_TYPE.PIE, x, y, w, h, cats, [("Share", vals)],
                   [NAVY], size=size, labels=True, num_fmt="0%", legend=legend)
    pl = gf.chart.plots[0]
    pl.vary_by_categories = True
    dl = pl.data_labels
    dl.show_value = False
    dl.show_percentage = True
    dl.number_format, dl.number_format_is_linked = "0%", False
    if labels is None:
        pl.has_data_labels = False
    if labels == "name":
        dl.show_category_name = True
        if not doughnut:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for i, col in enumerate(colors[:len(vals)]):
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
        pt.format.line.color.rgb = rgb(WHITE)
        if labels == "pct":
            lc = (label_colors or [WHITE] * len(vals))[i]
            pt.data_label.font.size = Pt(size)
            pt.data_label.font.bold = True
            pt.data_label.font.color.rgb = rgb(lc)
    if doughnut:
        doughnut_hole(gf.chart, hole)
        if centre:
            text(s, x + w / 2 - 0.6, y + h / 2 - 0.25, 1.2, 0.5, centre, size=size + 0.5, bold=True, color=NAVY,
                 align="c", anchor="m")
    return gf


def swatches(s, x, y, items, size=8, dx=1.3, vertical=False, dy=0.26):
    for i, (lab, col) in enumerate(items):
        xx, yy = (x, y + i * dy) if vertical else (x + i * dx, y)
        rect(s, xx, yy + 0.04, 0.13, 0.13, fill=col)
        text(s, xx + 0.18, yy, dx - 0.2 if not vertical else 2.2, 0.22, lab, size=size, color=DARK)


def heat(v, lo, mid, hi):
    """Red – white – green colour scale, returns hex."""
    def mix(c1, c2, t):
        return "".join(f"{int(round(int(c1[i:i + 2], 16) * (1 - t) + int(c2[i:i + 2], 16) * t)):02X}" for i in (0, 2, 4))
    if v <= mid:
        return mix("F4B9AE", "FFFFFF", 0 if mid == lo else (v - lo) / (mid - lo))
    return mix("FFFFFF", "B7DDB0", 1 if hi == mid else (v - mid) / (hi - mid))


def chip(s, x, y, w, label, fill, color=WHITE, size=8, h=0.22):
    b = rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    shape_text(b, label, size=size, color=color, bold=True, margin=0)
    return b


def elbow(s, x1, y1, x2, y2, color=GRIDGREY, width=1.0):
    """Right-angle connector drawn as three straight segments (horizontal – vertical – horizontal)."""
    xm = (x1 + x2) / 2
    line(s, x1, y1, xm, y1, color=color, width=width)
    line(s, xm, y1, xm, y2, color=color, width=width)
    line(s, xm, y2, x2, y2, color=color, width=width)

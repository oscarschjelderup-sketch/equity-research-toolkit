"""Chart library – part 3: market-view charts (reverse DCF, growth engine, consensus gap, locations)."""
from pptx.enum.chart import XL_MARKER_STYLE
from cl_core import *
from cl_a1 import FOOT
from deck_data import HC, FC

XS2, W2 = [0.47, 6.87], 6.0
YS2, H2 = [1.40, 4.18], 2.68


def _g2(i):
    return XS2[i % 2], YS2[i // 2]


def market_view_charts(prs, d):
    s = std(prs, "Chart library #9", "Market-view charts",
            "What the share price implies, where the growth comes from and where we differ from consensus – all read from the model",
            FOOT + " Reverse DCF, consensus and the location engine come from the Reverse_DCF, Consensus and Model sheets.")
    R = lambda k: d.c("Reverse_DCF", k)
    price, dcf = d.c("Inputs", "price"), d.c("DCF", "dcf_ps")
    yl = d.years([FC[-1]])[0]
    # 1 reverse DCF: value vs. margin, with the market-implied and base-case points
    x, y = _g2(0)
    box = cpanel(s, x, y, W2, H2, f"Reverse DCF – value vs. EBIT margin in {yl} (NOK)",
                 "the price-implied margin vs. yours – where the debate is", 1)
    m_imp, m_our, k = R("rd_m_last"), R("rd_m_last_ours"), R("rd_k")
    grid = [m_our + sh / 100 for sh in range(-5, 4)]
    if isinstance(m_imp, (int, float)):
        near = min(range(len(grid)), key=lambda i: abs(grid[i] - m_imp))
        if abs(grid[near] - m_imp) < 0.0015:
            grid[near] = m_imp
        else:
            grid = sorted(grid + [m_imp])
    val = lambda m: dcf + (m - m_our) * 100 * k
    pick = lambda m0: [val(m) if isinstance(m0, (int, float)) and abs(m - m0) < 1e-12 else None for m in grid]
    gf = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, *box, [f"{m * 100:.1f}%" for m in grid],
                   [("DCF value per share", [val(m) for m in grid]), ("Share price", [price] * len(grid)),
                    ("Market-implied", pick(m_imp)), ("Our base case", pick(m_our))],
                   [NAVY, RED, RED, NAVY], size=7, legend="t", labels=False, num_fmt="0", val_axis=True, gridlines=True)
    for j, ser in enumerate(gf.chart.plots[0].series):
        if j < 2:
            ser.marker.style = XL_MARKER_STYLE.NONE
            if j == 1:
                ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                ser.format.line.width = Pt(1.25)
        else:
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 10
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(RED if j == 2 else NAVY)
            ser.marker.format.line.color.rgb = rgb(WHITE)
            ser.format.line.fill.background()
    gf.chart.value_axis.tick_labels.number_format, gf.chart.value_axis.tick_labels.number_format_is_linked = "0", False
    # 2 growth by source
    x, y = _g2(1)
    box = cpanel(s, x, y, W2, H2, "Revenue growth by source", "capacity vs. like-for-like vs. price – what drives the top line", 2)
    gcols = HC[3:] + FC[:5]
    gf = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, *box, d.years(gcols),
                   [("New locations", d.row("Model", "g_newc", gcols)), ("Like-for-like", d.row("Model", "g_lflc", gcols)),
                    ("Price/mix", d.row("Model", "g_pxc", gcols))], [NAVY, MIDBLUE, LBLUE], size=7, legend="t", labels=True,
                   num_fmt="0%", label_pos=XL_LABEL_POSITION.CENTER, gap=45, overlap=100, label_color=WHITE)
    series_labels_off(gf.chart, 1)
    # 3 consensus gap
    x, y = _g2(2)
    box = cpanel(s, x, y, W2, H2, "Our estimates vs. consensus", "show where – and how much – you differ from the market", 3)
    yrs = d.years(FC[:3])
    num = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None    # "n.a." -> gap in the chart
    series = [(MLAB, [num(d.c("Consensus", f"diff_{m}_{j}")) for j in (1, 2, 3)])
              for m, MLAB in [("rev", "Revenue"), ("ebit", "EBIT"), ("eps", "EPS")]]
    add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, *box, yrs, series, [LBLUE, MIDBLUE, NAVY], size=7, legend="t", labels=True,
              num_fmt="+0%;-0%;0%", label_pos=XL_LABEL_POSITION.OUTSIDE_END, gap=70, overlap=-10)
    # 4 locations and revenue per location
    x, y = _g2(3)
    box = cpanel(s, x, y, W2, H2, "Locations and revenue per location (NOKm)", "the roll-out and what each location earns", 4)
    lcols = HC[1:] + FC[:5]
    locs, rpl = d.row("Model", "loc", lcols), d.row("Model", "rev_loc", lcols)
    gf = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, *box, d.years(lcols),
                   [("Locations", locs), ("Revenue per location (rhs)", rpl)], [LBLUE, NAVY], size=7, legend="t", labels=True,
                   num_fmt="0", label_pos=XL_LABEL_POSITION.INSIDE_END, gap=40, val_min=0, val_max=max(locs) * 1.45,
                   label_color=NAVY)
    to_combo(gf.chart, 1, color=NAVY, fmt="0", size=7, val_min=0, val_max=max(rpl) * 1.2)
    notes(s, "Chart 1: the red point is the EBIT margin at which the DCF equals today's share price (Reverse_DCF sheet); every "
             "point shifts the margin by the same amount in all years. Charts 2-4 read the Model and Consensus sheets.")
    return s

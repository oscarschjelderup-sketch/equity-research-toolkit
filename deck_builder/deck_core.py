"""Helpers for building the pitch deck with python-pptx (native charts and tables).

With a template (any deck built on the Pareto case template) the slides use its layouts and placeholders. Without one
the same layout is drawn on blank slides: same positions, fonts and colours, no logo (the "neutral" layout)."""
import copy
import os
import re
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.oxml.ns import qn

FONT = "Arial"
NAVY, LBLUE, MIDBLUE, PALEBLUE = "003255", "81B0C0", "4F7FA0", "D6E6EC"
PANEL, GRIDGREY, DARK, MUTED = "ECEEF1", "BFC5CC", "1F1F1F", "6B7480"
GREEN, RED, AMBER, BLACK, WHITE = "2E7D32", "C00000", "C9A227", "000000", "FFFFFF"
SLIDE_W, SLIDE_H = 13.333, 7.5
TITLE_FONT = "Cambria"
BRAND = "Equity Research Toolkit"   # footer text on the neutral layout
LEFT, RIGHT = 0.47, 12.87
CONTENT_W = RIGHT - LEFT


def rgb(h):
    return RGBColor.from_string(h)


# ------------------------------------------------------------------ presentation / slides
def open_template(path):
    """Open a template and drop its slides; with no template (None, "none" or a missing file) start a neutral deck."""
    if not path or os.path.basename(str(path)).lower() == "none" or not os.path.exists(path):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        prs._eqr_neutral = True
        return prs
    prs = Presentation(path)
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.get(qn("r:id")))
        lst.remove(sld)
    return prs


def layout(prs, name):
    for lay in prs.slide_masters[0].slide_layouts:
        if lay.name.strip() == name:
            return lay
    raise KeyError(name)


def is_neutral(prs):
    return getattr(prs, "_eqr_neutral", False)


def _neutral_slide(prs, lay_name, title, subtitle, chapter, footnote, ph_text):
    """Draw the template's three layouts on a blank slide (positions taken from the case template)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    n = len(prs.slides._sldIdLst)
    if lay_name.startswith(("Title slide", "Chapter slide")):
        v = ph_text or {}
        rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
        line(s, 1.336, 2.726, 2.478, 2.726, color=LBLUE, width=1.5)
        if lay_name.startswith("Title slide"):
            text(s, 0.913, 0.95, 6.0, 0.5, BRAND, size=20, color=WHITE, font=TITLE_FONT)
            for k, (y, h, sz) in {0: (3.081, 1.181, 36), 1: (4.595, 0.905, 20), 12: (6.191, 0.337, 12)}.items():
                if v.get(k):
                    text(s, 1.336, y, 9.6, h, v[k], size=sz, color=WHITE, font=TITLE_FONT if k == 0 else FONT,
                         anchor="b" if k == 0 else "t")
        else:
            if v.get(0):
                text(s, 1.339, 2.95, 9.0, 0.75, v[0], size=36, color=WHITE, font=TITLE_FONT, anchor="b")
            if v.get(1):
                text(s, 1.344, 4.516, 9.0, 1.06, v[1], size=18, color=LBLUE)
        return s
    if chapter:
        text(s, 0.474, 0.197, 5.96, 0.223, chapter, size=8, bold=True, color=NAVY)
        line(s, 0.474, 0.404, 1.124, 0.404, color=NAVY, width=1.0)
    if title:
        text(s, 0.469, 0.50, 12.395, 0.48, title, size=26, color=NAVY, font=TITLE_FONT, anchor="b")
    if subtitle:
        text(s, 0.469, 0.993, 12.4, 0.3, subtitle, size=14, color=NAVY)
    line(s, 0.459, 6.979, 11.29, 6.979, color=GRIDGREY, width=0.75)
    line(s, 11.748, 6.979, 12.872, 6.979, color=GRIDGREY, width=0.75)
    if footnote:
        text(s, 0.474, 7.036, 10.758, 0.382, footnote, size=7, italic=True, color=DARK)
    text(s, 11.70, 7.04, 0.85, 0.36, ["Equity Research", "Toolkit"], size=7, color=NAVY, font=TITLE_FONT, bold=True)
    text(s, 12.529, 7.106, 0.34, 0.236, str(n), size=9, color=NAVY, align="r")
    return s


def new_slide(prs, lay_name, title=None, subtitle=None, chapter=None, footnote=None, ph_text=None):
    if is_neutral(prs):
        return _neutral_slide(prs, lay_name, title, subtitle, chapter, footnote, ph_text)
    s = prs.slides.add_slide(layout(prs, lay_name))
    vals = ph_text if ph_text is not None else {0: title, 11: subtitle, 13: chapter, 12: footnote}
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        v = vals.get(idx)
        if v is None:
            ph._element.getparent().remove(ph._element)
            continue
        set_text(ph.text_frame, v)
    return s


def set_text(tf, text):
    tf.text = ""
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ------------------------------------------------------------------ text
_MARK = re.compile(r"(\*\*.+?\*\*)")


def _runs(p, text, size, color, bold, italic, font):
    end = p._p.get_or_add_endParaRPr()
    end.set("sz", str(int(round(size * 100))))
    parts = _MARK.split(text) if isinstance(text, str) else [text]
    for part in parts:
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        r = p.add_run()
        r.text = part
        f = r.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), b, italic
        f.color.rgb = rgb(color)


def _bullet(p, char="▪", color=NAVY, indent_in=0.16, level_in=0.0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in + level_in))))
    pPr.set("indent", str(int(-Inches(indent_in))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buSzPct"), val="100000")
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char=char)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def text(slide, x, y, w, h, paras, size=10, color=DARK, bold=False, italic=False, align="l", anchor="t",
         font=FONT, margin=0.02, space_after=0, line_spacing=None, bullets=False, wrap=True, name=None):
    """paras: str or list of str / (str, opts). opts: size,color,bold,italic,align,bullet,space_after,level."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        t, o = (item, {}) if not isinstance(item, tuple) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[o.get("align", align)]
        sa = o.get("space_after", space_after)
        if sa:
            p.space_after = Pt(sa)
        if line_spacing:
            p.line_spacing = line_spacing
        if o.get("bullet", bullets):
            _bullet(p, level_in=0.18 * o.get("level", 0))
        else:
            _no_bullet(p)
        _runs(p, t, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", italic), font)
    return tb


# ------------------------------------------------------------------ shapes
def rect(slide, x, y, w, h, fill=PANEL, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE, name=None, dash=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def shape_text(s, t, size=10, color=WHITE, bold=True, align="c", anchor="m", margin=0.03, italic=False):
    tf = s.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = t if isinstance(t, list) else [t]
    for i, item in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        _runs(p, item, size, color, bold, italic, FONT)
    return s


def line(slide, x1, y1, x2, y2, color=NAVY, width=1.0, dash=None, arrow_end=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        c.name = name
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if arrow_end:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def badge(slide, x, y, num, size=0.30, fill=NAVY, font_size=12):
    s = rect(slide, x, y, size, size, fill=fill)
    shape_text(s, str(num), size=font_size, color=WHITE, bold=False, margin=0)
    return s


def circle(slide, x, y, d, fill=NAVY, txt=None, font_size=11, color=WHITE, line_col=None, bold=True):
    s = rect(slide, x, y, d, d, fill=fill, line=line_col, shape=MSO_SHAPE.OVAL)
    if txt is not None:
        shape_text(s, txt, size=font_size, color=color, bold=bold, margin=0)
    return s


def panel_header(slide, x, y, w, title, num=None, size=12):
    tw = w - (0.42 if num else 0)
    text(slide, x, y, tw, 0.30, title, size=size, bold=True, color=NAVY, anchor="b", margin=0)
    line(slide, x, y + 0.33, x + tw - 0.05, y + 0.33, color=NAVY, width=1.25)
    if num is not None:
        badge(slide, x + w - 0.34, y + 0.05, num)


def panel(slide, x, y, w, h, fill=PANEL):
    return rect(slide, x, y, w, h, fill=fill)


def key_box(slide, y, t, h=0.52, x=LEFT, w=CONTENT_W, size=11):
    s = rect(slide, x, y, w, h, fill=NAVY)
    tf = s.text_frame
    tf.word_wrap = True
    for side, m in (("left", 0.18), ("right", 0.18), ("top", 0.04), ("bottom", 0.04)):
        setattr(tf, f"margin_{side}", Inches(m))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _runs(p, t, size, WHITE, False, False, FONT)
    return s


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


# ------------------------------------------------------------------ tables
_TC_ORDER = ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D", "noFill", "solidFill", "gradFill",
             "blipFill", "pattFill", "grpFill", "headers", "extLst"]


def _tc_insert(tcPr, el):
    name = etree.QName(el).localname
    pos = _TC_ORDER.index(name)
    for i, child in enumerate(list(tcPr)):
        cn = etree.QName(child).localname
        if cn in _TC_ORDER and _TC_ORDER.index(cn) > pos:
            tcPr.insert(i, el)
            return
    tcPr.append(el)


def cell_border(cell, side, color=GRIDGREY, width=0.75, dash="solid"):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "lnL", "R": "lnR", "T": "lnT", "B": "lnB"}[side]
    for e in tcPr.findall(qn("a:" + tag)):
        tcPr.remove(e)
    ln = etree.Element(qn("a:" + tag), w=str(int(width * 12700)), cap="flat", cmpd="sng", algn="ctr")
    if color is None:
        etree.SubElement(ln, qn("a:noFill"))
    else:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val=dash)
    _tc_insert(tcPr, ln)


def table(slide, x, y, w, rows, col_w, row_h=0.2, size=8, align=None, name=None):
    """rows: list of dict(cells=[...], fill=None, color=DARK, bold=False, italic=False, line_top=None,
    line_bottom=None, size=None, fills={col: hex}, colors={col: hex}, bolds={col: bool}, h=None)."""
    nr, nc = len(rows), len(col_w)
    heights = [r.get("h", row_h) for r in rows]
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(heights)))
    if name:
        gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(k, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    scale = w / sum(col_w)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw * scale)
    align = align or (["l"] + ["r"] * (nc - 1))
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(heights[i])
        for j in range(nc):
            v = r["cells"][j] if j < len(r["cells"]) else ""
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.04)
            c.margin_top = c.margin_bottom = Inches(0.0)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            a = r.get("align", {}).get(j, align[j])
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[a]
            _runs(p, "" if v is None else str(v), r.get("size", size), r.get("colors", {}).get(j, r.get("color", DARK)),
                  r.get("bolds", {}).get(j, r.get("bold", False)), r.get("italic", False), FONT)
            for side in "LRTB":
                cell_border(c, side, color=None)
            if r.get("line_top"):
                cell_border(c, "T", color=r["line_top"], width=r.get("line_top_w", 0.75))
            if r.get("line_bottom"):
                cell_border(c, "B", color=r["line_bottom"], width=r.get("line_bottom_w", 0.75))
            fill = r.get("fills", {}).get(j, r.get("fill"))
            if fill:
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(fill)
            else:
                c.fill.background()
    return gf


# ------------------------------------------------------------------ charts
def _font_all(chart, size=8, color=DARK):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(color)


def _axis_line(axis, color=GRIDGREY, visible=True):
    if visible:
        axis.format.line.color.rgb = rgb(color)
        axis.format.line.width = Pt(0.75)
    else:
        axis.format.line.fill.background()


def _plot_area(chart):
    return chart._chartSpace.find(".//" + qn("c:plotArea"))


def manual_plot_layout(chart, x, y, w, h):
    pa = _plot_area(chart)
    lay = pa.find(qn("c:layout"))
    if lay is None:
        lay = etree.Element(qn("c:layout"))
        pa.insert(0, lay)
    for ch in list(lay):
        lay.remove(ch)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    etree.SubElement(ml, qn("c:layoutTarget"), val="inner")
    etree.SubElement(ml, qn("c:xMode"), val="edge")
    etree.SubElement(ml, qn("c:yMode"), val="edge")
    for k, v in (("x", x), ("y", y), ("w", w), ("h", h)):
        etree.SubElement(ml, qn("c:" + k), val=f"{v:.4f}")


def add_chart(slide, ctype, x, y, w, h, cats, series, colors, size=8, legend=None, labels=True,
              num_fmt="#,##0", label_pos=None, gap=60, overlap=None, val_axis=False, cat_axis=True,
              val_min=None, val_max=None, gridlines=False, label_color=None, name=None):
    cd = CategoryChartData(number_format=num_fmt)
    cd.categories = cats
    for sname, vals in series:
        cd.add_series(sname, vals)
    gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    if name:
        gf.name = name
    ch = gf.chart
    _font_all(ch, size)
    ch.has_title = False
    if legend:
        ch.has_legend = True
        ch.legend.position = {"t": XL_LEGEND_POSITION.TOP, "b": XL_LEGEND_POSITION.BOTTOM,
                              "r": XL_LEGEND_POSITION.RIGHT, "l": XL_LEGEND_POSITION.LEFT}[legend]
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    else:
        ch.has_legend = False
    plot = ch.plots[0]
    if hasattr(plot, "gap_width") and gap is not None:
        plot.gap_width = gap
    if overlap is not None and hasattr(plot, "overlap"):
        plot.overlap = overlap
    for i, s in enumerate(plot.series):
        col = colors[i % len(colors)]
        if ctype in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
            s.format.line.color.rgb = rgb(col)
            s.format.line.width = Pt(2)
            s.smooth = False
            if ctype == XL_CHART_TYPE.LINE_MARKERS:
                from pptx.enum.chart import XL_MARKER_STYLE
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = 6
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = rgb(col)
                s.marker.format.line.fill.background()
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(col)
            if hasattr(s, "invert_if_negative"):
                s.invert_if_negative = False
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_value = True
        dl.number_format = num_fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(size)
        if label_color:
            dl.font.color.rgb = rgb(label_color)
        if label_pos:
            dl.position = label_pos
    if ctype not in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE):
        va = ch.value_axis
        va.has_major_gridlines = gridlines
        if gridlines:
            va.major_gridlines.format.line.color.rgb = rgb("E3E6EA")
        va.visible = True
        if not val_axis:
            va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        va.tick_labels.font.size = Pt(size)
        if val_min is not None:
            va.minimum_scale = val_min
        if val_max is not None:
            va.maximum_scale = val_max
        _axis_line(va, visible=False)
        va.major_tick_mark = XL_TICK_MARK.NONE
        ca = ch.category_axis
        ca.visible = cat_axis
        ca.tick_labels.font.size = Pt(size)
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.has_major_gridlines = False
        _axis_line(ca, visible=True)
    return gf


def hide_small_labels(chart, series_idx, thr=0.005):
    """Blank the data labels of points whose absolute value is below thr (a '0%' sitting on the axis is noise)."""
    s = chart.plots[0].series[series_idx]
    for i, v in enumerate(s.values):
        if v is not None and abs(v) < thr:
            dl = s.points[i].data_label
            dl.has_text_frame = True
            dl.text_frame.text = ""


def color_points(chart, colors, series_idx=0):
    s = chart.plots[0].series[series_idx]
    for i, col in enumerate(colors):
        pt = s.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
    for dpt in s._element.findall(qn("c:dPt")):          # keep the colour for negative values
        if dpt.find(qn("c:invertIfNegative")) is None:
            inv = etree.Element(qn("c:invertIfNegative"), val="0")
            dpt.find(qn("c:idx")).addnext(inv)


def hide_series(chart, idx):
    s = chart.plots[0].series[idx]
    s.format.fill.background()
    s.format.line.fill.background()


def doughnut_hole(chart, pct=58):
    dn = chart._chartSpace.find(".//" + qn("c:doughnutChart"))
    hs = dn.find(qn("c:holeSize"))
    if hs is None:
        hs = etree.SubElement(dn, qn("c:holeSize"))
    hs.set("val", str(pct))


def series_labels_off(chart, idx):
    """Delete data labels for one series (plot-level labels stay on for the others)."""
    ser = chart.plots[0].series[idx]._element
    d = ser.find(qn("c:dLbls"))
    if d is not None:
        ser.remove(d)
    d = etree.Element(qn("c:dLbls"))
    etree.SubElement(d, qn("c:delete"), val="1")
    # insert after dPt / before trendline, errBars, cat, val
    anchor = None
    for tag in ("c:trendline", "c:errBars", "c:cat", "c:val"):
        anchor = ser.find(qn(tag))
        if anchor is not None:
            break
    anchor.addprevious(d)


def _txpr(size, color, bold=False):
    tx = etree.Element(qn("c:txPr"))
    bp = etree.SubElement(tx, qn("a:bodyPr"))
    etree.SubElement(tx, qn("a:lstStyle"))
    p = etree.SubElement(tx, qn("a:p"))
    ppr = etree.SubElement(p, qn("a:pPr"))
    d = etree.SubElement(ppr, qn("a:defRPr"), sz=str(int(size * 100)), b="1" if bold else "0")
    sf = etree.SubElement(d, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(d, qn("a:latin"), typeface=FONT)
    etree.SubElement(p, qn("a:endParaRPr"), lang="en-US")
    return tx


def to_combo(chart, line_idx, color, fmt="0%", size=8, secondary=True, label_color=None, val_min=None, val_max=None):
    """Move bar-chart series `line_idx` into a line chart on a hidden secondary axis."""
    pa = _plot_area(chart)
    bar = pa.find(qn("c:barChart"))
    sers = bar.findall(qn("c:ser"))
    ser = sers[line_idx]
    bar.remove(ser)
    for tag in ("c:invertIfNegative", "c:pictureOptions", "c:shape"):
        for e in ser.findall(qn(tag)):
            ser.remove(e)
    # series formatting: line only
    sp = ser.find(qn("c:spPr"))
    if sp is not None:
        ser.remove(sp)
    sp = etree.Element(qn("c:spPr"))
    ln = etree.SubElement(sp, qn("a:ln"), w=str(int(2.25 * 12700)), cap="rnd")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(ln, qn("a:round"))
    tx = ser.find(qn("c:tx"))
    (tx if tx is not None else ser.find(qn("c:order"))).addnext(sp)
    mk = etree.Element(qn("c:marker"))
    etree.SubElement(mk, qn("c:symbol"), val="circle")
    etree.SubElement(mk, qn("c:size"), val="5")
    msp = etree.SubElement(mk, qn("c:spPr"))
    msf = etree.SubElement(msp, qn("a:solidFill"))
    etree.SubElement(msf, qn("a:srgbClr"), val=color)
    mln = etree.SubElement(msp, qn("a:ln"))
    etree.SubElement(mln, qn("a:noFill"))
    sp.addnext(mk)
    old = ser.find(qn("c:dLbls"))
    if old is not None:
        ser.remove(old)
    dl = etree.Element(qn("c:dLbls"))
    etree.SubElement(dl, qn("c:numFmt"), formatCode=fmt, sourceLinked="0")
    dsp = etree.SubElement(dl, qn("c:spPr"))
    etree.SubElement(dsp, qn("a:noFill"))
    dln = etree.SubElement(dsp, qn("a:ln"))
    etree.SubElement(dln, qn("a:noFill"))
    dl.append(_txpr(size, label_color or color, bold=True))
    etree.SubElement(dl, qn("c:dLblPos"), val="t")
    for k, v in (("showLegendKey", "0"), ("showVal", "1"), ("showCatName", "0"), ("showSerName", "0"),
                 ("showPercent", "0"), ("showBubbleSize", "0")):
        etree.SubElement(dl, qn("c:" + k), val=v)
    cat = ser.find(qn("c:cat"))
    cat.addprevious(dl)
    val = ser.find(qn("c:val"))
    sm = etree.Element(qn("c:smooth"), val="0")
    val.addnext(sm)
    # line chart element directly after the bar chart
    lc = etree.Element(qn("c:lineChart"))
    etree.SubElement(lc, qn("c:grouping"), val="standard")
    etree.SubElement(lc, qn("c:varyColors"), val="0")
    lc.append(ser)
    etree.SubElement(lc, qn("c:marker"), val="1")
    etree.SubElement(lc, qn("c:axId"), val="900001")
    etree.SubElement(lc, qn("c:axId"), val="900002")
    bar.addnext(lc)
    # secondary axes (hidden) placed after the existing axes
    last_ax = pa.findall(qn("c:valAx"))[-1]
    cat_ax = etree.Element(qn("c:catAx"))
    etree.SubElement(cat_ax, qn("c:axId"), val="900001")
    sc = etree.SubElement(cat_ax, qn("c:scaling"))
    etree.SubElement(sc, qn("c:orientation"), val="minMax")
    etree.SubElement(cat_ax, qn("c:delete"), val="1")
    etree.SubElement(cat_ax, qn("c:axPos"), val="b")
    etree.SubElement(cat_ax, qn("c:majorTickMark"), val="none")
    etree.SubElement(cat_ax, qn("c:minorTickMark"), val="none")
    etree.SubElement(cat_ax, qn("c:tickLblPos"), val="nextTo")
    etree.SubElement(cat_ax, qn("c:crossAx"), val="900002")
    etree.SubElement(cat_ax, qn("c:crosses"), val="autoZero")
    etree.SubElement(cat_ax, qn("c:auto"), val="1")
    etree.SubElement(cat_ax, qn("c:lblAlgn"), val="ctr")
    etree.SubElement(cat_ax, qn("c:lblOffset"), val="100")
    etree.SubElement(cat_ax, qn("c:noMultiLvlLbl"), val="0")
    val_ax = etree.Element(qn("c:valAx"))
    etree.SubElement(val_ax, qn("c:axId"), val="900002")
    sc = etree.SubElement(val_ax, qn("c:scaling"))
    etree.SubElement(sc, qn("c:orientation"), val="minMax")
    if val_max is not None:
        etree.SubElement(sc, qn("c:max"), val=str(val_max))
    if val_min is not None:
        etree.SubElement(sc, qn("c:min"), val=str(val_min))
    etree.SubElement(val_ax, qn("c:delete"), val="0")
    etree.SubElement(val_ax, qn("c:axPos"), val="r")
    etree.SubElement(val_ax, qn("c:numFmt"), formatCode=fmt, sourceLinked="0")
    etree.SubElement(val_ax, qn("c:majorTickMark"), val="none")
    etree.SubElement(val_ax, qn("c:minorTickMark"), val="none")
    etree.SubElement(val_ax, qn("c:tickLblPos"), val="none")
    vsp = etree.SubElement(val_ax, qn("c:spPr"))
    vln = etree.SubElement(vsp, qn("a:ln"))
    etree.SubElement(vln, qn("a:noFill"))
    etree.SubElement(val_ax, qn("c:crossAx"), val="900001")
    etree.SubElement(val_ax, qn("c:crosses"), val="max")
    etree.SubElement(val_ax, qn("c:crossBetween"), val="between")
    last_ax.addnext(cat_ax)
    cat_ax.addnext(val_ax)
    return lc


def set_cat_reverse(chart):
    ax = chart._chartSpace.find(".//" + qn("c:catAx"))
    sc = ax.find(qn("c:scaling"))
    o = sc.find(qn("c:orientation"))
    o.set("val", "maxMin")
    # keep value axis at the bottom
    vax = chart._chartSpace.find(".//" + qn("c:valAx"))
    cr = vax.find(qn("c:crosses"))
    if cr is not None:
        cr.set("val", "max")
